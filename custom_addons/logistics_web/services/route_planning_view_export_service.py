import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from odoo import fields
from odoo.exceptions import ValidationError


class RoutePlanningViewExportError(ValidationError):
    def __init__(self, error_code, message):
        self.error_code = error_code
        super().__init__(message)


class RoutePlanningViewExportService:
    DOWNLOAD_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    NOTE_FILL = PatternFill(fill_type="solid", fgColor="FFF3CD")
    NOTE_FONT = Font(bold=True, color="9A6700")
    HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
    HEADER_FONT = Font(bold=True, color="1F2937")
    ERROR_FILL = PatternFill(fill_type="solid", fgColor="FDE2E1")

    ORDER_DETAIL_HEADERS = [
        "批次号",
        "排线序号",
        "运单号",
        "销售订单号",
        "订单号",
        "客户名称",
        "订单备注",
        "收货地址",
        "整件数",
        "散件数",
        "总重量_kg",
        "总体积_m3",
        "集货位",
        "司机姓名",
        "车牌号",
    ]
    STORE_DETAIL_HEADERS = [
        "批次号",
        "排线序号",
        "运单号",
        "客户名称",
        "收货地址",
        "联系人",
        "联系电话",
        "集货位",
        "不收货时间段",
        "门店备注",
        "整件数",
        "散件数",
        "总重量_kg",
        "总体积_m3",
        "经度",
        "纬度",
        "司机姓名",
        "车牌号",
    ]
    GOODS_DETAIL_HEADERS = [
        "批次号",
        "排线序号",
        "运单号",
        "销售订单号",
        "客户名称",
        "行号",
        "商品名称",
        "商品编号",
        "商品条码",
        "包装规格",
        "发货数量",
        "小单位销售单位",
        "小单位销售单价",
        "金额",
        "收货地址",
        "联系人",
        "联系电话",
        "业务员",
        "业务员联系方式",
    ]

    EXPORT_SPECS = {
        "order_detail": {
            "sheet_name": "订单详情",
            "file_name": "route_order_details_{batch_no}.xlsx",
            "headers": ORDER_DETAIL_HEADERS,
            "critical_headers": ["运单号", "客户名称", "收货地址"],
        },
        "store_detail": {
            "sheet_name": "门店详情",
            "file_name": "route_store_details_{batch_no}.xlsx",
            "headers": STORE_DETAIL_HEADERS,
            "critical_headers": ["运单号", "客户名称", "收货地址", "经度", "纬度"],
        },
        "goods_detail": {
            "sheet_name": "门店货物信息三联单",
            "file_name": "route_goods_details_{batch_no}.xlsx",
            "headers": GOODS_DETAIL_HEADERS,
            "critical_headers": ["运单号", "客户名称", "商品名称"],
        },
    }

    @classmethod
    def export_view_by_batch(cls, env, *, export_key="", batch_id=0, file_locale="zh_CN"):
        del file_locale
        cls._ensure_access(env)
        spec = cls.EXPORT_SPECS.get((export_key or "").strip())
        if not spec:
            raise RoutePlanningViewExportError("ROUTE_VIEW_EXPORT_KEY_INVALID", f"Unsupported export key: {export_key}")
        batch = env["logistics.route.planning.batch"].sudo().browse(int(batch_id or 0)).exists()
        if not batch:
            raise RoutePlanningViewExportError("ROUTE_VIEW_EXPORT_BATCH_MISSING", "排线批次不存在。")
        rows, anomaly_summary = cls._build_rows(env, export_key=export_key, batch=batch)
        workbook_bytes = cls._build_workbook_bytes(spec, batch=batch, rows=rows, anomaly_summary=anomaly_summary)
        return {
            "file_name": spec["file_name"].format(batch_no=(batch.batch_no or f"batch_{batch.id}")),
            "content_type": cls.DOWNLOAD_CONTENT_TYPE,
            "file_bytes": workbook_bytes,
        }

    @classmethod
    def _ensure_access(cls, env):
        if env.user.has_group("logistics_dispatch.group_logistics_export_user"):
            return
        raise RoutePlanningViewExportError("EXPORT_PERMISSION_DENIED", "当前用户没有排线视图导出权限。")

    @classmethod
    def _build_rows(cls, env, *, export_key, batch):
        if export_key == "order_detail":
            return cls._build_order_rows(batch)
        if export_key == "store_detail":
            return cls._build_store_rows(batch)
        if export_key == "goods_detail":
            return cls._build_goods_rows(env, batch)
        raise RoutePlanningViewExportError("ROUTE_VIEW_EXPORT_KEY_INVALID", f"Unsupported export key: {export_key}")

    @classmethod
    def _build_order_rows(cls, batch):
        rows = []
        anomaly_fields = set()
        for line in batch._get_related_waybills().mapped("order_line_ids").sorted(
            key=lambda rec: (rec.stop_seq_in_waybill, rec.sales_order_no or rec.order_no or "", rec.sequence, rec.id)
        ):
            row = {
                "批次号": line.batch_no or batch.batch_no or "",
                "排线序号": line.stop_seq_in_waybill or "",
                "运单号": line.waybill_id.name or "",
                "销售订单号": line.sales_order_no or "",
                "订单号": line.order_no or "",
                "客户名称": line.customer_name_snapshot or line.customer_name or "",
                "订单备注": line.order_remark or "",
                "收货地址": line.address_full_snapshot or "",
                "整件数": line.whole_package_count or "",
                "散件数": line.loose_package_count or "",
                "总重量_kg": line.weight_summary or "",
                "总体积_m3": line.volume_summary or "",
                "集货位": line.gathering_location or "",
                "司机姓名": line.driver_name_snapshot or "",
                "车牌号": line.vehicle_no_snapshot or "",
            }
            row_anomalies = cls._collect_missing_fields(row, cls.EXPORT_SPECS["order_detail"]["critical_headers"])
            if row_anomalies:
                anomaly_fields.update(row_anomalies)
            rows.append({"values": row, "anomalies": row_anomalies})
        return rows, cls._build_anomaly_summary(rows, anomaly_fields)

    @classmethod
    def _build_store_rows(cls, batch):
        rows = []
        anomaly_fields = set()
        for stop_line in batch.stop_line_ids.sorted(key=lambda rec: (rec.stop_seq, rec.id)):
            row = {
                "批次号": stop_line.batch_no or batch.batch_no or "",
                "排线序号": stop_line.stop_seq or "",
                "运单号": stop_line.waybill_no or "",
                "客户名称": stop_line.store_name or "",
                "收货地址": stop_line.address_detail or "",
                "联系人": stop_line.contact_name_snapshot or "",
                "联系电话": stop_line.contact_phone or "",
                "集货位": stop_line.gathering_location_summary or "",
                "不收货时间段": stop_line.no_receive_time_slots_text or "",
                "门店备注": stop_line.cargo_summary or "",
                "整件数": stop_line.whole_package_count_summary or stop_line.package_count_total or "",
                "散件数": stop_line.loose_package_count_summary or "",
                "总重量_kg": stop_line.weight_summary_total or stop_line.goods_weight_total or "",
                "总体积_m3": stop_line.volume_summary_total or stop_line.goods_volume_total or "",
                "经度": stop_line.longitude or "",
                "纬度": stop_line.latitude or "",
                "司机姓名": stop_line.driver_name or "",
                "车牌号": stop_line.vehicle_no or "",
            }
            row_anomalies = cls._collect_missing_fields(row, cls.EXPORT_SPECS["store_detail"]["critical_headers"])
            if row_anomalies:
                anomaly_fields.update(row_anomalies)
            rows.append({"values": row, "anomalies": row_anomalies})
        return rows, cls._build_anomaly_summary(rows, anomaly_fields)

    @classmethod
    def _build_goods_rows(cls, env, batch):
        goods_model = env["logistics.dispatch.waybill.customer.goods.line"].sudo()
        goods_lines = goods_model.search(
            [("waybill_id", "in", batch._get_related_waybills().ids)],
            order="waybill_id asc, stop_seq_in_waybill asc, sequence asc, id asc",
        )
        rows = []
        anomaly_fields = set()
        for line in goods_lines:
            row = {
                "批次号": line.batch_no or batch.batch_no or "",
                "排线序号": line.stop_seq_in_waybill or "",
                "运单号": line.waybill_id.name or "",
                "销售订单号": line.sales_order_no or "",
                "客户名称": line.customer_name or "",
                "行号": line.sequence or "",
                "商品名称": line.goods_name or "",
                "商品编号": line.goods_code or "",
                "商品条码": line.barcode_snapshot or "",
                "包装规格": line.specification or "",
                "发货数量": line.delivery_qty_text or line.quantity or "",
                "小单位销售单位": line.small_unit_name or "",
                "小单位销售单价": line.small_unit_price or "",
                "金额": line.amount or "",
                "收货地址": line.address_full_snapshot or "",
                "联系人": line.contact_name_snapshot or "",
                "联系电话": line.contact_phone_snapshot or "",
                "业务员": line.salesperson_name_snapshot or "",
                "业务员联系方式": line.salesperson_phone_snapshot or "",
            }
            row_anomalies = cls._collect_missing_fields(row, cls.EXPORT_SPECS["goods_detail"]["critical_headers"])
            if row_anomalies:
                anomaly_fields.update(row_anomalies)
            rows.append({"values": row, "anomalies": row_anomalies})
        return rows, cls._build_anomaly_summary(rows, anomaly_fields)

    @classmethod
    def _collect_missing_fields(cls, row, critical_fields):
        anomalies = []
        for header in critical_fields:
            if row.get(header) in (None, "", False):
                anomalies.append(header)
        return anomalies

    @classmethod
    def _build_anomaly_summary(cls, rows, anomaly_fields):
        row_issue_count = sum(1 for row in rows if row["anomalies"])
        if not row_issue_count:
            return {"has_warning": False, "message": "无异常提示。"}
        focus_fields = ", ".join(sorted(anomaly_fields))
        return {
            "has_warning": True,
            "message": f"异常提醒：共有 {row_issue_count} 行需要人工复核，重点检查 {focus_fields}。",
        }

    @classmethod
    def _build_workbook_bytes(cls, spec, *, batch, rows, anomaly_summary):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = spec["sheet_name"]

        headers = spec["headers"]
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        note_cell = sheet.cell(
            row=1,
            column=1,
            value=(
                f"批次号：{batch.batch_no or batch.id}；"
                f"配送日期：{fields.Date.to_string(batch.delivery_date) or ''}；"
                f"{anomaly_summary['message']}"
            ),
        )
        note_cell.fill = cls.NOTE_FILL
        note_cell.font = cls.NOTE_FONT
        note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=2, column=index, value=header)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_index, row_item in enumerate(rows, start=3):
            values = row_item["values"]
            anomalies = set(row_item["anomalies"])
            for col_index, header in enumerate(headers, start=1):
                cell = sheet.cell(row=row_index, column=col_index, value=values.get(header, ""))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if header in anomalies:
                    cell.fill = cls.ERROR_FILL

        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = f"A2:{cls._column_letter(len(headers))}{max(sheet.max_row, 2)}"
        for index, _header in enumerate(headers, start=1):
            sheet.column_dimensions[cls._column_letter(index)].width = 18 if index < 8 else 22

        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    @classmethod
    def _column_letter(cls, index):
        letters = ""
        while index:
            index, remainder = divmod(index - 1, 26)
            letters = chr(65 + remainder) + letters
        return letters
