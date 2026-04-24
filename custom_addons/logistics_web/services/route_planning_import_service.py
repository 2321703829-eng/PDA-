import io
from collections import defaultdict

from openpyxl import load_workbook

from odoo import fields

from .waybill_standard_import_service_v2 import WaybillStandardImportService


class RoutePlanningImportService(WaybillStandardImportService):
    TEMPLATE_CODE = "TSL-IMPORT-ROUTE-PLANNING-V1"
    TEMPLATE_VERSION = "v1"
    LEGACY_TEMPLATE_CODES = set()
    LEGACY_TEMPLATE_VERSIONS = set()
    OBJECT_TYPE = "route_planning"
    OBJECT_TYPE_LABEL = "排线用数据导入"
    TEMPLATE_FILE_NAME = "TSL-IMPORT-ROUTE-PLANNING-V1.xlsx"
    TEMPLATE_VARIANTS = {
        "en_US": {
            "label": "Route Planning Template (English headers)",
            "file_name": "TSL-IMPORT-ROUTE-PLANNING-V1.xlsx",
            "field_locale": "en_US",
        },
        "zh_CN": {
            "label": "排线模板（中文列头）",
            "file_name": "TSL-IMPORT-ROUTE-PLANNING-V1.zh_CN.xlsx",
            "field_locale": "zh_CN",
        },
    }
    DEFAULT_TEMPLATE_LOCALE = "zh_CN"
    INPUT_MODE = "route_planning_single_sheet"
    PRIMARY_SHEET = {
        "key": "rows",
        "sheet_name": "路线输入表",
        "fields": [
            "batch_no",
            "waybill_no",
            "stop_seq",
            "store_name",
            "longitude",
            "latitude",
            "address_detail",
            "store_contact_phone",
            "cargo_summary",
            "delivery_date",
            "driver_name",
            "driver_phone",
            "vehicle_no",
            "supplier_name",
            "warehouse_name",
        ],
        "required_fields": {
            "batch_no",
            "waybill_no",
            "stop_seq",
            "store_name",
            "longitude",
            "latitude",
            "address_detail",
            "delivery_date",
        },
    }
    SHEETS = [PRIMARY_SHEET]
    FIELD_LABELS = {
        "batch_no": "批次号",
        "waybill_no": "运单号",
        "stop_seq": "停靠点顺序",
        "store_name": "门店名称",
        "longitude": "经度",
        "latitude": "纬度",
        "address_detail": "详细地址",
        "store_contact_phone": "门店联系电话",
        "cargo_summary": "货物信息",
        "delivery_date": "配送日期",
        "driver_name": "司机姓名",
        "driver_phone": "司机电话",
        "vehicle_no": "车牌号",
        "supplier_name": "供应商名称",
        "warehouse_name": "仓库名称",
        "template_file": "模板文件",
    }
    FIELD_ALIASES = {
        "batch_no": ["batch_no", "batch no"],
        "waybill_no": ["waybill_no", "waybill no"],
        "stop_seq": ["stop_seq", "stop seq", "停靠顺序", "送货顺序"],
        "store_name": ["store_name", "store name", "门店", "门店名字"],
        "longitude": ["lng", "lon"],
        "latitude": ["lat"],
        "address_detail": ["address_detail", "address", "详细地址", "地址"],
        "store_contact_phone": ["store_contact_phone", "contact_phone", "门店电话", "联系电话"],
        "cargo_summary": ["cargo_summary", "goods_info", "货物摘要"],
        "delivery_date": ["delivery_date", "配送时间"],
        "driver_name": ["driver_name"],
        "driver_phone": ["driver_phone"],
        "vehicle_no": ["vehicle_no", "plate_no", "车牌"],
        "supplier_name": ["supplier_name", "供应商"],
        "warehouse_name": ["warehouse_name", "仓库"],
    }
    WARNING_FIELDS = {
        "store_contact_phone",
        "cargo_summary",
        "driver_name",
        "driver_phone",
        "vehicle_no",
        "supplier_name",
        "warehouse_name",
    }

    @classmethod
    def _build_template_download_url(cls, *, template_locale):
        return (
            "/api/admin/logistics/imports/route-planning/template/download"
            f"?template_code={cls.TEMPLATE_CODE}"
            f"&template_version={cls.TEMPLATE_VERSION}"
            f"&template_locale={template_locale}"
        )

    @classmethod
    def _empty_source_data(cls):
        return {"rows": [], "input_mode": cls.INPUT_MODE}

    @classmethod
    def _build_template_samples(cls):
        return {
            "rows": [
                {
                    "batch_no": "PL-BT-20260424-01",
                    "waybill_no": "PL-WB-20260424-01",
                    "stop_seq": "1",
                    "store_name": "世纪大道门店",
                    "longitude": "121.54420",
                    "latitude": "31.22110",
                    "address_detail": "上海市浦东新区世纪大道100号",
                    "store_contact_phone": "13910000001",
                    "cargo_summary": "饮料 20 件；常温",
                    "delivery_date": "2026-04-24",
                    "driver_name": "张师傅",
                    "driver_phone": "13910000099",
                    "vehicle_no": "沪A12345",
                    "supplier_name": "浦东配送组",
                    "warehouse_name": "浦东仓",
                },
                {
                    "batch_no": "PL-BT-20260424-01",
                    "waybill_no": "PL-WB-20260424-01",
                    "stop_seq": "2",
                    "store_name": "张江门店",
                    "longitude": "121.59910",
                    "latitude": "31.20320",
                    "address_detail": "上海市浦东新区张江路200号",
                    "store_contact_phone": "13910000002",
                    "cargo_summary": "常温货 12 件",
                    "delivery_date": "2026-04-24",
                    "driver_name": "张师傅",
                    "driver_phone": "13910000099",
                    "vehicle_no": "沪A12345",
                    "supplier_name": "浦东配送组",
                    "warehouse_name": "浦东仓",
                },
            ]
        }

    @classmethod
    def precheck(cls, env, raw_bytes, *, filename="", template_code="", template_version=""):
        template_errors = cls._validate_template_identity(template_code, template_version)
        parsed_data, parse_errors = cls._parse_workbook(raw_bytes)
        errors = template_errors + parse_errors
        warnings = []
        if not parse_errors:
            row_errors, row_warnings = cls._validate_route_rows(parsed_data)
            errors.extend(row_errors)
            warnings.extend(row_warnings)

        result = cls._build_route_precheck_result(parsed_data, errors, warnings)
        source_file = cls._create_source_file(
            env,
            raw_bytes=raw_bytes,
            filename=filename or cls.TEMPLATE_FILE_NAME,
        )
        task = cls._create_precheck_task(env, source_file=source_file, result=result)
        row_line_map, row_data_map = cls._create_precheck_task_lines(task, parsed_data, result["errors"])
        cls._create_task_error_lines(task, result["errors"], row_line_map=row_line_map, row_data_map=row_data_map)
        error_report_url = cls._build_task_error_report_url(task) if result["errors"] else False
        result["task_no"] = task.task_no
        result["import_batch_no"] = task.task_no
        result["object_type"] = task.object_type
        result["object_type_label"] = cls._get_selection_label(task._fields["object_type"].selection, task.object_type)
        result["status"] = task.status
        result["status_label"] = cls._get_selection_label(task._fields["status"].selection, task.status)
        result["summary_message"] = task.summary_message
        result["precheck_token"] = False
        result["error_report_url"] = error_report_url
        result["error_report"] = {
            "download_ready": bool(result["errors"]),
            "download_url": error_report_url,
        }
        result["source_file"] = cls._build_source_file_payload(source_file)
        return result

    @classmethod
    def _parse_workbook(cls, raw_bytes):
        if not raw_bytes:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name="模板文件",
                    row_no=0,
                    field_code="template_file",
                    error_code="PRECHECK_PARSE_FAILED",
                    error_message="模板文件内容为空，无法执行排线导入预校验。",
                )
            ]
        try:
            workbook = load_workbook(io.BytesIO(raw_bytes), data_only=True)
        except Exception:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name="模板文件",
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_FILE_TYPE_INVALID",
                    error_message="排线导入仅支持 .xlsx 文件。",
                )
            ]

        target_sheet_name = cls._find_target_sheet_name(workbook)
        if not target_sheet_name:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name="模板文件",
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_SHEET_MISSING",
                    error_message="未识别到可用的排线输入表，请至少提供命中排线表头规则的工作表。",
                )
            ]

        worksheet = workbook[target_sheet_name]
        parsed_rows, errors = cls._parse_sheet_rows(worksheet, cls.PRIMARY_SHEET, sheet_name_override=target_sheet_name)
        if errors:
            return cls._empty_source_data(), errors
        return {"rows": parsed_rows, "input_mode": cls.INPUT_MODE}, []

    @classmethod
    def _find_target_sheet_name(cls, workbook):
        preferred_sheet_name = cls.PRIMARY_SHEET["sheet_name"]
        if preferred_sheet_name in workbook.sheetnames:
            return preferred_sheet_name
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows = list(worksheet.iter_rows(values_only=True, min_row=1, max_row=1))
            if not rows:
                continue
            header_values = [cls._cell_text(value) for value in rows[0]]
            _header_map, header_errors = cls._build_header_map(cls.PRIMARY_SHEET, header_values)
            if not header_errors:
                return sheet_name
        return False

    @classmethod
    def _validate_route_rows(cls, source_data):
        errors = []
        warnings = []
        waybill_snapshot_map = {}
        stop_seq_keys = set()
        duplicate_location_keys = set()
        for row in source_data.get("rows", []):
            errors.extend(cls._validate_required_fields(row, cls.PRIMARY_SHEET["required_fields"]))
            errors.extend(cls._validate_positive_integer(row, "stop_seq"))
            errors.extend(cls._validate_geo_field(row, "longitude", min_value=-180.0, max_value=180.0))
            errors.extend(cls._validate_geo_field(row, "latitude", min_value=-90.0, max_value=90.0))
            errors.extend(cls._validate_date_field(row, "delivery_date"))

            if row.get("store_contact_phone") and not cls._is_valid_phone(row.get("store_contact_phone")):
                warnings.append(
                    cls._make_warning(row, "store_contact_phone", "PHONE_FORMAT_REVIEW", "门店联系电话格式异常，请人工复查。")
                )
            if row.get("driver_phone") and not cls._is_valid_phone(row.get("driver_phone")):
                warnings.append(
                    cls._make_warning(row, "driver_phone", "PHONE_FORMAT_REVIEW", "司机电话格式异常，请人工复查。")
                )

            snapshot_key = row.get("waybill_no")
            snapshot = {
                "batch_no": row.get("batch_no"),
                "delivery_date": row.get("delivery_date"),
            }
            if snapshot_key in waybill_snapshot_map and waybill_snapshot_map[snapshot_key] != snapshot:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="waybill_no",
                        error_code="WAYBILL_SNAPSHOT_CONFLICT",
                        error_message="同一运单号下的批次号和配送日期必须保持一致。",
                    )
                )
            else:
                waybill_snapshot_map[snapshot_key] = snapshot

            stop_key = (row.get("waybill_no"), row.get("stop_seq"))
            if all(stop_key):
                if stop_key in stop_seq_keys:
                    errors.append(
                        cls._make_error(
                            sheet_name=row["_sheet_name"],
                            row_no=row["_source_row_no"],
                            field_code="stop_seq",
                            error_code="STOP_SEQ_DUPLICATED",
                            error_message="同一运单号下的停靠点顺序不允许重复。",
                        )
                    )
                else:
                    stop_seq_keys.add(stop_key)

            location_key = (row.get("waybill_no"), row.get("store_name"), row.get("address_detail"))
            if all(location_key):
                if location_key in duplicate_location_keys:
                    warnings.append(
                        cls._make_warning(
                            row,
                            "store_name",
                            "STOP_LOCATION_DUPLICATED",
                            "同一运单号下出现重复门店名称和地址，请人工复查。",
                        )
                    )
                else:
                    duplicate_location_keys.add(location_key)

            for field_code in cls.WARNING_FIELDS:
                if not row.get(field_code):
                    warnings.append(
                        cls._make_warning(
                            row,
                            field_code,
                            "FIELD_REVIEW_RECOMMENDED",
                            f"{cls.FIELD_LABELS[field_code]}为空，建议人工复查后再继续排线。",
                        )
                    )
        return errors, warnings

    @classmethod
    def _validate_positive_integer(cls, row, field_code):
        raw_value = (row.get(field_code) or "").strip()
        if not raw_value:
            return []
        try:
            parsed_value = int(raw_value)
        except Exception:
            parsed_value = 0
        if parsed_value > 0:
            return []
        return [
            cls._make_error(
                sheet_name=row["_sheet_name"],
                row_no=row["_source_row_no"],
                field_code=field_code,
                error_code="FIELD_FORMAT_INVALID",
                error_message=f"{cls.FIELD_LABELS[field_code]}必须为正整数。",
            )
        ]

    @classmethod
    def _validate_geo_field(cls, row, field_code, *, min_value, max_value):
        raw_value = (row.get(field_code) or "").strip()
        if not raw_value:
            return []
        try:
            parsed_value = float(raw_value)
        except Exception:
            parsed_value = None
        if parsed_value is not None and min_value <= parsed_value <= max_value:
            return []
        return [
            cls._make_error(
                sheet_name=row["_sheet_name"],
                row_no=row["_source_row_no"],
                field_code=field_code,
                error_code="FIELD_FORMAT_INVALID",
                error_message=f"{cls.FIELD_LABELS[field_code]}必须为合法坐标数值。",
            )
        ]

    @classmethod
    def _validate_date_field(cls, row, field_code):
        raw_value = (row.get(field_code) or "").strip()
        if not raw_value:
            return []
        try:
            fields.Date.from_string(raw_value)
            return []
        except Exception:
            return [
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code="FIELD_FORMAT_INVALID",
                    error_message=f"{cls.FIELD_LABELS[field_code]}必须为合法日期。",
                )
            ]

    @classmethod
    def _is_valid_phone(cls, raw_value):
        normalized = "".join(ch for ch in (raw_value or "") if ch.isdigit())
        return 7 <= len(normalized) <= 20

    @classmethod
    def _make_warning(cls, row, field_code, warning_code, warning_message):
        return {
            "sheet_name": row["_sheet_name"],
            "row_no": row["_source_row_no"],
            "field_code": field_code,
            "field_label": cls.FIELD_LABELS.get(field_code, field_code),
            "warning_code": warning_code,
            "warning_message": warning_message,
        }

    @classmethod
    def _build_route_precheck_result(cls, source_data, errors, warnings):
        result = cls._build_precheck_result(source_data, errors)
        warning_row_keys = {
            (warning.get("sheet_name"), warning.get("row_no"))
            for warning in warnings
            if warning.get("row_no")
        }
        result["warning_count"] = len(warnings)
        result["warning_row_count"] = len(warning_row_keys)
        result["has_review_warning"] = bool(warnings)
        result["warnings"] = sorted(
            warnings,
            key=lambda item: (
                item.get("sheet_name", ""),
                item.get("row_no", 0),
                item.get("field_code", ""),
                item.get("warning_code", ""),
            ),
        )
        return result

    @classmethod
    def _build_row_business_key(cls, row):
        detail_parts = [
            row.get("batch_no"),
            row.get("waybill_no"),
            row.get("stop_seq"),
            row.get("store_name"),
        ]
        detail = " / ".join(filter(None, detail_parts)) or "--"
        business_key = f"{row.get('_sheet_name') or cls.PRIMARY_SHEET['sheet_name']}@{row.get('_source_row_no') or 0} {detail}"
        return business_key[:128]

    @classmethod
    def _execute_task_import(cls, env, task, source_data):
        task_line_map = {line.business_key: line for line in task.task_line_ids}
        error_line_model = env["logistics.import.error.line"].sudo()
        draft_batch_model = env["logistics.route.planning.batch"].sudo()
        stop_line_model = env["logistics.route.planning.stop.line"].sudo()

        success_count = 0
        fail_count = 0
        written_batch_keys = set()
        created_batch_keys = set()
        rows_by_batch = defaultdict(list)
        for row in sorted(
            source_data.get("rows", []),
            key=lambda item: (
                item.get("delivery_date") or "",
                item.get("batch_no") or "",
                cls._safe_int(item.get("stop_seq"), default=0),
                item.get("_source_row_no") or 0,
            ),
        ):
            rows_by_batch[(row.get("delivery_date"), row.get("batch_no"))].append(row)

        for (delivery_date, batch_no), rows in rows_by_batch.items():
            try:
                batch = draft_batch_model.search(
                    [("batch_no", "=", batch_no), ("delivery_date", "=", delivery_date)],
                    limit=1,
                )
                first_row = rows[0]
                batch_vals = {
                    "batch_no": batch_no,
                    "delivery_date": fields.Date.from_string(delivery_date),
                    "driver_name": first_row.get("driver_name") or False,
                    "driver_phone": first_row.get("driver_phone") or False,
                    "vehicle_no": first_row.get("vehicle_no") or False,
                    "supplier_name": first_row.get("supplier_name") or False,
                    "warehouse_name": first_row.get("warehouse_name") or False,
                    "import_task_id": task.id,
                }
                if batch:
                    batch.write(batch_vals)
                    batch.stop_line_ids.unlink()
                else:
                    batch = draft_batch_model.create(batch_vals)
                    created_batch_keys.add((delivery_date, batch_no))
                written_batch_keys.add((delivery_date, batch_no))
            except Exception as exc:
                for row in rows:
                    task_line = task_line_map.get(cls._build_row_business_key(row))
                    fail_count += 1
                    cls._update_task_line(task_line, status="failed", message=f"排线批次写入失败：{exc}")
                    error_line_model.create(
                        cls._make_runtime_error_line_vals(
                            task=task,
                            task_line=task_line,
                            row=row,
                            field_name="batch_no",
                            error_message=f"排线批次写入失败：{exc}",
                        )
                    )
                continue

            for row in rows:
                task_line = task_line_map.get(cls._build_row_business_key(row))
                try:
                    stop_line = stop_line_model.create(
                        {
                            "batch_id": batch.id,
                            "import_task_id": task.id,
                            "import_task_line_id": task_line.id if task_line else False,
                            "waybill_no": row.get("waybill_no"),
                            "stop_seq": cls._safe_int(row.get("stop_seq"), default=0),
                            "store_name": row.get("store_name"),
                            "longitude": cls._safe_float(row.get("longitude"), default=0.0),
                            "latitude": cls._safe_float(row.get("latitude"), default=0.0),
                            "address_detail": row.get("address_detail"),
                            "contact_phone": row.get("store_contact_phone") or False,
                            "cargo_summary": row.get("cargo_summary") or False,
                            "driver_name": row.get("driver_name") or False,
                            "driver_phone": row.get("driver_phone") or False,
                            "vehicle_no": row.get("vehicle_no") or False,
                            "supplier_name": row.get("supplier_name") or False,
                            "warehouse_name": row.get("warehouse_name") or False,
                        }
                    )
                    success_count += 1
                    cls._update_task_line(
                        task_line,
                        status="success",
                        message="排线停靠点导入成功。",
                        target_model="logistics.route.planning.stop.line",
                        target_res_id=stop_line.id,
                    )
                except Exception as exc:
                    fail_count += 1
                    cls._update_task_line(task_line, status="failed", message=f"排线停靠点写入失败：{exc}")
                    error_line_model.create(
                        cls._make_runtime_error_line_vals(
                            task=task,
                            task_line=task_line,
                            row=row,
                            field_name="stop_seq",
                            error_message=f"排线停靠点写入失败：{exc}",
                        )
                    )

        task_status = "success"
        if fail_count and success_count:
            task_status = "partial_failed"
        elif fail_count and not success_count:
            task_status = "failed"
        return {
            "task_status": task_status,
            "success_count": success_count,
            "fail_count": fail_count,
            "skipped_count": 0,
            "created_route_batch_count": len(created_batch_keys),
            "written_route_batch_count": len(written_batch_keys),
            "written_route_stop_count": success_count,
            "summary_message": f"排线导入完成，共 {task.total_count} 行，成功 {success_count} 行，失败 {fail_count} 行。",
        }

    @classmethod
    def _build_task_result_payload(cls, task):
        task.ensure_one()
        counts = cls._get_task_line_status_counts(task)
        route_batch_model = task.env["logistics.route.planning.batch"].sudo()
        route_stop_model = task.env["logistics.route.planning.stop.line"].sudo()
        route_batches = route_batch_model.search([("import_task_id", "=", task.id)])
        route_stops = route_stop_model.search([("import_task_id", "=", task.id)])
        source_file = task.source_file_id
        return {
            "task_no": task.task_no,
            "import_batch_no": task.task_no,
            "object_type": task.object_type,
            "object_type_label": cls._get_selection_label(task._fields["object_type"].selection, task.object_type),
            "status": task.status,
            "status_label": cls._get_selection_label(task._fields["status"].selection, task.status),
            "total_count": task.total_count,
            "success_count": task.success_count,
            "fail_count": task.fail_count,
            "summary_message": task.summary_message,
            "template_code": cls.TEMPLATE_CODE,
            "template_version": cls.TEMPLATE_VERSION,
            "file_name": source_file.file_name if source_file else False,
            "total_row_count": task.total_count,
            "passed_row_count": task.success_count,
            "failed_row_count": task.fail_count,
            "started_at": fields.Datetime.to_string(task.started_at) if task.started_at else False,
            "finished_at": fields.Datetime.to_string(task.finished_at) if task.finished_at else False,
            "created_route_batch_count": len(route_batches),
            "written_route_stop_count": len(route_stops),
            "skipped_record_count": counts["skipped"],
            "failed_record_count": task.fail_count,
            "source_file": cls._build_source_file_payload(source_file) if source_file else False,
            "operator": cls._build_operator_payload(task.operator_id),
            "error_report": {
                "download_ready": bool(task.error_line_ids),
                "download_url": cls._build_task_error_report_url(task) if task.error_line_ids else False,
            },
            "error_report_url": cls._build_task_error_report_url(task) if task.error_line_ids else False,
            "failure_reason": cls._build_task_failure_reason(task),
        }
