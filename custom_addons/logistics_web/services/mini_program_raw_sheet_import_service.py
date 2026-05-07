import csv
import io
import json
import re

from openpyxl import load_workbook

from odoo import fields
from odoo.exceptions import ValidationError

from .waybill_standard_import_service_v2 import WaybillStandardImportService


class MiniProgramRawSheetImportService(WaybillStandardImportService):
    TEMPLATE_CODE = "MINI-PROGRAM-RAW-SHEET-V1"
    TEMPLATE_VERSION = "v1"
    LEGACY_TEMPLATE_CODES = set()
    LEGACY_TEMPLATE_VERSIONS = set()
    OBJECT_TYPE = "mini_program_raw_sheet"
    OBJECT_TYPE_LABEL = "Mini Program Raw Sheet"
    TEMPLATE_FILE_NAME = "mini_program_raw_sheet.xlsx"
    DEFAULT_TEMPLATE_LOCALE = "zh_CN"
    INPUT_MODE = "mini_program_raw_sheet"
    PRIMARY_HEADERS = {
        "store_name": ("收货单位", "收货门店", "客户名称", "门店名称"),
        "address": ("收货地址", "收获地址", "门店地址", "客户地址"),
        "batch_no": ("发车单号", "发车单", "批次号", "batch_no"),
    }
    FIELD_LABELS = {
        "template_file": "template_file",
        "batch_no": "batch_no",
        "store_name": "store_name",
        "address": "address",
    }
    ERROR_REPORT_HEADERS = (
        ("task_no", "task_no"),
        ("source_row_no", "source_row_no"),
        ("field_name", "field_name"),
        ("raw_value", "raw_value"),
        ("mapped_value", "mapped_value"),
        ("error_code", "error_code"),
        ("error_message", "error_message"),
    )

    @classmethod
    def build_template_payload(cls):
        return {
            "import_type": cls.OBJECT_TYPE,
            "import_type_label": "小程序原始单表导入",
            "template_code": cls.TEMPLATE_CODE,
            "template_version": cls.TEMPLATE_VERSION,
            "allowed_extensions": [".xlsx"],
            "required_headers": ["收货单位", "收货地址"],
            "required_meta_fields": ["发车单号"],
            "match_rule": "name_single_hit",
            "dedupe_rule": "batch_no + normalized_store_name + normalized_address",
            "confirm_rule": "ambiguous_or_unmatched_blocked",
            "template_file_name": "小程序导入模板V1.xlsx",
            "template_download_url": False,
        }

    @classmethod
    def precheck(cls, env, raw_bytes, *, filename="", template_code="", template_version=""):
        errors = []
        filename = filename or cls.TEMPLATE_FILE_NAME
        file_ext = cls._guess_file_ext(filename)
        if file_ext != "xlsx":
            errors.append(
                cls._make_error(
                    sheet_name="template_file",
                    row_no=0,
                    field_code="template_file",
                    error_code="FILE_EXTENSION_INVALID",
                    error_message="Only .xlsx files are supported.",
                )
            )
        if not raw_bytes:
            errors.append(
                cls._make_error(
                    sheet_name="template_file",
                    row_no=0,
                    field_code="template_file",
                    error_code="FILE_EMPTY",
                    error_message="Uploaded file is empty.",
                )
            )

        source_data = cls._empty_source_data()
        if not errors:
            source_data, parse_errors = cls._parse_raw_workbook(raw_bytes)
            errors.extend(parse_errors)

        candidates = []
        if not errors:
            candidates, candidate_errors = cls._build_candidates(env, source_data)
            errors.extend(candidate_errors)
        source_data["candidates"] = candidates

        summary = cls._build_summary(source_data, candidates, errors)
        source_file = cls._create_source_file(env, raw_bytes=raw_bytes, filename=filename)
        task = cls._create_mini_precheck_task(env, source_file=source_file, summary=summary)
        cls._write_mini_task_lines(task, candidates)
        cls._write_mini_error_lines(task, errors)

        return {
            "task_no": task.task_no,
            "task_status": cls._public_task_status(task),
            "import_type": cls.OBJECT_TYPE,
            "batch_no": summary.get("batch_no") or False,
            "source_file_name": source_file.file_name,
            "summary": summary,
            "result_flags": {
                "has_ambiguous": summary["ambiguous_row_count"] > 0,
                "has_unmatched": summary["unmatched_row_count"] > 0,
                "has_fatal_error": any((error.get("row_no") or 0) == 0 for error in errors),
            },
            "preview_lines": [cls._public_candidate_payload(item) for item in candidates[:20]],
            "error_count": len(errors),
            "summary_message": task.summary_message,
            "source_file": cls._build_source_file_payload(source_file),
            "error_report_url": cls._build_task_error_report_url(task) if errors else False,
        }

    @classmethod
    def confirm_import(cls, env, *, precheck_token="", import_batch_no="", task_no=""):
        task = cls._get_task_by_task_no(env, task_no or import_batch_no)
        if not task:
            raise ValidationError("Import task not found.")
        return cls._confirm_task_import(env, task)

    @classmethod
    def get_import_task_result(cls, env, *, task_no="", import_batch_no=""):
        task = cls._get_task_by_task_no(env, task_no or import_batch_no)
        if not task:
            raise ValidationError("Import task not found.")
        return cls._build_mini_task_result_payload(task)

    @classmethod
    def get_import_task_lines(cls, env, *, task_no, page=1, page_size=20, status="", match_status="", keyword=""):
        task = cls._get_task_by_task_no(env, task_no)
        if not task:
            raise ValidationError("Import task not found.")
        page = max(int(page or 1), 1)
        page_size = min(max(int(page_size or 20), 1), 200)
        normalized_status = (match_status or status or "").strip()
        normalized_keyword = (keyword or "").strip()
        items = []
        for task_line in task.task_line_ids.sorted(key=lambda record: (record.line_no, record.id)):
            payload = cls._deserialize_task_line_payload(task_line)
            public_item = cls._public_candidate_payload(payload)
            public_item["task_line_status"] = task_line.status
            public_item["task_line_status_label"] = cls._get_selection_label(
                task_line._fields["status"].selection, task_line.status
            )
            public_item["line_no"] = task_line.line_no
            if normalized_status and public_item.get("match_status") != normalized_status:
                continue
            if (
                normalized_keyword
                and normalized_keyword not in (public_item.get("raw_store_name") or "")
                and normalized_keyword not in (public_item.get("partner_name") or "")
            ):
                continue
            items.append(public_item)
        return cls._paginate_task_payload(task, items, page=page, page_size=page_size)

    @classmethod
    def get_import_task_errors(cls, env, *, task_no, page=1, page_size=50, error_code=""):
        task = cls._get_task_by_task_no(env, task_no)
        if not task:
            raise ValidationError("Import task not found.")
        page = max(int(page or 1), 1)
        page_size = min(max(int(page_size or 50), 1), 200)
        code_filter = (error_code or "").strip()
        items = []
        for error in task.error_line_ids.sorted(key=lambda record: (record.source_row_no, record.id)):
            if code_filter and error.error_code != code_filter:
                continue
            items.append(
                {
                    "error_no": len(items) + 1,
                    "line_no": error.task_line_id.line_no if error.task_line_id else False,
                    "source_row_nos": [error.source_row_no] if error.source_row_no else [],
                    "source_row_no": error.source_row_no,
                    "error_level": cls._error_level(error.error_code),
                    "error_code": error.error_code,
                    "error_message": error.error_message,
                    "field_name": error.field_name,
                    "raw_value": error.raw_value or False,
                    "mapped_value": error.mapped_value or False,
                    "suggestion": cls._error_suggestion(error.error_code),
                }
            )
        return cls._paginate_task_payload(task, items, page=page, page_size=page_size)

    @classmethod
    def build_import_task_error_report(cls, env, *, task_no="", import_batch_no="", precheck_token=""):
        task = cls._get_task_by_task_no(env, task_no or import_batch_no)
        if not task:
            raise ValidationError("Import task not found.")
        errors = task.error_line_ids.sorted(key=lambda record: (record.source_row_no, record.id))
        if not errors:
            raise ValidationError("No error report available for the current task.")
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([label for _field_code, label in cls.ERROR_REPORT_HEADERS])
        for error in errors:
            row = {
                "task_no": task.task_no,
                "source_row_no": error.source_row_no,
                "field_name": error.field_name,
                "raw_value": error.raw_value or "",
                "mapped_value": error.mapped_value or "",
                "error_code": error.error_code,
                "error_message": error.error_message,
            }
            writer.writerow([row[field_code] for field_code, _label in cls.ERROR_REPORT_HEADERS])
        return {
            "file_name": f"{task.task_no}-error-report.csv",
            "file_bytes": buffer.getvalue().encode("utf-8-sig"),
            "content_type": "text/csv; charset=utf-8",
        }

    @classmethod
    def _empty_source_data(cls):
        return {
            "input_mode": cls.INPUT_MODE,
            "batch_no": False,
            "raw_rows": [],
            "candidates": [],
        }

    @classmethod
    def _parse_raw_workbook(cls, raw_bytes):
        try:
            workbook = load_workbook(io.BytesIO(raw_bytes), data_only=True)
        except Exception:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name="template_file",
                    row_no=0,
                    field_code="template_file",
                    error_code="FILE_PARSE_ERROR",
                    error_message="Failed to parse the uploaded workbook.",
                )
            ]
        worksheet = workbook[workbook.sheetnames[0]] if workbook.sheetnames else False
        if not worksheet:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name="template_file",
                    row_no=0,
                    field_code="template_file",
                    error_code="WORKSHEET_NOT_FOUND",
                    error_message="No worksheet was found in the workbook.",
                )
            ]

        merged_value_map = cls._build_merged_value_map(worksheet)
        batch_no = cls._extract_batch_no(worksheet, merged_value_map)
        if not batch_no:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name=worksheet.title,
                    row_no=0,
                    field_code="batch_no",
                    error_code="META_BATCH_NO_MISSING",
                    error_message="Failed to locate batch number in the workbook.",
                )
            ]

        header_row_no, header_map = cls._resolve_header_row(worksheet, merged_value_map)
        if not header_row_no:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name=worksheet.title,
                    row_no=0,
                    field_code="template_file",
                    error_code="HEADER_ROW_NOT_FOUND",
                    error_message="Failed to locate the header row.",
                )
            ]
        if "store_name" not in header_map:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name=worksheet.title,
                    row_no=header_row_no,
                    field_code="store_name",
                    error_code="COLUMN_STORE_NAME_MISSING",
                    error_message="Store name column is missing.",
                )
            ]
        if "address" not in header_map:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name=worksheet.title,
                    row_no=header_row_no,
                    field_code="address",
                    error_code="COLUMN_ADDRESS_MISSING",
                    error_message="Address column is missing.",
                )
            ]

        raw_rows = []
        for row_no in range(header_row_no + 1, worksheet.max_row + 1):
            raw_store_name = cls._cell_text(cls._get_cell_value(worksheet, merged_value_map, row_no, header_map["store_name"]))
            raw_address = cls._cell_text(cls._get_cell_value(worksheet, merged_value_map, row_no, header_map["address"]))
            if not raw_store_name and not raw_address:
                continue
            if cls._is_footer_noise_row(raw_store_name, raw_address):
                continue
            raw_rows.append(
                {
                    "_sheet_name": worksheet.title,
                    "_source_row_no": row_no,
                    "batch_no": batch_no,
                    "raw_store_name": raw_store_name,
                    "raw_address": raw_address,
                }
            )
        if not raw_rows:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name=worksheet.title,
                    row_no=header_row_no,
                    field_code="template_file",
                    error_code="DATA_REGION_EMPTY",
                    error_message="No valid data rows were found below the header.",
                )
            ]
        return {
            "input_mode": cls.INPUT_MODE,
            "batch_no": batch_no,
            "raw_rows": raw_rows,
            "candidates": [],
        }, []

    @classmethod
    def _build_candidates(cls, env, source_data):
        grouped = {}
        for row in source_data.get("raw_rows", []):
            normalized_store_name = cls._normalize_store_name(row.get("raw_store_name"))
            normalized_address = cls._normalize_address(row.get("raw_address"))
            business_key = "|".join(
                [
                    cls._normalize_batch_no(row.get("batch_no")),
                    normalized_store_name,
                    normalized_address,
                ]
            )
            bucket = grouped.setdefault(
                business_key,
                {
                    "batch_no": cls._normalize_batch_no(row.get("batch_no")),
                    "raw_store_name": row.get("raw_store_name") or "",
                    "raw_address": row.get("raw_address") or "",
                    "normalized_store_name": normalized_store_name,
                    "normalized_address": normalized_address,
                    "source_row_nos": [],
                    "source_row_count": 0,
                    "business_key": business_key,
                },
            )
            bucket["source_row_nos"].append(row["_source_row_no"])
            bucket["source_row_count"] += 1

        candidates = []
        errors = []
        for business_key, bucket in grouped.items():
            candidate = {
                "business_key": business_key,
                "batch_no": bucket["batch_no"],
                "raw_store_name": bucket["raw_store_name"],
                "raw_address": bucket["raw_address"],
                "normalized_store_name": bucket["normalized_store_name"],
                "normalized_address": bucket["normalized_address"],
                "source_row_nos": bucket["source_row_nos"],
                "source_row_count": bucket["source_row_count"],
                "duplicate_merged": bucket["source_row_count"] > 1,
                "match_status": "row_invalid",
                "partner_id": False,
                "partner_name": False,
                "route_snapshot": {},
                "can_confirm_line": False,
            }
            source_row_no = bucket["source_row_nos"][0] if bucket["source_row_nos"] else 0
            if not bucket["normalized_store_name"]:
                errors.append(
                    cls._make_error(
                        sheet_name="raw_sheet",
                        row_no=source_row_no,
                        field_code="store_name",
                        error_code="ROW_STORE_NAME_EMPTY",
                        error_message="Store name is required.",
                    )
                )
                candidates.append(candidate)
                continue
            partners = cls._search_partner_matches(env, bucket["normalized_store_name"])
            if len(partners) > 1:
                candidate["match_status"] = "matched_ambiguous"
                errors.append(
                    cls._make_error(
                        sheet_name="raw_sheet",
                        row_no=source_row_no,
                        field_code="store_name",
                        error_code="NAME_MATCH_AMBIGUOUS",
                        error_message="Store name matched multiple customer master records.",
                    )
                )
                candidates.append(candidate)
                continue
            if not partners:
                candidate["match_status"] = "unmatched"
                errors.append(
                    cls._make_error(
                        sheet_name="raw_sheet",
                        row_no=source_row_no,
                        field_code="store_name",
                        error_code="NAME_MATCH_NOT_FOUND",
                        error_message="Store name did not match any customer master record.",
                    )
                )
                candidates.append(candidate)
                continue

            partner = partners[0]
            route_snapshot = cls._build_route_snapshot(env, partner, candidate)
            candidate.update(
                {
                    "match_status": "matched_exact",
                    "partner_id": partner.id,
                    "partner_name": partner.name,
                    "route_snapshot": route_snapshot,
                }
            )
            if not route_snapshot.get("longitude") or not route_snapshot.get("latitude"):
                candidate["match_status"] = "row_invalid"
                errors.append(
                    cls._make_error(
                        sheet_name="raw_sheet",
                        row_no=source_row_no,
                        field_code="store_name",
                        error_code="MATCHED_PROFILE_GEO_MISSING",
                        error_message="Matched customer profile is missing longitude or latitude.",
                    )
                )
            else:
                candidate["can_confirm_line"] = True
            candidates.append(candidate)
        candidates.sort(key=lambda item: item["source_row_nos"][0] if item["source_row_nos"] else 0)
        return candidates, errors

    @classmethod
    def _create_mini_precheck_task(cls, env, *, source_file, summary):
        summary_message = cls._build_precheck_summary_message(summary)
        return env["logistics.import.task"].sudo().create(
            {
                "object_type": cls.OBJECT_TYPE,
                "source_file_id": source_file.id,
                "status": "pending",
                "total_count": summary["deduplicated_row_count"],
                "success_count": summary["matched_row_count"],
                "fail_count": summary["ambiguous_row_count"] + summary["unmatched_row_count"] + summary["error_row_count"],
                "operator_id": env.user.id,
                "summary_message": summary_message,
            }
        )

    @classmethod
    def _write_mini_task_lines(cls, task, candidates):
        if not candidates:
            return
        vals_list = []
        for index, candidate in enumerate(candidates, start=1):
            vals_list.append(
                {
                    "task_id": task.id,
                    "line_no": index,
                    "source_row_no": candidate["source_row_nos"][0] if candidate["source_row_nos"] else 0,
                    "object_type": cls.OBJECT_TYPE,
                    "status": "pending" if candidate.get("can_confirm_line") else "failed",
                    "business_key": candidate["business_key"],
                    "message": json.dumps(candidate, ensure_ascii=False),
                }
            )
        task.env["logistics.import.task.line"].sudo().create(vals_list)

    @classmethod
    def _write_mini_error_lines(cls, task, errors):
        if not errors:
            return
        task_lines = task.task_line_ids.sudo()
        task_line_by_row_no = {line.source_row_no: line for line in task_lines}
        vals_list = []
        for error in errors:
            task_line = task_line_by_row_no.get(error.get("row_no") or 0)
            vals_list.append(
                {
                    "task_id": task.id,
                    "task_line_id": task_line.id if task_line else False,
                    "source_row_no": error.get("row_no") or 0,
                    "field_name": error.get("field_code") or "template_file",
                    "raw_value": error.get("raw_value") or False,
                    "mapped_value": error.get("mapped_value") or False,
                    "error_code": error.get("error_code") or "IMPORT_PRECHECK_FAILED",
                    "error_message": error.get("error_message") or "Import precheck failed.",
                }
            )
        task.env["logistics.import.error.line"].sudo().create(vals_list)

    @classmethod
    def _confirm_task_import(cls, env, task):
        task.ensure_one()
        if task.status == "success":
            return cls._build_confirm_result_payload(task)
        if task.status == "running":
            raise ValidationError("Import task is running. Please retry later.")
        if task.status != "pending":
            raise ValidationError("Import task status does not allow confirm.")

        task_lines = task.task_line_ids.sorted(key=lambda record: (record.line_no, record.id))
        candidates = [cls._deserialize_task_line_payload(task_line) for task_line in task_lines]
        if not candidates:
            raise ValidationError("No prechecked rows are available for confirm.")
        if any(not item.get("can_confirm_line") for item in candidates):
            raise ValidationError("Precheck has blocking issues. Confirm is not allowed.")

        batch_no = next((item.get("batch_no") for item in candidates if item.get("batch_no")), "")
        if not batch_no:
            raise ValidationError("Batch number is required.")
        delivery_date = fields.Date.context_today(env.user)
        batch_model = env["logistics.route.planning.batch"].sudo()
        batch = batch_model.search(
            [
                ("batch_no", "=", batch_no),
                ("delivery_date", "=", delivery_date),
            ],
            limit=1,
        )
        if batch and batch.stop_line_ids:
            raise ValidationError(
                "A route planning batch with stop lines already exists for the same batch number and delivery date."
            )
        if not batch:
            batch = batch_model.create(
                {
                    "batch_no": batch_no,
                    "delivery_date": delivery_date,
                    "import_task_id": task.id,
                }
            )
        else:
            batch.write({"import_task_id": task.id})

        created_count = 0
        for index, task_line in enumerate(task_lines, start=1):
            candidate = cls._deserialize_task_line_payload(task_line)
            route_snapshot = candidate.get("route_snapshot") or {}
            stop_line = env["logistics.route.planning.stop.line"].sudo().create(
                {
                    "batch_id": batch.id,
                    "import_task_id": task.id,
                    "import_task_line_id": task_line.id,
                    "waybill_no": f"{batch_no}-RP-{index:03d}",
                    "stop_seq": index,
                    "store_name": candidate.get("partner_name") or candidate.get("raw_store_name") or f"stop-{index}",
                    "longitude": float(route_snapshot.get("longitude") or 0.0),
                    "latitude": float(route_snapshot.get("latitude") or 0.0),
                    "address_detail": route_snapshot.get("address_detail") or candidate.get("raw_address") or "N/A",
                    "contact_phone": route_snapshot.get("contact_phone") or False,
                    "cargo_summary": False,
                    "driver_name": False,
                    "driver_phone": False,
                    "vehicle_no": False,
                    "supplier_name": False,
                    "warehouse_name": False,
                }
            )
            cls._update_task_line(
                task_line,
                status="success",
                message=task_line.message,
                target_model="logistics.route.planning.stop.line",
                target_res_id=stop_line.id,
            )
            created_count += 1

        task.sudo().write(
            {
                "status": "success",
                "total_count": len(task_lines),
                "success_count": created_count,
                "fail_count": max(len(task_lines) - created_count, 0),
                "finished_at": fields.Datetime.now(),
                "summary_message": cls._build_confirm_summary_message(batch_no, created_count),
            }
        )
        return cls._build_confirm_result_payload(task)

    @classmethod
    def _build_mini_task_result_payload(cls, task):
        candidates = [
            cls._deserialize_task_line_payload(task_line)
            for task_line in task.task_line_ids.sorted(key=lambda record: (record.line_no, record.id))
        ]
        errors = list(task.error_line_ids)
        summary = cls._build_summary_from_task(task, candidates, errors)
        source_file = task.source_file_id
        return {
            "task_no": task.task_no,
            "task_status": cls._public_task_status(task),
            "import_type": task.object_type,
            "batch_no": summary.get("batch_no") or False,
            "source_file_name": source_file.file_name if source_file else False,
            "summary": summary,
            "result_flags": {
                "has_ambiguous": summary["ambiguous_row_count"] > 0,
                "has_unmatched": summary["unmatched_row_count"] > 0,
                "has_fatal_error": any((error.source_row_no or 0) == 0 for error in errors),
            },
            "error_count": len(errors),
            "summary_message": task.summary_message,
        }

    @classmethod
    def _build_confirm_result_payload(cls, task):
        task.ensure_one()
        batch = task.env["logistics.route.planning.batch"].sudo().search([("import_task_id", "=", task.id)], limit=1)
        return {
            "task_no": task.task_no,
            "task_status": cls._public_task_status(task),
            "batch_no": batch.batch_no if batch else False,
            "route_planning_batch_id": batch.id if batch else False,
            "route_planning_batch_no": batch.batch_no if batch else False,
            "created_stop_line_count": len(batch.stop_line_ids) if batch else 0,
            "skipped_stop_line_count": 0,
            "can_view_result": bool(batch),
        }

    @classmethod
    def _build_summary(cls, source_data, candidates, errors):
        return {
            "batch_no": source_data.get("batch_no") or False,
            "total_row_count": len(source_data.get("raw_rows", [])),
            "parsed_row_count": len(source_data.get("raw_rows", [])),
            "deduplicated_row_count": len(candidates),
            "duplicate_row_count": sum(max(item.get("source_row_count", 1) - 1, 0) for item in candidates),
            "matched_row_count": sum(
                1
                for item in candidates
                if item.get("match_status") == "matched_exact" and item.get("can_confirm_line")
            ),
            "ambiguous_row_count": sum(1 for item in candidates if item.get("match_status") == "matched_ambiguous"),
            "unmatched_row_count": sum(1 for item in candidates if item.get("match_status") == "unmatched"),
            "error_row_count": len({error.get("row_no") for error in errors if (error.get("row_no") or 0) > 0}),
            "can_confirm_import": bool(candidates)
            and all(item.get("can_confirm_line") for item in candidates)
            and not any((error.get("row_no") or 0) == 0 for error in errors),
        }

    @classmethod
    def _build_summary_from_task(cls, task, candidates, errors):
        batch_no = next((item.get("batch_no") for item in candidates if item.get("batch_no")), False)
        return {
            "batch_no": batch_no,
            "total_row_count": sum(item.get("source_row_count", 1) for item in candidates),
            "parsed_row_count": sum(item.get("source_row_count", 1) for item in candidates),
            "deduplicated_row_count": len(candidates),
            "duplicate_row_count": sum(max(item.get("source_row_count", 1) - 1, 0) for item in candidates),
            "matched_row_count": sum(
                1
                for item in candidates
                if item.get("match_status") == "matched_exact" and item.get("can_confirm_line")
            ),
            "ambiguous_row_count": sum(1 for item in candidates if item.get("match_status") == "matched_ambiguous"),
            "unmatched_row_count": sum(1 for item in candidates if item.get("match_status") == "unmatched"),
            "error_row_count": len({error.source_row_no for error in errors if (error.source_row_no or 0) > 0}),
            "can_confirm_import": task.status == "pending"
            and bool(candidates)
            and all(item.get("can_confirm_line") for item in candidates)
            and not any((error.source_row_no or 0) == 0 for error in errors),
        }

    @classmethod
    def _build_merged_value_map(cls, worksheet):
        merged_value_map = {}
        for merged_range in worksheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merged_range.bounds
            value = worksheet.cell(row=min_row, column=min_col).value
            for row_no in range(min_row, max_row + 1):
                for col_no in range(min_col, max_col + 1):
                    merged_value_map[(row_no, col_no)] = value
        return merged_value_map

    @classmethod
    def _get_cell_value(cls, worksheet, merged_value_map, row_no, col_no):
        if (row_no, col_no) in merged_value_map:
            return merged_value_map[(row_no, col_no)]
        return worksheet.cell(row=row_no, column=col_no).value

    @classmethod
    def _extract_batch_no(cls, worksheet, merged_value_map):
        for row_no in range(1, min(worksheet.max_row, 15) + 1):
            for col_no in range(1, min(worksheet.max_column, 20) + 1):
                cell_text = cls._cell_text(cls._get_cell_value(worksheet, merged_value_map, row_no, col_no))
                normalized = cls._normalize_header(cell_text)
                if not normalized:
                    continue
                if any(cls._normalize_header(alias) in normalized for alias in cls.PRIMARY_HEADERS["batch_no"]):
                    for candidate_col in range(col_no + 1, min(worksheet.max_column, col_no + 4) + 1):
                        candidate_value = cls._cell_text(
                            cls._get_cell_value(worksheet, merged_value_map, row_no, candidate_col)
                        )
                        if candidate_value:
                            return cls._normalize_batch_no(candidate_value)
                    if ":" in cell_text or "：" in cell_text:
                        parts = re.split(r"[:：]", cell_text, maxsplit=1)
                        if len(parts) == 2 and parts[1].strip():
                            return cls._normalize_batch_no(parts[1].strip())
        return False

    @classmethod
    def _resolve_header_row(cls, worksheet, merged_value_map):
        for row_no in range(1, min(worksheet.max_row, 30) + 1):
            header_map = {}
            for col_no in range(1, worksheet.max_column + 1):
                cell_text = cls._cell_text(cls._get_cell_value(worksheet, merged_value_map, row_no, col_no))
                normalized = cls._normalize_header(cell_text)
                if not normalized:
                    continue
                if "store_name" not in header_map and any(
                    cls._normalize_header(alias) == normalized for alias in cls.PRIMARY_HEADERS["store_name"]
                ):
                    header_map["store_name"] = col_no
                if "address" not in header_map and any(
                    cls._normalize_header(alias) == normalized for alias in cls.PRIMARY_HEADERS["address"]
                ):
                    header_map["address"] = col_no
            if "store_name" in header_map and "address" in header_map:
                return row_no, header_map
        return False, {}

    @classmethod
    def _normalize_store_name(cls, value):
        return cls._normalize_text(value)

    @classmethod
    def _normalize_address(cls, value):
        return cls._normalize_text(value)

    @classmethod
    def _normalize_batch_no(cls, value):
        text = cls._normalize_text(value)
        if not text:
            return ""
        text = re.sub(r"^(?:发车单号|发车单|批次号|batch[_ ]?no)\s*[:：]?\s*", "", text, flags=re.IGNORECASE)
        return text.strip()

    @classmethod
    def _normalize_text(cls, value):
        text = (value or "").replace("\r", " ").replace("\n", " ").replace("\u3000", " ")
        text = re.sub(r"\s+", " ", text).strip()
        text = text.replace("（", "(").replace("）", ")")
        text = re.sub(r"\s*\(\s*", "(", text)
        text = re.sub(r"\s*\)\s*", ")", text)
        return text.strip()

    @classmethod
    def _is_footer_noise_row(cls, raw_store_name, raw_address):
        store_name = cls._normalize_text(raw_store_name)
        address = cls._normalize_text(raw_address)
        if address:
            return False
        if not store_name:
            return False
        return "\u6253\u5370\u4eba" in store_name or "\u6253\u5370\u65f6\u95f4" in store_name

    @classmethod
    def _search_partner_matches(cls, env, normalized_store_name):
        if not normalized_store_name:
            return []
        search_term = normalized_store_name
        records = env["res.partner"].sudo().search([("name", "ilike", search_term)])
        if not records:
            search_term = re.split(r"[\s()（）]+", normalized_store_name, maxsplit=1)[0]
            if search_term:
                records = env["res.partner"].sudo().search([("name", "ilike", search_term)])
        matched = [record for record in records if cls._normalize_store_name(record.name) == normalized_store_name]
        return matched

    @classmethod
    def _build_route_snapshot(cls, env, partner, candidate):
        store_profile = env["logistics.store.profile"].sudo().search([("partner_id", "=", partner.id)], limit=1)
        longitude = (store_profile.longitude if store_profile else 0.0) or getattr(partner, "partner_longitude", 0.0) or 0.0
        latitude = (store_profile.latitude if store_profile else 0.0) or getattr(partner, "partner_latitude", 0.0) or 0.0
        address_detail = (
            (store_profile.address_full if store_profile else False)
            or getattr(partner, "address_full", False)
            or candidate.get("raw_address")
            or candidate.get("normalized_address")
            or False
        )
        contact_phone = (
            getattr(partner, "contact_phone", False)
            or getattr(partner, "phone", False)
            or getattr(partner, "mobile", False)
            or False
        )
        return {
            "partner_id": partner.id,
            "partner_name": partner.name,
            "longitude": longitude,
            "latitude": latitude,
            "address_detail": address_detail,
            "contact_phone": contact_phone,
        }

    @classmethod
    def _deserialize_task_line_payload(cls, task_line):
        payload = {}
        if task_line.message:
            try:
                payload = json.loads(task_line.message)
            except Exception:
                payload = {}
        payload.setdefault("business_key", task_line.business_key or False)
        payload.setdefault("source_row_nos", [task_line.source_row_no] if task_line.source_row_no else [])
        payload.setdefault("source_row_count", 1)
        return payload

    @classmethod
    def _public_candidate_payload(cls, candidate):
        source_row_nos = candidate.get("source_row_nos") or []
        return {
            "source_row_nos": source_row_nos,
            "batch_no": candidate.get("batch_no") or False,
            "raw_store_name": candidate.get("raw_store_name") or False,
            "raw_address": candidate.get("raw_address") or False,
            "normalized_store_name": candidate.get("normalized_store_name") or False,
            "normalized_address": candidate.get("normalized_address") or False,
            "match_status": candidate.get("match_status") or False,
            "partner_id": candidate.get("partner_id") or False,
            "partner_name": candidate.get("partner_name") or False,
            "duplicate_merged": bool(candidate.get("duplicate_merged")),
            "duplicate_source_count": max(int(candidate.get("source_row_count") or 1) - 1, 0),
            "can_confirm_line": bool(candidate.get("can_confirm_line")),
        }

    @classmethod
    def _paginate_task_payload(cls, task, items, *, page, page_size):
        total = len(items)
        total_pages = max((total + page_size - 1) // page_size, 1)
        page = min(page, total_pages)
        start = (page - 1) * page_size
        end = start + page_size
        return {
            "task_no": task.task_no,
            "object_type": task.object_type,
            "object_type_label": cls.OBJECT_TYPE_LABEL,
            "task_status": cls._public_task_status(task),
            "page": page,
            "page_size": page_size,
            "total": total,
            "total_pages": total_pages,
            "has_prev": page > 1,
            "has_next": page < total_pages,
            "showing_from": start + 1 if total else 0,
            "showing_to": min(end, total),
            "items": items[start:end],
        }

    @classmethod
    def _public_task_status(cls, task):
        mapping = {
            "pending": "prechecked",
            "running": "confirming",
            "success": "imported",
            "partial_failed": "failed",
            "failed": "failed",
            "cancelled": "failed",
        }
        return mapping.get(task.status, task.status)

    @classmethod
    def _error_level(cls, error_code):
        code = (error_code or "").upper()
        if code in {"NAME_MATCH_AMBIGUOUS", "NAME_MATCH_NOT_FOUND", "ROW_ADDRESS_EMPTY"}:
            return "warning"
        return "error"

    @classmethod
    def _error_suggestion(cls, error_code):
        suggestion_map = {
            "NAME_MATCH_AMBIGUOUS": "Clean duplicate master data or add alias rules.",
            "NAME_MATCH_NOT_FOUND": "Create or sync the matching customer master data first.",
            "MATCHED_PROFILE_GEO_MISSING": "Complete longitude and latitude on the matched store profile.",
            "ROW_STORE_NAME_EMPTY": "Fill in store name in the source workbook and retry.",
        }
        return suggestion_map.get(error_code or "", False)

    @classmethod
    def _build_precheck_summary_message(cls, summary):
        return (
            f"Precheck finished, deduplicated {summary['deduplicated_row_count']} rows, "
            f"matched {summary['matched_row_count']} rows, "
            f"ambiguous {summary['ambiguous_row_count']} rows, "
            f"unmatched {summary['unmatched_row_count']} rows."
        )

    @classmethod
    def _build_confirm_summary_message(cls, batch_no, created_count):
        return f"Confirm finished for {batch_no}, created {created_count} route planning stop lines."
