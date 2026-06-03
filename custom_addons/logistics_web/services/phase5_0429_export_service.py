import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from odoo import fields
from odoo.exceptions import ValidationError


class Phase5ExportError(ValidationError):
    def __init__(self, error_code, message):
        self.error_code = error_code
        super().__init__(message)


class Phase5ExportService:
    DOWNLOAD_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
    HEADER_FONT = Font(bold=True, color="1F2937")
    NOTE_FILL = PatternFill(fill_type="solid", fgColor="FFF3CD")
    NOTE_FONT = Font(bold=True, color="9A6700")

    EXPORT_CONFIG = {
        "order_detail": {
            "file_name": "排线-订单详情_名创+遇见0429.xlsx",
            "sheets": [
                {
                    "title": "Sheet1",
                    "headers": [
                        "批次号",
                        "排线序号",
                        "运单号",
                        "销售订单号",
                        "客户名称",
                        "订单备注",
                        "收货地址",
                        "整件数",
                        "散件数",
                        "总重量_kg",
                        "总体积_m3",
                        "集货位",
                        "司机姓名",
                        "车牌号",
                    ],
                    "widths": {
                        "A": 18,
                        "B": 10,
                        "C": 20,
                        "D": 20,
                        "E": 18,
                        "F": 24,
                        "G": 38,
                        "H": 12,
                        "I": 12,
                        "J": 14,
                        "K": 14,
                        "L": 16,
                        "M": 14,
                        "N": 14,
                    },
                }
            ],
        },
        "store_detail": {
            "file_name": "排线-门店详情_名创+遇见0429.xlsx",
            "sheets": [
                {
                    "title": "Sheet1",
                    "headers": [
                        "批次号",
                        "排线序号",
                        "运单号",
                        "客户名称",
                        "收货地址",
                        "联系人",
                        "联系电话",
                        "集货位",
                        "不收货时间段",
                        "门店备注",
                        "整件数",
                        "散件数",
                        "总重量_kg",
                        "总体积_m3",
                        "经度",
                        "纬度",
                        "司机姓名",
                        "车牌号",
                    ],
                    "widths": {
                        "A": 18,
                        "B": 10,
                        "C": 20,
                        "D": 18,
                        "E": 38,
                        "F": 14,
                        "G": 16,
                        "H": 16,
                        "I": 22,
                        "J": 24,
                        "K": 12,
                        "L": 12,
                        "M": 14,
                        "N": 14,
                        "O": 14,
                        "P": 14,
                        "Q": 14,
                        "R": 14,
                    },
                }
            ],
        },
        "store_goods_triplet": {
            "file_name": "排线-门店货物信息三联单名创+遇见0429.xlsx",
            "sheets": [
                {
                    "title": "主要字段",
                    "headers": [
                        "批次号",
                        "排线序号",
                        "运单号",
                        "销售订单号",
                        "客户名称",
                        "行号",
                        "商品名称",
                        "商品编号",
                        "商品条码",
                        "包装规格",
                        "发货数量",
                        "小单位销售单位",
                        "小单位销售单价",
                        "金额",
                        "收货地址",
                        "联系人",
                        "联系电话",
                        "业务员",
                        "业务员联系方式",
                    ],
                    "widths": {
                        "A": 18,
                        "B": 10,
                        "C": 20,
                        "D": 20,
                        "E": 18,
                        "F": 10,
                        "G": 22,
                        "H": 16,
                        "I": 18,
                        "J": 18,
                        "K": 18,
                        "L": 16,
                        "M": 16,
                        "N": 14,
                        "O": 38,
                        "P": 14,
                        "Q": 16,
                        "R": 14,
                        "S": 18,
                    },
                },
                {
                    "title": "表头字段",
                    "headers": [
                        "销售订单号",
                        "客户名称",
                        "收货地址",
                        "联系人",
                        "联系电话",
                        "业务员",
                        "业务员联系方式",
                    ],
                    "widths": {
                        "A": 20,
                        "B": 18,
                        "C": 38,
                        "D": 14,
                        "E": 16,
                        "F": 14,
                        "G": 18,
                    },
                },
            ],
        },
    }

    @classmethod
    def export_by_delivery_date(cls, env, *, export_key="", delivery_date="", batch_no="", file_locale="zh_CN"):
        del file_locale
        cls._ensure_export_access(env)
        config = cls.EXPORT_CONFIG.get(export_key)
        if not config:
            raise Phase5ExportError("PHASE5_EXPORT_KEY_INVALID", f"Unsupported export key: {export_key}")
        normalized_date = cls._normalize_delivery_date(delivery_date)
        normalized_batch_no = (batch_no or "").strip()

        dataset = cls._collect_dataset(env, delivery_date=normalized_date, batch_no=normalized_batch_no)
        sheet_rows = cls._build_sheet_rows(export_key, dataset)
        workbook_bytes = cls._build_workbook_bytes(
            config=config,
            sheet_rows=sheet_rows,
            delivery_date=normalized_date,
            batch_no=normalized_batch_no,
        )
        return {
            "file_name": config["file_name"],
            "content_type": cls.DOWNLOAD_CONTENT_TYPE,
            "file_bytes": workbook_bytes,
        }

    @classmethod
    def _ensure_export_access(cls, env):
        if env.user.has_group("logistics_dispatch.group_logistics_export_user"):
            return
        raise Phase5ExportError("EXPORT_PERMISSION_DENIED", "You do not have permission to export the 0429 workbook.")

    @classmethod
    def _normalize_delivery_date(cls, delivery_date):
        if not delivery_date:
            raise Phase5ExportError("DELIVERY_DATE_REQUIRED", "delivery_date is required.")
        normalized = fields.Date.to_date(delivery_date)
        if not normalized:
            raise Phase5ExportError("DELIVERY_DATE_INVALID", f"Invalid delivery_date: {delivery_date}")
        return normalized.isoformat()

    @classmethod
    def _collect_dataset(cls, env, *, delivery_date, batch_no):
        waybill_domain = [("delivery_date", "=", delivery_date)]
        if batch_no:
            waybill_domain.append(("batch_id.name", "=", batch_no))
        waybills = env["logistics.dispatch.waybill"].sudo().search(
            waybill_domain,
            order="batch_id asc, route_seq asc, name asc, id asc",
        )
        if not waybills:
            detail = f" and batch_no {batch_no}" if batch_no else ""
            raise Phase5ExportError(
                "PHASE5_EXPORT_EMPTY",
                f"No export data found for delivery_date {delivery_date}{detail}.",
            )

        customer_lines = env["logistics.dispatch.waybill.customer.line"].sudo().search(
            [("waybill_id", "in", waybills.ids)],
            order="waybill_id asc, stop_seq_in_waybill asc, sequence asc, id asc",
        )
        order_lines = env["logistics.dispatch.waybill.order.line"].sudo().search(
            [("waybill_id", "in", waybills.ids)],
            order="waybill_id asc, id asc",
        )
        goods_lines = env["logistics.dispatch.waybill.customer.goods.line"].sudo().search(
            [("waybill_id", "in", waybills.ids)],
            order="waybill_id asc, customer_line_id asc, order_line_id asc, sequence asc, id asc",
        )
        stop_domain = [("delivery_date", "=", delivery_date)]
        if batch_no:
            stop_domain.append(("batch_no", "=", batch_no))
        stop_lines = env["logistics.route.planning.stop.line"].sudo().search(
            stop_domain,
            order="batch_no asc, stop_seq asc, id asc",
        )
        store_profiles = env["logistics.store.profile"].sudo().search(
            [("partner_id", "in", customer_lines.mapped("store_id").ids)],
        )

        stop_by_waybill = {}
        for stop_line in stop_lines:
            stop_by_waybill.setdefault(stop_line.waybill_no, stop_line)

        customer_line_map = {line.id: line for line in customer_lines}
        order_line_map = {line.id: line for line in order_lines}
        order_lines_by_customer_line = {}
        for order_line in order_lines:
            if order_line.customer_line_id:
                order_lines_by_customer_line.setdefault(order_line.customer_line_id.id, []).append(order_line)
        store_profile_by_partner = {profile.partner_id.id: profile for profile in store_profiles}

        return {
            "waybills": waybills,
            "customer_lines": customer_lines,
            "order_lines": order_lines,
            "goods_lines": goods_lines,
            "stop_by_waybill": stop_by_waybill,
            "customer_line_map": customer_line_map,
            "order_line_map": order_line_map,
            "order_lines_by_customer_line": order_lines_by_customer_line,
            "store_profile_by_partner": store_profile_by_partner,
        }

    @classmethod
    def _build_sheet_rows(cls, export_key, dataset):
        if export_key == "order_detail":
            return {"Sheet1": cls._build_order_detail_rows(dataset)}
        if export_key == "store_detail":
            return {"Sheet1": cls._build_store_detail_rows(dataset)}
        if export_key == "store_goods_triplet":
            main_rows, header_rows = cls._build_store_goods_triplet_rows(dataset)
            return {"主要字段": main_rows, "表头字段": header_rows}
        raise Phase5ExportError("PHASE5_EXPORT_KEY_INVALID", f"Unsupported export key: {export_key}")

    @classmethod
    def _build_order_detail_rows(cls, dataset):
        rows = []
        for order_line in dataset["order_lines"]:
            waybill = order_line.waybill_id
            customer_line = order_line.customer_line_id
            stop_line = dataset["stop_by_waybill"].get(waybill.name)
            rows.append(
                {
                    "批次号": cls._batch_no(waybill, stop_line),
                    "排线序号": cls._stop_seq(waybill, customer_line, stop_line),
                    "运单号": waybill.name or "",
                    "销售订单号": order_line.sales_order_no or order_line.order_no or order_line.source_doc_no or "",
                    "客户名称": cls._customer_name(customer_line, waybill),
                    "订单备注": order_line.order_remark or "",
                    "收货地址": cls._address(customer_line, stop_line),
                    "整件数": cls._record_value(order_line, "whole_package_count"),
                    "散件数": cls._record_value(order_line, "loose_package_count"),
                    "总重量_kg": cls._safe_number(order_line.weight_summary),
                    "总体积_m3": cls._safe_number(order_line.volume_summary),
                    "集货位": cls._record_value(order_line, "gathering_location"),
                    "司机姓名": cls._driver_name(waybill, stop_line),
                    "车牌号": cls._vehicle_no(waybill, stop_line),
                }
            )
        return rows

    @classmethod
    def _build_store_detail_rows(cls, dataset):
        rows = []
        for customer_line in dataset["customer_lines"]:
            waybill = customer_line.waybill_id
            stop_line = dataset["stop_by_waybill"].get(waybill.name)
            store_partner = customer_line.store_id or customer_line.partner_id or customer_line.customer_id
            store_profile = dataset["store_profile_by_partner"].get(store_partner.id) if store_partner else False
            related_orders = dataset["order_lines_by_customer_line"].get(customer_line.id, [])
            rows.append(
                {
                    "批次号": cls._batch_no(waybill, stop_line),
                    "排线序号": cls._stop_seq(waybill, customer_line, stop_line),
                    "运单号": waybill.name or "",
                    "客户名称": cls._customer_name(customer_line, waybill),
                    "收货地址": cls._address(customer_line, stop_line),
                    "联系人": customer_line.contact_name_snapshot or "",
                    "联系电话": (customer_line.contact_phone_snapshot or "") or (stop_line.contact_phone if stop_line else ""),
                    "集货位": cls._join_distinct(
                        [cls._record_value(order_line, "gathering_location") for order_line in related_orders]
                    ),
                    "不收货时间段": cls._no_receive_time_slots(store_profile, store_partner),
                    "门店备注": waybill.remark or waybill.delivery_remark_snapshot or "",
                    "整件数": cls._sum_numeric([cls._record_value(order_line, "whole_package_count") for order_line in related_orders]),
                    "散件数": cls._sum_numeric([cls._record_value(order_line, "loose_package_count") for order_line in related_orders]),
                    "总重量_kg": cls._safe_number(customer_line.total_weight),
                    "总体积_m3": cls._safe_number(customer_line.total_volume),
                    "经度": customer_line.longitude_snapshot or (stop_line.longitude if stop_line else ""),
                    "纬度": customer_line.latitude_snapshot or (stop_line.latitude if stop_line else ""),
                    "司机姓名": cls._driver_name(waybill, stop_line),
                    "车牌号": cls._vehicle_no(waybill, stop_line),
                }
            )
        return rows

    @classmethod
    def _build_store_goods_triplet_rows(cls, dataset):
        main_rows = []
        header_rows = []
        seen_header_keys = set()
        for goods_line in dataset["goods_lines"]:
            waybill = goods_line.waybill_id
            customer_line = goods_line.customer_line_id
            order_line = goods_line.order_line_id
            stop_line = dataset["stop_by_waybill"].get(waybill.name)
            sales_order_no = (
                order_line.sales_order_no or order_line.order_no or order_line.source_doc_no or ""
                if order_line
                else ""
            )
            customer_name = cls._customer_name(customer_line, waybill)
            address = cls._address(customer_line, stop_line)
            contact_name = customer_line.contact_name_snapshot or ""
            contact_phone = (customer_line.contact_phone_snapshot or "") or (stop_line.contact_phone if stop_line else "")
            salesperson_name = order_line.salesperson_name_snapshot if order_line else ""
            salesperson_phone = cls._record_value(order_line, "salesperson_phone")
            main_rows.append(
                {
                    "批次号": cls._batch_no(waybill, stop_line),
                    "排线序号": cls._stop_seq(waybill, customer_line, stop_line),
                    "运单号": waybill.name or "",
                    "销售订单号": sales_order_no,
                    "客户名称": customer_name,
                    "行号": goods_line.sequence or "",
                    "商品名称": goods_line.goods_name or goods_line.product_name_snapshot or "",
                    "商品编号": goods_line.goods_code or goods_line.external_product_code_snapshot or "",
                    "商品条码": goods_line.barcode_snapshot or "",
                    "包装规格": goods_line.specification or goods_line.spec_snapshot or "",
                    "发货数量": cls._record_value(goods_line, "delivery_qty_text"),
                    "小单位销售单位": goods_line.small_unit_name or "",
                    "小单位销售单价": cls._safe_number(goods_line.small_unit_price),
                    "金额": cls._safe_number(goods_line.amount),
                    "收货地址": address,
                    "联系人": contact_name,
                    "联系电话": contact_phone,
                    "业务员": salesperson_name or "",
                    "业务员联系方式": salesperson_phone,
                }
            )
            header_key = (
                sales_order_no,
                customer_name,
                address,
                contact_name,
                contact_phone,
                salesperson_name or "",
                salesperson_phone,
            )
            if header_key not in seen_header_keys:
                seen_header_keys.add(header_key)
                header_rows.append(
                    {
                        "销售订单号": sales_order_no,
                        "客户名称": customer_name,
                        "收货地址": address,
                        "联系人": contact_name,
                        "联系电话": contact_phone,
                        "业务员": salesperson_name or "",
                        "业务员联系方式": salesperson_phone,
                    }
                )
        return main_rows, header_rows

    @classmethod
    def _build_workbook_bytes(cls, *, config, sheet_rows, delivery_date, batch_no):
        workbook = Workbook()
        active_sheet = workbook.active
        note_text = f"导出日期：{delivery_date}"
        if batch_no:
            note_text += f"；批次号：{batch_no}"
        note_text += "；0429 模板导出。"
        for index, sheet_config in enumerate(config["sheets"]):
            sheet = active_sheet if index == 0 else workbook.create_sheet(title=sheet_config["title"])
            sheet.title = sheet_config["title"]
            headers = sheet_config["headers"]
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
            note_cell = sheet.cell(row=1, column=1, value=note_text)
            note_cell.fill = cls.NOTE_FILL
            note_cell.font = cls.NOTE_FONT
            note_cell.alignment = Alignment(horizontal="left", vertical="center")
            for col_index, header in enumerate(headers, start=1):
                cell = sheet.cell(row=2, column=col_index, value=header)
                cell.fill = cls.HEADER_FILL
                cell.font = cls.HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row_index, row in enumerate(sheet_rows.get(sheet_config["title"], []), start=3):
                for col_index, header in enumerate(headers, start=1):
                    value = row.get(header, "")
                    cell = sheet.cell(row=row_index, column=col_index, value=value)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.freeze_panes = "A3"
            last_row = max(sheet.max_row, 2)
            sheet.auto_filter.ref = f"A2:{cls._column_letter(len(headers))}{last_row}"
            for column_letter, width in sheet_config["widths"].items():
                sheet.column_dimensions[column_letter].width = width
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    @classmethod
    def _batch_no(cls, waybill, stop_line):
        batch = getattr(waybill, "batch_id", False)
        return (
            (stop_line.batch_no if stop_line else "")
            or (batch.batch_no if batch and hasattr(batch, "batch_no") else "")
            or (batch.name if batch else "")
            or ""
        )

    @classmethod
    def _stop_seq(cls, waybill, customer_line, stop_line):
        if stop_line and stop_line.stop_seq:
            return stop_line.stop_seq
        if customer_line and customer_line.stop_seq_in_waybill:
            return customer_line.stop_seq_in_waybill
        return waybill.route_seq or ""

    @classmethod
    def _customer_name(cls, customer_line, waybill):
        return (
            (customer_line.customer_name_snapshot or "").strip()
            or (customer_line.customer_id.name or "").strip()
            or (customer_line.partner_id.name or "").strip()
            or (waybill.partner_name or "").strip()
            or (waybill.customer_name or "").strip()
            or ""
        )

    @classmethod
    def _address(cls, customer_line, stop_line):
        return (
            (customer_line.address_full_snapshot or "").strip()
            or ((stop_line.address_detail or "").strip() if stop_line else "")
            or ""
        )

    @classmethod
    def _driver_name(cls, waybill, stop_line):
        batch = getattr(waybill, "batch_id", False)
        return (
            (stop_line.driver_name if stop_line else "")
            or (batch.driver_name_snapshot if batch else "")
            or (waybill.driver_employee_id.name if waybill.driver_employee_id else "")
            or ""
        )

    @classmethod
    def _vehicle_no(cls, waybill, stop_line):
        batch = getattr(waybill, "batch_id", False)
        vehicle = getattr(waybill, "vehicle_id", False)
        return (
            (stop_line.vehicle_no if stop_line else "")
            or (vehicle.license_plate if vehicle else "")
            or (batch.vehicle_id.license_plate if batch and batch.vehicle_id else "")
            or ""
        )

    @classmethod
    def _no_receive_time_slots(cls, store_profile, store_partner):
        return (
            cls._record_value(store_profile, "no_receive_time_slots_text")
            or cls._record_value(store_partner, "no_receive_time_slots_text")
            or ""
        )

    @classmethod
    def _record_value(cls, record, field_name):
        if not record:
            return ""
        if field_name not in record._fields:
            return ""
        value = record[field_name]
        return value if value not in (False, None) else ""

    @classmethod
    def _safe_number(cls, value):
        return "" if value in (False, None) else value

    @classmethod
    def _sum_numeric(cls, values):
        total = 0
        has_value = False
        for value in values:
            if value in (False, None, ""):
                continue
            has_value = True
            total += value
        return total if has_value else ""

    @classmethod
    def _join_distinct(cls, values):
        seen = []
        for value in values:
            text = str(value or "").strip()
            if text and text not in seen:
                seen.append(text)
        return " / ".join(seen)

    @classmethod
    def _column_letter(cls, index):
        result = []
        while index:
            index, remainder = divmod(index - 1, 26)
            result.append(chr(65 + remainder))
        return "".join(reversed(result))
