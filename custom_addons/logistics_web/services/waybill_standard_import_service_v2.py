import csv
import hashlib
import io
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from odoo import fields
from odoo.exceptions import ValidationError


class WaybillStandardImportService:
    TEMPLATE_CODE = "TSL-IMPORT-WAYBILL-V3"
    TEMPLATE_VERSION = "v3"
    LEGACY_TEMPLATE_CODES = {"TSL-IMPORT-WAYBILL-V2"}
    LEGACY_TEMPLATE_VERSIONS = {"v2"}
    OBJECT_TYPE = "dispatch_main"
    OBJECT_TYPE_LABEL = "主数据导入"
    TEMPLATE_FILE_NAME = "TSL-IMPORT-WAYBILL-V3.xlsx"
    TEMPLATE_VARIANTS = {
        "en_US": {
            "label": "标准模板（英文列头）",
            "file_name": "TSL-IMPORT-WAYBILL-V3.xlsx",
            "field_locale": "en_US",
        },
        "zh_CN": {
            "label": "标准模板（中文列头）",
            "file_name": "TSL-IMPORT-WAYBILL-V3.zh_CN.xlsx",
            "field_locale": "zh_CN",
        },
    }
    DEFAULT_TEMPLATE_LOCALE = "zh_CN"
    FORMAL_INPUT_MODE = "formal_four_sheet"
    FIELD_LABELS = {
        "warehouse_code": "仓库",
        "delivery_date": "预出库日期",
        "wave_no": "波次号",
        "batch_no": "批次号",
        "waybill_no": "运单号",
        "organization_name": "组织",
        "route_name": "线路",
        "driver_name": "司机",
        "driver_phone": "司机电话",
        "delivery_remark": "主表备注",
        "customer_line_no": "门店节点编号",
        "external_customer_code": "客户编号",
        "customer_name": "客户名称",
        "contact_name": "联系人",
        "contact_phone": "联系电话",
        "department_name": "部门名称",
        "salesperson_name": "业务员",
        "channel_name": "渠道",
        "customer_level": "客户等级",
        "allow_cash_on_delivery": "是否允许货到付款",
        "customer_status": "客户状态",
        "internal_counterparty_flag": "内部往来单位",
        "invoice_type": "发票类型",
        "registered_phone": "注册电话",
        "registered_address": "注册地址",
        "customer_seq_no": "客户序号",
        "address_full": "客户地址",
        "province_name": "省",
        "city_name": "市",
        "district_name": "区",
        "longitude": "经度",
        "latitude": "纬度",
        "route_preference": "线路偏好",
        "warehouse_preference": "仓库偏好",
        "stop_seq_in_waybill": "停靠点顺序",
        "receive_start_time": "开始收货时间",
        "receive_end_time": "截止收货时间",
        "receive_time_slots_text": "收货时间段",
        "no_receive_time_slots_text": "不收货时间段",
        "illegal_parking_flag": "违停",
        "free_parking_minutes": "免费停车时长_分钟",
        "parking_fee_per_hour": "停车费_小时",
        "parking_location_text": "停车位置",
        "parking_mode_text": "停车方式",
        "unload_entrance_text": "卸货入口",
        "unload_location_text": "卸货位置",
        "access_alley": "小巷子",
        "access_handcart": "手推车",
        "access_pallet_exchange": "托盘置换",
        "access_cooler_box_exchange": "保温箱置换",
        "upstairs_floor_count": "上楼层数",
        "basement_height_limit_text": "地库限高",
        "delivery_week_mon": "周一配送",
        "delivery_week_tue": "周二配送",
        "delivery_week_wed": "周三配送",
        "delivery_week_thu": "周四配送",
        "delivery_week_fri": "周五配送",
        "delivery_week_sat": "周六配送",
        "delivery_week_sun": "周日配送",
        "order_line_no": "订单行编号",
        "source_doc_no": "单据号",
        "sales_order_no": "销售订单号",
        "source_ref_no": "来源参考号",
        "third_party_doc_no": "第三方单据号",
        "doc_type": "单据类型",
        "business_type": "业务类型",
        "doc_source": "单据来源",
        "doc_date": "单据日期",
        "audited_at": "审核日期",
        "payment_status": "支付状态",
        "settlement_status": "结算状态",
        "doc_status": "单据状态",
        "logistics_status": "物流状态",
        "maker_name": "制单人",
        "auditor_name": "审核人",
        "made_at": "创建日期",
        "order_remark": "订单备注",
        "custom_field_1": "自定义字段1",
        "external_product_code": "货号",
        "product_name": "商品名称",
        "spec": "规格",
        "barcode": "条码",
        "brand_name": "品牌",
        "category_name": "类别",
        "base_unit_name": "基本单位",
        "doc_unit_name": "单据单位",
        "small_unit_name": "小单位",
        "base_qty": "基本数量",
        "doc_qty": "单据数量",
        "small_qty": "小单位数量",
        "box_qty": "箱数",
        "gift_qty": "赠品数量",
        "exchange_qty": "换货数量",
        "unit_price": "单据单价",
        "small_unit_price": "小单位单价",
        "amount": "金额",
        "settled_amount": "已结算金额",
        "unsettled_amount": "未结算金额",
        "tax_amount": "税额",
        "amount_ex_tax": "不含税金额",
        "cost_amount": "成本金额",
        "gross_profit": "毛利",
        "gross_profit_rate": "毛利率",
        "above_standard_price_flag": "高于标准价",
        "below_standard_price_flag": "低于标准价",
        "unit_weight": "商品单件毛重",
        "unit_volume": "商品单件体积",
        "total_weight": "总毛重",
        "total_volume": "总体积",
        "line_remark": "明细备注",
        "remark": "备注",
        "customer_no": "客户编号",
        "store_no": "门店编号",
        "store_name": "门店名称",
        "customer_ref": "客户参考号",
        "goods_code": "货物编码",
        "goods_name": "货物名称",
        "qty": "数量",
        "package_count": "件数",
        "uom_name": "单位",
        "weight": "重量",
        "volume": "体积",
        "temperature_zone": "温层",
        "package_type": "包装类型",
        "template_file": "模板文件",
    }
    WAYBILL_SHEET = {
        "key": "waybill_rows",
        "sheet_name": "Waybill",
        "fields": [
            "warehouse_code",
            "delivery_date",
            "wave_no",
            "batch_no",
            "waybill_no",
            "organization_name",
            "route_name",
            "driver_name",
            "driver_phone",
            "delivery_remark",
        ],
        "required_fields": {"warehouse_code", "delivery_date", "wave_no", "batch_no", "waybill_no"},
    }
    CUSTOMER_LINE_SHEET = {
        "key": "customer_line_rows",
        "sheet_name": "CustomerLine",
        "fields": [
            "waybill_no",
            "customer_line_no",
            "external_customer_code",
            "customer_name",
            "contact_name",
            "contact_phone",
            "organization_name",
            "department_name",
            "salesperson_name",
            "channel_name",
            "customer_level",
            "allow_cash_on_delivery",
            "customer_status",
            "internal_counterparty_flag",
            "invoice_type",
            "registered_phone",
            "registered_address",
            "customer_seq_no",
            "address_full",
            "province_name",
            "city_name",
            "district_name",
            "longitude",
            "latitude",
            "route_preference",
            "warehouse_preference",
            "stop_seq_in_waybill",
            "receive_start_time",
            "receive_end_time",
            "receive_time_slots_text",
            "no_receive_time_slots_text",
            "illegal_parking_flag",
            "free_parking_minutes",
            "parking_fee_per_hour",
            "parking_location_text",
            "parking_mode_text",
            "unload_entrance_text",
            "unload_location_text",
            "access_alley",
            "access_handcart",
            "access_pallet_exchange",
            "access_cooler_box_exchange",
            "upstairs_floor_count",
            "basement_height_limit_text",
            "delivery_week_mon",
            "delivery_week_tue",
            "delivery_week_wed",
            "delivery_week_thu",
            "delivery_week_fri",
            "delivery_week_sat",
            "delivery_week_sun",
        ],
        "required_fields": {"waybill_no", "customer_line_no", "customer_name", "address_full"},
    }
    ORDER_LINE_SHEET = {
        "key": "order_line_rows",
        "sheet_name": "OrderLine",
        "fields": [
            "order_line_no",
            "waybill_no",
            "customer_line_no",
            "source_doc_no",
            "sales_order_no",
            "source_ref_no",
            "third_party_doc_no",
            "doc_type",
            "business_type",
            "doc_source",
            "doc_date",
            "audited_at",
            "department_name",
            "channel_name",
            "salesperson_name",
            "payment_status",
            "settlement_status",
            "doc_status",
            "logistics_status",
            "maker_name",
            "auditor_name",
            "made_at",
            "order_remark",
            "custom_field_1",
        ],
        "required_fields": {"order_line_no", "waybill_no", "customer_line_no"},
    }
    GOODS_LINE_SHEET = {
        "key": "goods_line_rows",
        "sheet_name": "GoodsLine",
        "fields": [
            "order_line_no",
            "external_product_code",
            "product_name",
            "spec",
            "barcode",
            "brand_name",
            "category_name",
            "base_unit_name",
            "doc_unit_name",
            "small_unit_name",
            "base_qty",
            "doc_qty",
            "small_qty",
            "box_qty",
            "gift_qty",
            "exchange_qty",
            "unit_price",
            "small_unit_price",
            "amount",
            "settled_amount",
            "unsettled_amount",
            "tax_amount",
            "amount_ex_tax",
            "cost_amount",
            "gross_profit",
            "gross_profit_rate",
            "above_standard_price_flag",
            "below_standard_price_flag",
            "unit_weight",
            "unit_volume",
            "total_weight",
            "total_volume",
            "line_remark",
        ],
        "required_fields": {"order_line_no", "product_name", "doc_qty"},
    }
    PRIMARY_SHEET = {
        "key": "rows",
        "sheet_name": "????",
        "fields": [
            "warehouse_code",
            "delivery_date",
            "wave_no",
            "batch_no",
            "waybill_no",
            "driver_name",
            "driver_phone",
            "remark",
            "customer_line_no",
            "customer_no",
            "customer_name",
            "store_no",
            "store_name",
            "delivery_remark",
            "signoff_requirement",
            "customer_ref",
            "goods_code",
            "goods_name",
            "spec",
            "qty",
            "package_count",
            "uom_name",
            "weight",
            "volume",
            "temperature_zone",
            "package_type",
        ],
        "required_fields": {"warehouse_code", "delivery_date", "batch_no", "waybill_no", "goods_name", "qty"},
    }
    SHEETS = [WAYBILL_SHEET, CUSTOMER_LINE_SHEET, ORDER_LINE_SHEET, GOODS_LINE_SHEET]
    LEGACY_SHEETS = [
        {
            "key": "waybill_rows",
            "sheet_name": "??",
            "fields": [
                "waybill_no",
                "batch_no",
                "wave_no",
                "delivery_date",
                "warehouse_code",
                "driver_name",
                "driver_phone",
                "remark",
            ],
            "required_fields": {"waybill_no"},
        },
        {
            "key": "customer_rows",
            "sheet_name": "????",
            "fields": [
                "waybill_no",
                "customer_no",
                "customer_name",
                "store_no",
                "store_name",
                "delivery_remark",
                "signoff_requirement",
                "customer_ref",
            ],
            "required_fields": {"waybill_no", "customer_no"},
        },
        {
            "key": "goods_rows",
            "sheet_name": "????",
            "fields": [
                "waybill_no",
                "customer_no",
                "store_no",
                "goods_code",
                "goods_name",
                "spec",
                "qty",
                "package_count",
                "uom_name",
                "weight",
                "volume",
                "temperature_zone",
                "package_type",
                "remark",
            ],
            "required_fields": {"waybill_no", "customer_no", "goods_name", "qty"},
        },
    ]
    STATUS_LABELS = {
        "prechecked": "????",
        "importing": "???",
        "finished": "???",
        "failed": "???",
        "expired": "???",
    }
    ALLOWED_TEMPERATURE_ZONES = {"ambient", "chilled", "frozen", "other"}
    BOOLEAN_TRUE_VALUES = {"?", "Y", "YES", "TRUE", "1"}
    BOOLEAN_FALSE_VALUES = {"?", "N", "NO", "FALSE", "0"}
    ERROR_REPORT_HEADERS = [
        ("sheet_name", "Sheet"),
        ("row_no", "??"),
        ("field_code", "????"),
        ("field_label", "????"),
        ("error_code", "????"),
        ("error_message", "????"),
    ]
    FIELD_ALIASES = {
        "warehouse_code": ["warehouse_code", "warehouse code", "??"],
        "delivery_date": ["delivery_date", "delivery date", "?????"],
        "wave_no": ["wave_no", "wave no", "???"],
        "batch_no": ["batch_no", "batch no", "???"],
        "waybill_no": ["waybill_no", "waybill no", "???"],
        "organization_name": ["organization_name", "organization name", "??"],
        "route_name": ["route_name", "route name", "??"],
        "driver_name": ["driver_name", "driver name", "??"],
        "driver_phone": ["driver_phone", "driver phone", "????"],
        "delivery_remark": ["delivery_remark", "delivery remark", "????", "????"],
        "customer_line_no": ["customer_line_no", "customer line no", "??????"],
        "external_customer_code": ["external_customer_code", "external customer code", "????"],
        "customer_name": ["customer_name", "customer name", "????"],
        "contact_name": ["contact_name", "contact name", "???"],
        "contact_phone": ["contact_phone", "contact phone", "????"],
        "department_name": ["department_name", "department name", "????"],
        "salesperson_name": ["salesperson_name", "salesperson name", "???"],
        "channel_name": ["channel_name", "channel name", "??"],
        "customer_level": ["customer_level", "customer level", "????"],
        "allow_cash_on_delivery": ["allow_cash_on_delivery", "allow cash on delivery", "????????"],
        "customer_status": ["customer_status", "customer status", "????"],
        "internal_counterparty_flag": ["internal_counterparty_flag", "internal counterparty flag", "??????"],
        "invoice_type": ["invoice_type", "invoice type", "????"],
        "registered_phone": ["registered_phone", "registered phone", "????"],
        "registered_address": ["registered_address", "registered address", "????"],
        "customer_seq_no": ["customer_seq_no", "customer seq no", "????"],
        "address_full": ["address_full", "address full", "????"],
        "province_name": ["province_name", "province name", "?"],
        "city_name": ["city_name", "city name", "?"],
        "district_name": ["district_name", "district name", "?"],
        "longitude": ["longitude", "??"],
        "latitude": ["latitude", "??"],
        "route_preference": ["route_preference", "route preference", "????"],
        "warehouse_preference": ["warehouse_preference", "warehouse preference", "????"],
        "stop_seq_in_waybill": ["stop_seq_in_waybill", "stop seq in waybill", "?????"],
        "receive_start_time": ["receive_start_time", "receive start time", "??????"],
        "receive_end_time": ["receive_end_time", "receive end time", "??????"],
        "receive_time_slots_text": ["receive_time_slots_text", "receive time slots text", "?????"],
        "no_receive_time_slots_text": ["no_receive_time_slots_text", "no receive time slots text", "??????"],
        "illegal_parking_flag": ["illegal_parking_flag", "illegal parking flag", "??"],
        "free_parking_minutes": ["free_parking_minutes", "free parking minutes", "??????_??"],
        "parking_fee_per_hour": ["parking_fee_per_hour", "parking fee per hour", "???_??"],
        "parking_location_text": ["parking_location_text", "parking location text", "????"],
        "parking_mode_text": ["parking_mode_text", "parking mode text", "????"],
        "unload_entrance_text": ["unload_entrance_text", "unload entrance text", "????"],
        "unload_location_text": ["unload_location_text", "unload location text", "????"],
        "access_alley": ["access_alley", "access alley", "???"],
        "access_handcart": ["access_handcart", "access handcart", "???"],
        "access_pallet_exchange": ["access_pallet_exchange", "access pallet exchange", "????"],
        "access_cooler_box_exchange": ["access_cooler_box_exchange", "access cooler box exchange", "?????"],
        "upstairs_floor_count": ["upstairs_floor_count", "upstairs floor count", "????"],
        "basement_height_limit_text": ["basement_height_limit_text", "basement height limit text", "????"],
        "delivery_week_mon": ["delivery_week_mon", "delivery week mon", "????"],
        "delivery_week_tue": ["delivery_week_tue", "delivery week tue", "????"],
        "delivery_week_wed": ["delivery_week_wed", "delivery week wed", "????"],
        "delivery_week_thu": ["delivery_week_thu", "delivery week thu", "????"],
        "delivery_week_fri": ["delivery_week_fri", "delivery week fri", "????"],
        "delivery_week_sat": ["delivery_week_sat", "delivery week sat", "????"],
        "delivery_week_sun": ["delivery_week_sun", "delivery week sun", "????"],
        "order_line_no": ["order_line_no", "order line no", "?????"],
        "source_doc_no": ["source_doc_no", "source doc no", "???"],
        "sales_order_no": ["sales_order_no", "sales order no", "?????"],
        "source_ref_no": ["source_ref_no", "source ref no", "?????"],
        "third_party_doc_no": ["third_party_doc_no", "third party doc no", "?????"],
        "doc_type": ["doc_type", "doc type", "????"],
        "business_type": ["business_type", "business type", "????"],
        "doc_source": ["doc_source", "doc source", "????"],
        "doc_date": ["doc_date", "doc date", "????"],
        "audited_at": ["audited_at", "audited at", "????"],
        "payment_status": ["payment_status", "payment status", "????"],
        "settlement_status": ["settlement_status", "settlement status", "????"],
        "doc_status": ["doc_status", "doc status", "????"],
        "logistics_status": ["logistics_status", "logistics status", "????"],
        "maker_name": ["maker_name", "maker name", "???"],
        "auditor_name": ["auditor_name", "auditor name", "???"],
        "made_at": ["made_at", "made at", "????"],
        "order_remark": ["order_remark", "order remark", "????"],
        "custom_field_1": ["custom_field_1", "custom field 1", "?????1"],
        "external_product_code": ["external_product_code", "external product code", "??"],
        "product_name": ["product_name", "product name", "????"],
        "spec": ["spec", "specification", "????"],
        "barcode": ["barcode", "??"],
        "brand_name": ["brand_name", "brand name", "??"],
        "category_name": ["category_name", "category name", "??"],
        "base_unit_name": ["base_unit_name", "base unit name", "????"],
        "doc_unit_name": ["doc_unit_name", "doc unit name", "????"],
        "small_unit_name": ["small_unit_name", "small unit name", "???"],
        "base_qty": ["base_qty", "base qty", "????"],
        "doc_qty": ["doc_qty", "doc qty", "????"],
        "small_qty": ["small_qty", "small qty", "?????"],
        "box_qty": ["box_qty", "box qty", "??"],
        "gift_qty": ["gift_qty", "gift qty", "????"],
        "exchange_qty": ["exchange_qty", "exchange qty", "????"],
        "unit_price": ["unit_price", "unit price", "????"],
        "small_unit_price": ["small_unit_price", "small unit price", "?????"],
        "amount": ["amount", "??"],
        "settled_amount": ["settled_amount", "settled amount", "?????"],
        "unsettled_amount": ["unsettled_amount", "unsettled amount", "?????"],
        "tax_amount": ["tax_amount", "tax amount", "??"],
        "amount_ex_tax": ["amount_ex_tax", "amount ex tax", "?????"],
        "cost_amount": ["cost_amount", "cost amount", "????"],
        "gross_profit": ["gross_profit", "gross profit", "??"],
        "gross_profit_rate": ["gross_profit_rate", "gross profit rate", "???"],
        "above_standard_price_flag": ["above_standard_price_flag", "above standard price flag", "?????"],
        "below_standard_price_flag": ["below_standard_price_flag", "below standard price flag", "?????"],
        "unit_weight": ["unit_weight", "unit weight", "??????"],
        "unit_volume": ["unit_volume", "unit volume", "??????"],
        "total_weight": ["total_weight", "total weight", "???"],
        "total_volume": ["total_volume", "total volume", "???"],
        "line_remark": ["line_remark", "line remark", "????"],
        "remark": ["remark", "??"],
        "customer_no": ["customer_no", "customer no", "????", "???"],
        "store_no": ["store_no", "store no", "????", "???"],
        "store_name": ["store_name", "store name", "????"],
        "customer_ref": ["customer_ref", "customer ref", "???????"],
        "goods_code": ["goods_code", "goods code", "????"],
        "goods_name": ["goods_name", "goods name", "????"],
        "qty": ["qty", "quantity", "??"],
        "package_count": ["package_count", "package count", "??"],
        "uom_name": ["uom_name", "uom name", "??"],
        "weight": ["weight", "??"],
        "volume": ["volume", "??"],
        "temperature_zone": ["temperature_zone", "temperature zone", "??"],
        "package_type": ["package_type", "package type", "????"],
    }

    @classmethod
    def build_template_payload(cls):
        default_locale = cls.DEFAULT_TEMPLATE_LOCALE
        return {
            "template_code": cls.TEMPLATE_CODE,
            "template_version": cls.TEMPLATE_VERSION,
            "default_template_locale": default_locale,
            "file_name": cls.TEMPLATE_VARIANTS[default_locale]["file_name"],
            "download_url": cls._build_template_download_url(template_locale=default_locale),
            "available_templates": [
                {
                    "template_locale": template_locale,
                    "template_label": template_meta["label"],
                    "file_name": template_meta["file_name"],
                    "download_url": cls._build_template_download_url(template_locale=template_locale),
                }
                for template_locale, template_meta in cls.TEMPLATE_VARIANTS.items()
            ],
            "sheets": [
                {
                    "sheet_name": sheet_meta["sheet_name"],
                    "fields": [
                        {"field_code": field_code, "field_label": cls.FIELD_LABELS[field_code]}
                        for field_code in sheet_meta["fields"]
                    ],
                }
                for sheet_meta in cls.SHEETS
            ],
        }

    @classmethod
    def load_template_bytes(cls, *, template_locale=""):
        template_variant = cls.get_template_variant(template_locale)
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        field_locale = template_variant["field_locale"]
        samples = cls._build_template_samples()
        for sheet_meta in cls.SHEETS:
            worksheet = workbook.create_sheet(sheet_meta["sheet_name"])
            headers = [cls._get_field_header(field_code, field_locale) for field_code in sheet_meta["fields"]]
            worksheet.append(headers)
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
            for sample_row in samples[sheet_meta["key"]]:
                worksheet.append([sample_row.get(field_code, "") for field_code in sheet_meta["fields"]])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @classmethod
    def get_template_variant(cls, template_locale=""):
        normalized_locale = cls.normalize_template_locale(template_locale)
        variant = cls.TEMPLATE_VARIANTS.get(normalized_locale)
        if not variant:
            supported_locales = " / ".join(cls.TEMPLATE_VARIANTS.keys())
            raise ValidationError(f"模板语言非法，请使用 {supported_locales} 之一。")
        return {"template_locale": normalized_locale, **variant}

    @classmethod
    def normalize_template_locale(cls, template_locale=""):
        return (template_locale or cls.DEFAULT_TEMPLATE_LOCALE).strip() or cls.DEFAULT_TEMPLATE_LOCALE

    @classmethod
    def precheck(cls, env, raw_bytes, *, filename="", template_code="", template_version=""):
        template_errors = cls._validate_template_identity(template_code, template_version)
        parsed_data, parse_errors = cls._parse_workbook(raw_bytes)
        errors = template_errors + parse_errors
        if not parse_errors:
            errors.extend(cls._validate_rows(env, parsed_data))

        result = cls._build_precheck_result(parsed_data, errors)
        source_file = cls._create_source_file(
            env,
            raw_bytes=raw_bytes,
            filename=filename or cls.TEMPLATE_FILE_NAME,
        )
        task = cls._create_precheck_task(env, source_file=source_file, result=result)
        row_line_map, row_data_map = cls._create_precheck_task_lines(task, parsed_data, result["errors"])
        cls._create_task_error_lines(task, result["errors"], row_line_map=row_line_map, row_data_map=row_data_map)
        error_report_url = cls._build_task_error_report_url(task) if result["errors"] else False
        result["task_no"] = task.task_no
        result["import_batch_no"] = task.task_no
        result["object_type"] = task.object_type
        result["object_type_label"] = cls._get_selection_label(task._fields["object_type"].selection, task.object_type)
        result["status"] = task.status
        result["status_label"] = cls._get_selection_label(task._fields["status"].selection, task.status)
        result["summary_message"] = task.summary_message
        result["precheck_token"] = False
        result["error_report_url"] = error_report_url
        result["error_report"] = {
            "download_ready": bool(result["errors"]),
            "download_url": error_report_url,
        }
        result["source_file"] = cls._build_source_file_payload(source_file)
        return result

    @classmethod
    def confirm_import(cls, env, *, precheck_token, import_batch_no="", task_no=""):
        if task_no or import_batch_no:
            task_ref = task_no or import_batch_no
            task = cls._get_task_by_task_no(env, task_ref)
            if task:
                return cls._confirm_task_import(env, task)

        batch = cls._get_batch_by_precheck_token(env, precheck_token)
        batch.mark_expired_if_needed()
        if batch.state == "finished":
            return cls._build_result_payload(batch)
        if batch.state == "expired":
            raise ValidationError("预校验令牌已失效，请重新执行预校验。")
        if batch.state == "failed":
            raise ValidationError(batch.failure_reason or "这批数据导入失败，请重新执行预校验。")
        if batch.state != "prechecked":
            raise ValidationError("当前批次状态不允许正式导入。")
        if not batch.can_confirm_import:
            raise ValidationError("预校验未通过，不能执行正式导入。")

        if task_no and not import_batch_no:
            import_batch_no = task_no
        if import_batch_no and import_batch_no != batch.name:
            existing_batch = env["logistics.import.batch"].sudo().search([("name", "=", import_batch_no)], limit=1)
            if existing_batch:
                raise ValidationError("导入批次号已存在，请更换后重试。")
            batch.sudo().write({"name": import_batch_no})

        source_data = batch.get_source_rows()
        stats = {
            "created_waybill_count": 0,
            "created_customer_line_count": 0,
            "created_goods_line_count": 0,
            "updated_record_count": 0,
            "skipped_record_count": 0,
            "failed_record_count": 0,
        }
        batch.sudo().write({"state": "importing", "confirmed_at": fields.Datetime.now(), "failure_reason": False})
        try:
            cls._apply_import(env, source_data, stats)
        except Exception as exc:
            batch.sudo().write(
                {
                    "state": "failed",
                    "finished_at": fields.Datetime.now(),
                    "failed_record_count": cls._count_total_rows(source_data),
                    "failure_reason": str(exc),
                }
            )
            raise ValidationError(f"正式导入失败：{exc}") from exc

        batch.sudo().write({"state": "finished", "finished_at": fields.Datetime.now(), **stats})
        return cls._build_result_payload(batch)

    @classmethod
    def get_import_task_result(cls, env, *, task_no="", import_batch_no=""):
        task = cls._get_task_by_task_no(env, task_no or import_batch_no)
        if task:
            return cls._build_task_result_payload(task)
        batch = cls._get_batch_by_task_ref(env, task_no=task_no, import_batch_no=import_batch_no)
        batch.mark_expired_if_needed()
        return cls._build_result_payload(batch)

    @classmethod
    def get_import_result(cls, env, *, import_batch_no="", task_no=""):
        return cls.get_import_task_result(env, task_no=task_no, import_batch_no=import_batch_no)

    @classmethod
    def build_import_task_error_report(cls, env, *, task_no="", import_batch_no="", precheck_token=""):
        task = cls._get_task_by_task_no(env, task_no or import_batch_no)
        if task:
            return cls._build_task_error_report(task)

        batch = cls._get_batch_for_error_report(
            env,
            task_no=task_no,
            import_batch_no=import_batch_no,
            precheck_token=precheck_token,
        )
        errors = batch.get_errors()
        if not errors:
            raise ValidationError("当前批次没有可导出的预校验错误报告。")

        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([header_label for _, header_label in cls.ERROR_REPORT_HEADERS])
        for error in errors:
            writer.writerow([error.get(field_code, "") for field_code, _label in cls.ERROR_REPORT_HEADERS])

        return {
            "file_name": f"{batch.name}-error-report.csv",
            "file_bytes": buffer.getvalue().encode("utf-8-sig"),
            "content_type": "text/csv; charset=utf-8",
        }

    @classmethod
    def build_error_report(cls, env, *, import_batch_no="", precheck_token="", task_no=""):
        return cls.build_import_task_error_report(
            env,
            task_no=task_no,
            import_batch_no=import_batch_no,
            precheck_token=precheck_token,
        )

    @classmethod
    def get_import_task_lines(cls, env, *, task_no, page=1, page_size=20, status=""):
        task = cls._get_task_by_task_no(env, task_no)
        if not task:
            raise ValidationError("未找到对应的导入任务。")
        page = max(int(page or 1), 1)
        page_size = min(max(int(page_size or 20), 1), 200)
        line_status = (status or "").strip()
        domain = [("task_id", "=", task.id)]
        task_line_model = env["logistics.import.task.line"].sudo()
        valid_statuses = {value for value, _label in task_line_model._fields["status"].selection}
        if line_status:
            if line_status not in valid_statuses:
                raise ValidationError("瀵煎叆琛岀粨鏋滅姸鎬佺瓫閫夊€间笉鍚堟硶銆?")
            domain.append(("status", "=", line_status))
        total = task_line_model.search_count(domain)
        total_pages = max((total + page_size - 1) // page_size, 1)
        page = min(page, total_pages)
        records = task_line_model.search(domain, order="line_no asc", offset=(page - 1) * page_size, limit=page_size)
        showing_from = (page - 1) * page_size + 1 if total else 0
        showing_to = min(page * page_size, total)
        return {
            "task_no": task.task_no,
            "object_type": task.object_type,
            "object_type_label": cls._get_selection_label(task._fields["object_type"].selection, task.object_type),
            "task_status": task.status,
            "task_status_label": cls._get_selection_label(task._fields["status"].selection, task.status),
            "page": page,
            "page_size": page_size,
            "total": total,
            "total_pages": total_pages,
            "has_prev": page > 1,
            "has_next": page < total_pages,
            "showing_from": showing_from,
            "showing_to": showing_to,
            "status_filter": line_status or False,
            "status_filter_label": (
                cls._get_selection_label(task_line_model._fields["status"].selection, line_status) if line_status else False
            ),
            "items": [
                {
                    "line_no": record.line_no,
                    "source_row_no": record.source_row_no,
                    "object_type": record.object_type,
                    "status": record.status,
                    "status_label": cls._get_selection_label(record._fields["status"].selection, record.status),
                    "business_key": record.business_key or False,
                    "target_model": record.target_model or False,
                    "target_res_id": record.target_res_id or False,
                    "message": record.message or False,
                }
                for record in records
            ],
        }

    @classmethod
    def get_import_task_errors(cls, env, *, task_no, page=1, page_size=50):
        task = cls._get_task_by_task_no(env, task_no)
        if not task:
            raise ValidationError("未找到对应的导入任务。")
        page = max(int(page or 1), 1)
        page_size = min(max(int(page_size or 50), 1), 200)
        domain = [("task_id", "=", task.id)]
        error_line_model = env["logistics.import.error.line"].sudo()
        total = error_line_model.search_count(domain)
        total_pages = max((total + page_size - 1) // page_size, 1)
        page = min(page, total_pages)
        records = error_line_model.search(
            domain,
            order="source_row_no asc, id asc",
            offset=(page - 1) * page_size,
            limit=page_size,
        )
        showing_from = (page - 1) * page_size + 1 if total else 0
        showing_to = min(page * page_size, total)
        return {
            "task_no": task.task_no,
            "object_type": task.object_type,
            "object_type_label": cls._get_selection_label(task._fields["object_type"].selection, task.object_type),
            "task_status": task.status,
            "task_status_label": cls._get_selection_label(task._fields["status"].selection, task.status),
            "page": page,
            "page_size": page_size,
            "total": total,
            "total_pages": total_pages,
            "has_prev": page > 1,
            "has_next": page < total_pages,
            "showing_from": showing_from,
            "showing_to": showing_to,
            "items": [
                {
                    "line_no": record.task_line_id.line_no if record.task_line_id else False,
                    "business_key": record.task_line_id.business_key if record.task_line_id else False,
                    "source_row_no": record.source_row_no,
                    "field_name": record.field_name,
                    "raw_value": record.raw_value or False,
                    "mapped_value": record.mapped_value or False,
                    "error_code": record.error_code,
                    "error_message": record.error_message,
                }
                for record in records
            ],
        }

    @classmethod
    def _create_source_file(cls, env, *, raw_bytes, filename):
        file_name = filename or cls.TEMPLATE_FILE_NAME
        file_sha256 = hashlib.sha256(raw_bytes or b"").hexdigest()
        file_ext = cls._guess_file_ext(file_name) or "xlsx"
        source_file = env["logistics.import.source.file"].sudo().create(
            {
                "file_name": file_name,
                "file_ext": file_ext,
                "file_sha256": file_sha256,
                "uploader_id": env.user.id,
            }
        )
        storage_path = cls._write_source_file_bytes(
            source_file_no=source_file.source_file_no,
            file_name=file_name,
            raw_bytes=raw_bytes or b"",
        )
        source_file.sudo().write({"storage_path": storage_path})
        return source_file

    @classmethod
    def _create_precheck_task(cls, env, *, source_file, result):
        summary_message = (
            f"预校验完成，共 {result['total_row_count']} 行，"
            f"通过 {result['passed_row_count']} 行，失败 {result['failed_row_count']} 行。"
        )
        return env["logistics.import.task"].sudo().create(
            {
                "object_type": cls.OBJECT_TYPE,
                "source_file_id": source_file.id,
                "status": "pending",
                "total_count": result["total_row_count"],
                "success_count": result["passed_row_count"],
                "fail_count": result["failed_row_count"],
                "operator_id": env.user.id,
                "summary_message": summary_message,
            }
        )

    @classmethod
    def _create_precheck_task_lines(cls, task, source_data, errors):
        row_errors = cls._group_errors_by_row(errors)
        row_line_map = {}
        row_data_map = {}
        line_vals = []
        for line_no, row in enumerate(cls._iter_source_rows(source_data), start=1):
            row_key = (row["_sheet_name"], row["_source_row_no"])
            business_key = cls._build_row_business_key(row)
            row_line_map[row_key] = business_key
            row_data_map[row_key] = row
            has_error = bool(row_errors.get(row_key))
            line_vals.append(
                {
                    "task_id": task.id,
                    "line_no": line_no,
                    "source_row_no": row["_source_row_no"],
                    "object_type": cls.OBJECT_TYPE,
                    "status": "failed" if has_error else "pending",
                    "business_key": business_key,
                    "message": "预校验失败，请先修正该行错误。" if has_error else "预校验通过，等待正式导入。",
                }
            )
        if line_vals:
            task.env["logistics.import.task.line"].sudo().create(line_vals)
        return row_line_map, row_data_map

    @classmethod
    def _create_task_error_lines(cls, task, errors, *, row_line_map, row_data_map):
        if not errors:
            return
        task_line_model = task.env["logistics.import.task.line"].sudo()
        task_lines = task_line_model.search([("task_id", "=", task.id)])
        task_line_by_business_key = {line.business_key: line for line in task_lines}
        vals_list = []
        for error in errors:
            row_key = (error.get("sheet_name"), error.get("row_no"))
            business_key = row_line_map.get(row_key)
            task_line = task_line_by_business_key.get(business_key) if business_key else False
            row = row_data_map.get(row_key) or {}
            raw_value = row.get(error.get("field_code")) if row and error.get("field_code") in row else False
            vals_list.append(
                {
                    "task_id": task.id,
                    "task_line_id": task_line.id if task_line else False,
                    "source_row_no": error.get("row_no") or 0,
                    "field_name": error.get("field_code") or "template_file",
                    "raw_value": raw_value or False,
                    "mapped_value": raw_value or False,
                    "error_code": error.get("error_code") or "IMPORT_PRECHECK_FAILED",
                    "error_message": error.get("error_message") or "预校验失败。",
                }
            )
        if vals_list:
            task.env["logistics.import.error.line"].sudo().create(vals_list)

    @classmethod
    def _confirm_task_import(cls, env, task):
        task.ensure_one()
        if task.status in ("success", "partial_failed", "failed", "cancelled"):
            return cls._build_task_result_payload(task)
        if task.status == "running":
            raise ValidationError("当前导入任务正在执行，请稍后刷新结果。")
        if task.status != "pending":
            raise ValidationError("当前任务状态不允许正式导入。")
        if task.fail_count or task.error_line_ids:
            raise ValidationError("预校验未通过，不能执行正式导入。")

        source_file = task.source_file_id
        source_bytes = cls._load_source_file_bytes(source_file)
        source_data, parse_errors = cls._parse_workbook(source_bytes)
        if parse_errors:
            cls._create_task_error_lines(
                task,
                parse_errors,
                row_line_map={},
                row_data_map={},
            )
            task.sudo().write(
                {
                    "status": "failed",
                    "started_at": fields.Datetime.now(),
                    "finished_at": fields.Datetime.now(),
                    "success_count": 0,
                    "fail_count": max(task.total_count, 1),
                    "summary_message": "源文件重读失败，无法继续正式导入。",
                }
            )
            raise ValidationError("源文件重读失败，请重新上传并执行预校验。")

        task.sudo().write(
            {
                "status": "running",
                "started_at": fields.Datetime.now(),
                "finished_at": False,
                "summary_message": "任务执行中，请稍后刷新结果。",
            }
        )
        execute_stats = cls._execute_task_import(env, task, source_data)
        task.sudo().write(
            {
                "status": execute_stats["task_status"],
                "success_count": execute_stats["success_count"],
                "fail_count": execute_stats["fail_count"],
                "finished_at": fields.Datetime.now(),
                "summary_message": execute_stats["summary_message"],
            }
        )
        return cls._build_task_result_payload(task)

    @classmethod
    def _execute_task_import(cls, env, task, source_data):
        if source_data.get("input_mode") == cls.FORMAL_INPUT_MODE:
            return cls._execute_formal_task_import(env, task, source_data)
        task_line_map = {line.business_key: line for line in task.task_line_ids}
        error_line_model = env["logistics.import.error.line"].sudo()
        wave_model = env["logistics.dispatch.wave"].sudo()
        batch_model = env["logistics.dispatch.batch"].sudo()
        waybill_model = env["logistics.dispatch.waybill"].sudo()
        customer_line_model = env["logistics.dispatch.waybill.customer.line"].sudo()
        goods_line_model = env["logistics.dispatch.waybill.customer.goods.line"].sudo()

        success_count = 0
        fail_count = 0
        skipped_count = 0
        created_wave_count = 0
        created_batch_count = 0
        created_waybill_count = 0
        created_customer_line_count = 0
        created_goods_line_count = 0

        rows = source_data.get("rows", [])
        warehouse_map = cls._search_record_map(
            env,
            model_name="stock.warehouse",
            field_name="code",
            values={row["warehouse_code"] for row in rows if row.get("warehouse_code")},
        )
        customers_by_code = cls._search_partner_map(
            env,
            model_domain=[("is_logistics_customer", "=", True)],
            field_name="logistics_customer_code",
            values={row["customer_no"] for row in rows if row.get("customer_no")},
        )
        stores_by_code = cls._search_partner_map(
            env,
            model_domain=[("is_logistics_store", "=", True)],
            field_name="logistics_store_code",
            values={row["store_no"] for row in rows if row.get("store_no")},
        )
        existing_waves = wave_model.search([("name", "in", [row["wave_no"] for row in rows if row.get("wave_no")])])
        existing_batches = batch_model.search([("name", "in", [row["batch_no"] for row in rows if row.get("batch_no")])])
        existing_waybills = waybill_model.search([("name", "in", [row["waybill_no"] for row in rows if row.get("waybill_no")])])
        wave_record_map = {record.name: record for record in existing_waves}
        batch_record_map = {record.name: record for record in existing_batches}
        existing_waybill_map = {record.name: record for record in existing_waybills}
        waybill_record_map = {}
        customer_line_record_map = {}
        customer_line_no_map = {}
        customer_line_seq_by_waybill = {}

        for row in rows:
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                warehouse = warehouse_map.get(row.get("warehouse_code"))
                created_flags = []

                wave = False
                if row.get("wave_no"):
                    wave = wave_record_map.get(row["wave_no"])
                    if not wave:
                        wave = wave_model.create(
                            {
                                "wave_no": row["wave_no"],
                                "dispatch_date": row.get("delivery_date") or False,
                                "warehouse_id": warehouse.id,
                            }
                        )
                        wave_record_map[row["wave_no"]] = wave
                        created_wave_count += 1
                        created_flags.append("新建波次")

                batch = batch_record_map.get(row["batch_no"])
                if not batch:
                    batch_vals = {
                        "batch_no": row["batch_no"],
                        "warehouse_id": warehouse.id,
                    }
                    if wave:
                        batch_vals["wave_id"] = wave.id
                    batch = batch_model.create(batch_vals)
                    batch_record_map[row["batch_no"]] = batch
                    created_batch_count += 1
                    created_flags.append("新建批次")

                waybill = waybill_record_map.get(row["waybill_no"])
                if not waybill:
                    if row["waybill_no"] in existing_waybill_map:
                        raise ValidationError("运单号已存在，当前模式仅允许新建主链。")
                    waybill = waybill_model.create(
                        {
                            "waybill_no": row["waybill_no"],
                            "delivery_date": row.get("delivery_date") or False,
                            "remark": row.get("remark") or False,
                            "batch_id": batch.id,
                            "warehouse_id": warehouse.id,
                        }
                    )
                    waybill_record_map[row["waybill_no"]] = waybill
                    created_waybill_count += 1
                    created_flags.append("新建运单")

                customer_line_cache_key = cls._customer_group_key(row)
                customer_line = customer_line_record_map.get(customer_line_cache_key)
                if not customer_line:
                    effective_customer_line_no = cls._get_effective_customer_line_no(
                        row,
                        customer_line_no_map=customer_line_no_map,
                        customer_line_seq_by_waybill=customer_line_seq_by_waybill,
                    )
                    customer_line = customer_line_model.create(
                        cls._build_customer_line_create_vals(
                            row,
                            waybill=waybill,
                            customer_line_no=effective_customer_line_no,
                            customers_by_code=customers_by_code,
                            stores_by_code=stores_by_code,
                        )
                    )
                    customer_line_record_map[customer_line_cache_key] = customer_line
                    created_customer_line_count += 1
                    created_flags.append("新建门店节点")

                goods_line = goods_line_model.create(
                    {
                        "customer_line_id": customer_line.id,
                        "goods_code": row.get("goods_code") or False,
                        "goods_name": row["goods_name"],
                        "specification": row.get("spec") or False,
                        "quantity": cls._safe_float(row.get("qty"), default=0.0),
                        "package_count": cls._safe_int(row.get("package_count"), default=0),
                        "uom_name": row.get("uom_name") or False,
                        "weight": cls._safe_float(row.get("weight"), default=0.0),
                        "volume": cls._safe_float(row.get("volume"), default=0.0),
                        "temperature_zone": row.get("temperature_zone") or "ambient",
                        "package_type": row.get("package_type") or False,
                        "remark": row.get("remark") or row.get("delivery_remark") or False,
                    }
                )
                created_goods_line_count += 1
                success_count += 1
                cls._update_task_line(
                    task_line,
                    status="success",
                    message="、".join(created_flags) + "，本行导入成功。" if created_flags else "复用既有波次/批次，本行导入成功。",
                    target_model="logistics.dispatch.waybill.customer.goods.line",
                    target_res_id=goods_line.id,
                )
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"单表导入失败：{exc}")
                error_line_model.create(
                    cls._make_runtime_error_line_vals(
                        task=task,
                        task_line=task_line,
                        row=row,
                        field_name="waybill_no",
                        error_message=f"单表导入失败：{exc}",
                    )
                )

        task_status = "success"
        if fail_count and success_count:
            task_status = "partial_failed"
        elif fail_count and not success_count:
            task_status = "failed"
        summary_message = (
            f"本次任务已完成，共 {task.total_count} 行，"
            f"成功 {success_count} 行，失败 {fail_count} 行，跳过 {skipped_count} 行。"
        )
        return {
            "task_status": task_status,
            "success_count": success_count,
            "fail_count": fail_count,
            "skipped_count": skipped_count,
            "created_wave_count": created_wave_count,
            "created_batch_count": created_batch_count,
            "created_waybill_count": created_waybill_count,
            "created_customer_line_count": created_customer_line_count,
            "created_goods_line_count": created_goods_line_count,
            "summary_message": summary_message,
        }

    @classmethod
    def _build_task_result_payload(cls, task):
        task.ensure_one()
        counts = cls._get_task_line_status_counts(task)
        target_counts = cls._get_task_target_counts(task)
        business_counts = cls._get_task_business_counts(task)
        source_file = task.source_file_id
        return {
            "task_no": task.task_no,
            "import_batch_no": task.task_no,
            "object_type": task.object_type,
            "object_type_label": cls._get_selection_label(task._fields["object_type"].selection, task.object_type),
            "status": task.status,
            "status_label": cls._get_selection_label(task._fields["status"].selection, task.status),
            "total_count": task.total_count,
            "success_count": task.success_count,
            "fail_count": task.fail_count,
            "summary_message": task.summary_message or cls._build_task_summary_message(task, counts),
            "template_code": cls.TEMPLATE_CODE,
            "template_version": cls.TEMPLATE_VERSION,
            "file_name": source_file.file_name,
            "total_row_count": task.total_count,
            "passed_row_count": task.success_count,
            "failed_row_count": task.fail_count,
            "confirmed_at": fields.Datetime.to_string(task.started_at) if task.started_at else False,
            "started_at": fields.Datetime.to_string(task.started_at) if task.started_at else False,
            "finished_at": fields.Datetime.to_string(task.finished_at) if task.finished_at else False,
            "created_wave_count": business_counts["wave"],
            "created_batch_count": business_counts["batch"],
            "created_waybill_count": business_counts["waybill"] or target_counts["waybill"],
            "created_customer_line_count": business_counts["customer_line"] or target_counts["customer_line"],
            "created_order_line_count": business_counts["order_line"] or target_counts["order_line"],
            "created_goods_line_count": business_counts["goods_line"] or target_counts["goods_line"],
            "updated_record_count": 0,
            "skipped_record_count": counts["skipped"],
            "failed_record_count": task.fail_count,
            "business_summary": {
                "written_wave_count": business_counts["wave"],
                "written_batch_count": business_counts["batch"],
                "written_waybill_count": business_counts["waybill"] or target_counts["waybill"],
                "written_customer_line_count": business_counts["customer_line"] or target_counts["customer_line"],
                "written_order_line_count": business_counts["order_line"] or target_counts["order_line"],
                "written_goods_line_count": business_counts["goods_line"] or target_counts["goods_line"],
            },
            "source_file": cls._build_source_file_payload(source_file),
            "operator": cls._build_operator_payload(task.operator_id),
            "error_report": {
                "download_ready": bool(task.error_line_ids),
                "download_url": cls._build_task_error_report_url(task) if task.error_line_ids else False,
            },
            "next_actions": {
                "refresh": True,
                "go_import_center": True,
                "go_waybill_list": True,
                "review_by_task": True,
            },
            "error_report_url": cls._build_task_error_report_url(task) if task.error_line_ids else False,
            "failure_reason": cls._build_task_failure_reason(task),
        }

    @classmethod
    def _build_task_error_report(cls, task):
        errors = task.error_line_ids.sorted(key=lambda rec: (rec.source_row_no, rec.id))
        if not errors:
            raise ValidationError("当前任务没有可导出的错误报告。")
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(["任务号", "原始行号", "字段名", "原始值", "映射值", "错误码", "错误说明"])
        for error in errors:
            writer.writerow(
                [
                    task.task_no,
                    error.source_row_no,
                    error.field_name,
                    error.raw_value or "",
                    error.mapped_value or "",
                    error.error_code,
                    error.error_message,
                ]
            )
        return {
            "file_name": f"{task.task_no}-error-report.csv",
            "file_bytes": buffer.getvalue().encode("utf-8-sig"),
            "content_type": "text/csv; charset=utf-8",
        }

    @classmethod
    def _build_template_samples(cls):
        return {
            "waybill_rows": [
                {
                    "warehouse_code": "WH",
                    "delivery_date": "2026-04-24",
                    "wave_no": "FS-WV-SAMPLE-01",
                    "batch_no": "FS-BT-SAMPLE-01",
                    "waybill_no": "FS-WB-SAMPLE-01",
                    "organization_name": "测试组织",
                    "route_name": "沪东线A",
                    "driver_name": "",
                    "driver_phone": "",
                    "delivery_remark": "四 Sheet 中文样稿示例-运单1",
                },
                {
                    "warehouse_code": "WH",
                    "delivery_date": "2026-04-24",
                    "wave_no": "FS-WV-SAMPLE-01",
                    "batch_no": "FS-BT-SAMPLE-01",
                    "waybill_no": "FS-WB-SAMPLE-02",
                    "organization_name": "测试组织",
                    "route_name": "沪东线A",
                    "driver_name": "",
                    "driver_phone": "",
                    "delivery_remark": "四 Sheet 中文样稿示例-运单2",
                },
            ],
            "customer_line_rows": [
                {
                    "waybill_no": "FS-WB-SAMPLE-01",
                    "customer_line_no": "CL-001",
                    "external_customer_code": "FS-CUST-SAMPLE-01",
                    "customer_name": "四表门店A",
                    "contact_name": "张三",
                    "contact_phone": "13910000001",
                    "organization_name": "测试组织",
                    "department_name": "测试部门",
                    "salesperson_name": "业务员甲",
                    "channel_name": "商超",
                    "customer_level": "VIP",
                    "allow_cash_on_delivery": "N",
                    "customer_status": "正常",
                    "internal_counterparty_flag": "N",
                    "invoice_type": "专票",
                    "registered_phone": "021-10000001",
                    "registered_address": "上海市浦东新区测试路1号",
                    "customer_seq_no": "SEQ-SAMPLE-01",
                    "address_full": "上海市浦东新区测试路1号",
                    "province_name": "上海市",
                    "city_name": "上海市",
                    "district_name": "浦东新区",
                    "longitude": "121.5442",
                    "latitude": "31.2211",
                    "stop_seq_in_waybill": "1",
                    "route_preference": "沪东线A",
                    "warehouse_preference": "WH",
                    "receive_start_time": "08:00",
                    "receive_end_time": "18:00",
                    "receive_time_slots_text": "08:00-18:00",
                    "no_receive_time_slots_text": "12:00-13:00",
                    "illegal_parking_flag": "N",
                    "free_parking_minutes": "30",
                    "parking_fee_per_hour": "10",
                    "parking_location_text": "门口停车位",
                    "parking_mode_text": "路边停车",
                    "unload_entrance_text": "东门",
                    "unload_location_text": "收货区A",
                    "access_alley": "Y",
                    "access_handcart": "Y",
                    "access_pallet_exchange": "N",
                    "access_cooler_box_exchange": "N",
                    "upstairs_floor_count": "1",
                    "basement_height_limit_text": "2.2m",
                    "delivery_week_mon": "Y",
                    "delivery_week_tue": "Y",
                    "delivery_week_wed": "Y",
                    "delivery_week_thu": "Y",
                    "delivery_week_fri": "Y",
                    "delivery_week_sat": "N",
                    "delivery_week_sun": "N",
                },
                {
                    "waybill_no": "FS-WB-SAMPLE-01",
                    "customer_line_no": "CL-002",
                    "external_customer_code": "FS-CUST-SAMPLE-02",
                    "customer_name": "四表门店B",
                    "contact_name": "李四",
                    "contact_phone": "13910000002",
                    "organization_name": "测试组织",
                    "department_name": "测试部门",
                    "salesperson_name": "业务员乙",
                    "channel_name": "餐饮",
                    "customer_level": "standard",
                    "allow_cash_on_delivery": "Y",
                    "customer_status": "正常",
                    "internal_counterparty_flag": "N",
                    "invoice_type": "普票",
                    "registered_phone": "021-10000002",
                    "registered_address": "上海市浦东新区测试路2号",
                    "customer_seq_no": "SEQ-SAMPLE-02",
                    "address_full": "上海市浦东新区测试路2号",
                    "province_name": "上海市",
                    "city_name": "上海市",
                    "district_name": "浦东新区",
                    "longitude": "121.5002",
                    "latitude": "31.2102",
                    "stop_seq_in_waybill": "2",
                    "route_preference": "沪东线A",
                    "warehouse_preference": "WH",
                    "receive_start_time": "09:00",
                    "receive_end_time": "17:00",
                    "receive_time_slots_text": "09:00-17:00",
                    "illegal_parking_flag": "Y",
                    "free_parking_minutes": "15",
                    "parking_fee_per_hour": "12",
                    "parking_location_text": "商场停车场",
                    "parking_mode_text": "地下停车",
                    "unload_entrance_text": "南门",
                    "unload_location_text": "收货区B",
                    "access_alley": "N",
                    "access_handcart": "Y",
                    "access_pallet_exchange": "Y",
                    "access_cooler_box_exchange": "N",
                    "upstairs_floor_count": "0",
                    "basement_height_limit_text": "2.0m",
                    "delivery_week_mon": "Y",
                    "delivery_week_tue": "N",
                    "delivery_week_wed": "Y",
                    "delivery_week_thu": "N",
                    "delivery_week_fri": "Y",
                    "delivery_week_sat": "Y",
                    "delivery_week_sun": "N",
                },
                {
                    "waybill_no": "FS-WB-SAMPLE-02",
                    "customer_line_no": "CL-001",
                    "external_customer_code": "FS-CUST-SAMPLE-03",
                    "customer_name": "四表门店C",
                    "contact_name": "王五",
                    "contact_phone": "13910000003",
                    "organization_name": "测试组织",
                    "department_name": "测试部门",
                    "salesperson_name": "业务员丙",
                    "channel_name": "便利店",
                    "customer_level": "strategic",
                    "allow_cash_on_delivery": "N",
                    "customer_status": "正常",
                    "internal_counterparty_flag": "Y",
                    "invoice_type": "专票",
                    "registered_phone": "021-10000003",
                    "registered_address": "上海市闵行区测试路3号",
                    "customer_seq_no": "SEQ-SAMPLE-03",
                    "address_full": "上海市闵行区测试路3号",
                    "province_name": "上海市",
                    "city_name": "上海市",
                    "district_name": "闵行区",
                    "longitude": "121.4003",
                    "latitude": "31.1003",
                    "stop_seq_in_waybill": "1",
                    "route_preference": "沪南线B",
                    "warehouse_preference": "WH",
                    "receive_start_time": "07:30",
                    "receive_end_time": "15:30",
                    "receive_time_slots_text": "07:30-15:30",
                    "no_receive_time_slots_text": "11:30-12:30",
                    "illegal_parking_flag": "N",
                    "free_parking_minutes": "20",
                    "parking_fee_per_hour": "8",
                    "parking_location_text": "后场停车位",
                    "parking_mode_text": "露天停车",
                    "unload_entrance_text": "西门",
                    "unload_location_text": "收货区C",
                    "access_alley": "Y",
                    "access_handcart": "N",
                    "access_pallet_exchange": "N",
                    "access_cooler_box_exchange": "Y",
                    "upstairs_floor_count": "2",
                    "basement_height_limit_text": "2.4m",
                    "delivery_week_mon": "N",
                    "delivery_week_tue": "Y",
                    "delivery_week_wed": "N",
                    "delivery_week_thu": "Y",
                    "delivery_week_fri": "N",
                    "delivery_week_sat": "Y",
                    "delivery_week_sun": "Y",
                },
            ],
            "order_line_rows": [
                {
                    "order_line_no": "FS-OL-SAMPLE-01",
                    "waybill_no": "FS-WB-SAMPLE-01",
                    "customer_line_no": "CL-001",
                    "source_doc_no": "DOC-SAMPLE-01",
                    "sales_order_no": "SO-SAMPLE-01",
                    "source_ref_no": "REF-SAMPLE-01",
                    "third_party_doc_no": "TP-SAMPLE-01",
                    "doc_type": "sales",
                    "business_type": "to_store",
                    "doc_source": "erp",
                    "doc_date": "2026-04-23",
                    "audited_at": "2026-04-23",
                    "department_name": "测试部门",
                    "channel_name": "商超",
                    "salesperson_name": "业务员甲",
                    "payment_status": "unpaid",
                    "settlement_status": "unsettled",
                    "doc_status": "confirmed",
                    "logistics_status": "pending_outbound",
                    "maker_name": "制单员A",
                    "auditor_name": "审核员A",
                    "made_at": "2026-04-23 09:00:00",
                    "order_remark": "四 Sheet 中文样稿-订单1",
                    "custom_field_1": "扩展字段1",
                },
                {
                    "order_line_no": "FS-OL-SAMPLE-02",
                    "waybill_no": "FS-WB-SAMPLE-01",
                    "customer_line_no": "CL-002",
                    "source_doc_no": "DOC-SAMPLE-02",
                    "sales_order_no": "SO-SAMPLE-02",
                    "source_ref_no": "REF-SAMPLE-02",
                    "third_party_doc_no": "TP-SAMPLE-02",
                    "doc_type": "sales",
                    "business_type": "to_store",
                    "doc_source": "oms",
                    "doc_date": "2026-04-23",
                    "audited_at": "2026-04-23",
                    "department_name": "测试部门",
                    "channel_name": "餐饮",
                    "salesperson_name": "业务员乙",
                    "payment_status": "paid",
                    "settlement_status": "settled",
                    "doc_status": "confirmed",
                    "logistics_status": "outbounded",
                    "maker_name": "制单员B",
                    "auditor_name": "审核员B",
                    "made_at": "2026-04-23 09:10:00",
                    "order_remark": "四 Sheet 中文样稿-订单2",
                    "custom_field_1": "扩展字段2",
                },
                {
                    "order_line_no": "FS-OL-SAMPLE-03",
                    "waybill_no": "FS-WB-SAMPLE-02",
                    "customer_line_no": "CL-001",
                    "source_doc_no": "DOC-SAMPLE-03",
                    "sales_order_no": "SO-SAMPLE-03",
                    "source_ref_no": "REF-SAMPLE-03",
                    "third_party_doc_no": "TP-SAMPLE-03",
                    "doc_type": "transfer",
                    "business_type": "to_store",
                    "doc_source": "erp",
                    "doc_date": "2026-04-23",
                    "audited_at": "2026-04-23",
                    "department_name": "测试部门",
                    "channel_name": "便利店",
                    "salesperson_name": "业务员丙",
                    "payment_status": "partial_paid",
                    "settlement_status": "partial_settled",
                    "doc_status": "draft",
                    "logistics_status": "pending_outbound",
                    "maker_name": "制单员C",
                    "auditor_name": "审核员C",
                    "made_at": "2026-04-23 09:20:00",
                    "order_remark": "四 Sheet 中文样稿-订单3",
                    "custom_field_1": "扩展字段3",
                },
            ],
            "goods_line_rows": [
                {
                    "order_line_no": "FS-OL-SAMPLE-01",
                    "external_product_code": "SKU-SAMPLE-01",
                    "product_name": "可乐330ml",
                    "spec": "24瓶/箱",
                    "barcode": "690000000001",
                    "brand_name": "品牌A",
                    "category_name": "饮料",
                    "base_unit_name": "瓶",
                    "doc_unit_name": "箱",
                    "small_unit_name": "瓶",
                    "base_qty": "24",
                    "doc_qty": "1",
                    "small_qty": "24",
                    "box_qty": "1",
                    "gift_qty": "0",
                    "exchange_qty": "0",
                    "unit_price": "80",
                    "small_unit_price": "3.3333",
                    "amount": "80",
                    "settled_amount": "0",
                    "unsettled_amount": "80",
                    "tax_amount": "7.2",
                    "amount_ex_tax": "72.8",
                    "cost_amount": "60",
                    "gross_profit": "20",
                    "gross_profit_rate": "0.25",
                    "above_standard_price_flag": "N",
                    "below_standard_price_flag": "N",
                    "unit_weight": "0.5",
                    "unit_volume": "0.0012",
                    "total_weight": "12",
                    "total_volume": "0.0288",
                    "line_remark": "四 Sheet 中文样稿-货品1",
                },
                {
                    "order_line_no": "FS-OL-SAMPLE-01",
                    "external_product_code": "SKU-SAMPLE-02",
                    "product_name": "雪碧330ml",
                    "spec": "24瓶/箱",
                    "barcode": "690000000002",
                    "brand_name": "品牌A",
                    "category_name": "饮料",
                    "base_unit_name": "瓶",
                    "doc_unit_name": "箱",
                    "small_unit_name": "瓶",
                    "base_qty": "24",
                    "doc_qty": "2",
                    "small_qty": "48",
                    "box_qty": "2",
                    "gift_qty": "0",
                    "exchange_qty": "0",
                    "unit_price": "78",
                    "small_unit_price": "3.25",
                    "amount": "156",
                    "settled_amount": "0",
                    "unsettled_amount": "156",
                    "tax_amount": "14.04",
                    "amount_ex_tax": "141.96",
                    "cost_amount": "120",
                    "gross_profit": "36",
                    "gross_profit_rate": "0.2308",
                    "above_standard_price_flag": "N",
                    "below_standard_price_flag": "Y",
                    "unit_weight": "0.5",
                    "unit_volume": "0.0012",
                    "total_weight": "24",
                    "total_volume": "0.0576",
                    "line_remark": "四 Sheet 中文样稿-货品2",
                },
                {
                    "order_line_no": "FS-OL-SAMPLE-02",
                    "external_product_code": "SKU-SAMPLE-03",
                    "product_name": "矿泉水1L",
                    "spec": "12瓶/箱",
                    "barcode": "690000000003",
                    "brand_name": "品牌B",
                    "category_name": "饮料",
                    "base_unit_name": "瓶",
                    "doc_unit_name": "箱",
                    "small_unit_name": "瓶",
                    "base_qty": "12",
                    "doc_qty": "1",
                    "small_qty": "12",
                    "box_qty": "1",
                    "gift_qty": "0",
                    "exchange_qty": "0",
                    "unit_price": "36",
                    "small_unit_price": "3",
                    "amount": "36",
                    "settled_amount": "36",
                    "unsettled_amount": "0",
                    "tax_amount": "3.24",
                    "amount_ex_tax": "32.76",
                    "cost_amount": "24",
                    "gross_profit": "12",
                    "gross_profit_rate": "0.3333",
                    "above_standard_price_flag": "Y",
                    "below_standard_price_flag": "N",
                    "unit_weight": "1",
                    "unit_volume": "0.0015",
                    "total_weight": "12",
                    "total_volume": "0.018",
                    "line_remark": "四 Sheet 中文样稿-货品3",
                },
                {
                    "order_line_no": "FS-OL-SAMPLE-03",
                    "external_product_code": "SKU-SAMPLE-04",
                    "product_name": "茶饮料550ml",
                    "spec": "15瓶/箱",
                    "barcode": "690000000004",
                    "brand_name": "品牌C",
                    "category_name": "饮料",
                    "base_unit_name": "瓶",
                    "doc_unit_name": "箱",
                    "small_unit_name": "瓶",
                    "base_qty": "15",
                    "doc_qty": "3",
                    "small_qty": "45",
                    "box_qty": "3",
                    "gift_qty": "1",
                    "exchange_qty": "0",
                    "unit_price": "55",
                    "small_unit_price": "3.6667",
                    "amount": "165",
                    "settled_amount": "80",
                    "unsettled_amount": "85",
                    "tax_amount": "14.85",
                    "amount_ex_tax": "150.15",
                    "cost_amount": "100",
                    "gross_profit": "65",
                    "gross_profit_rate": "0.3939",
                    "above_standard_price_flag": "N",
                    "below_standard_price_flag": "N",
                    "unit_weight": "0.7",
                    "unit_volume": "0.0010",
                    "total_weight": "31.5",
                    "total_volume": "0.045",
                    "line_remark": "四 Sheet 中文样稿-货品4",
                },
            ],
        }

    @classmethod
    def _get_field_header(cls, field_code, field_locale):
        if field_locale == "en_US":
            return field_code
        return cls.FIELD_LABELS[field_code]

    @classmethod
    def _validate_template_identity(cls, template_code, template_version):
        errors = []
        normalized_code = (template_code or cls.TEMPLATE_CODE).strip()
        normalized_version = (template_version or cls.TEMPLATE_VERSION).strip().lower()
        valid_codes = {cls.TEMPLATE_CODE, *cls.LEGACY_TEMPLATE_CODES}
        valid_versions = {cls.TEMPLATE_VERSION, *cls.LEGACY_TEMPLATE_VERSIONS}
        if normalized_code not in valid_codes:
            errors.append(
                cls._make_error(
                    sheet_name="模板文件",
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_CODE_INVALID",
                    error_message=f"请使用标准模板 {cls.TEMPLATE_CODE}。",
                )
            )
        if normalized_version not in valid_versions:
            errors.append(
                cls._make_error(
                    sheet_name="模板文件",
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_VERSION_INVALID",
                    error_message=f"请使用模板版本 {cls.TEMPLATE_VERSION}。",
                )
            )
        return errors

    @classmethod
    def _parse_workbook(cls, raw_bytes):
        if not raw_bytes:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name="????",
                    row_no=0,
                    field_code="template_file",
                    error_code="PRECHECK_PARSE_FAILED",
                    error_message="??????????????????????",
                )
            ]

        try:
            workbook = load_workbook(io.BytesIO(raw_bytes), data_only=True)
        except Exception:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name="????",
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_FILE_TYPE_INVALID",
                    error_message="?????????? .xlsx ?????",
                )
            ]

        if all(sheet_meta["sheet_name"] in workbook.sheetnames for sheet_meta in cls.SHEETS):
            return cls._parse_formal_workbook(workbook)

        primary_sheet_name = cls.PRIMARY_SHEET["sheet_name"]
        if primary_sheet_name in workbook.sheetnames:
            return cls._parse_single_sheet_workbook(workbook, sheet_name=primary_sheet_name)

        first_sheet_name = workbook.sheetnames[0] if workbook.sheetnames else ""
        if first_sheet_name:
            single_data, single_errors = cls._parse_single_sheet_workbook(
                workbook,
                sheet_name=first_sheet_name,
                strict_sheet_name=False,
            )
            if not single_errors:
                return single_data, []

        if all(sheet_meta["sheet_name"] in workbook.sheetnames for sheet_meta in cls.LEGACY_SHEETS):
            return cls._parse_legacy_workbook(workbook)

        return cls._empty_source_data(), [
            cls._make_error(
                sheet_name="????",
                row_no=0,
                field_code="template_file",
                error_code="TEMPLATE_SHEET_MISSING",
                error_message="??????? Waybill / CustomerLine / OrderLine / GoodsLine ??????",
            )
        ]

    @classmethod
    def _parse_single_sheet_workbook(cls, workbook, *, sheet_name, strict_sheet_name=True):
        if strict_sheet_name and sheet_name not in workbook.sheetnames:
            return cls._empty_source_data(), [
                cls._make_error(
                    sheet_name="????",
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_SHEET_MISSING",
                    error_message=f"????????{cls.PRIMARY_SHEET['sheet_name']}?",
                )
            ]
        worksheet = workbook[sheet_name]
        parsed_rows, errors = cls._parse_sheet_rows(worksheet, cls.PRIMARY_SHEET, sheet_name_override=sheet_name)
        if errors:
            return cls._empty_source_data(), errors
        return {"rows": parsed_rows, "input_mode": "single_sheet"}, []

    @classmethod
    def _parse_legacy_workbook(cls, workbook):
        legacy_data = {sheet_meta["key"]: [] for sheet_meta in cls.LEGACY_SHEETS}
        errors = []
        for sheet_meta in cls.LEGACY_SHEETS:
            worksheet = workbook[sheet_meta["sheet_name"]]
            parsed_rows, sheet_errors = cls._parse_sheet_rows(worksheet, sheet_meta)
            if sheet_errors:
                errors.extend(sheet_errors)
                continue
            legacy_data[sheet_meta["key"]] = parsed_rows
        if errors:
            return cls._empty_source_data(), errors
        return {
            "rows": cls._normalize_legacy_rows(legacy_data),
            "input_mode": "legacy_patch",
        }, []

    @classmethod
    def _parse_formal_workbook(cls, workbook):
        source_data = {
            "input_mode": cls.FORMAL_INPUT_MODE,
            "waybill_rows": [],
            "customer_line_rows": [],
            "order_line_rows": [],
            "goods_line_rows": [],
        }
        errors = []
        for sheet_meta in cls.SHEETS:
            worksheet = workbook[sheet_meta["sheet_name"]]
            parsed_rows, sheet_errors = cls._parse_sheet_rows(worksheet, sheet_meta)
            if sheet_errors:
                errors.extend(sheet_errors)
                continue
            source_data[sheet_meta["key"]] = parsed_rows
        if errors:
            return cls._empty_source_data(), errors
        return source_data, []

    @classmethod
    def _parse_sheet_rows(cls, worksheet, sheet_meta, *, sheet_name_override=""):
        rows = list(worksheet.iter_rows(values_only=True))
        sheet_name = sheet_name_override or sheet_meta["sheet_name"]
        if not rows:
            return [], [
                cls._make_error(
                    sheet_name=sheet_name,
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_HEADER_MISMATCH",
                    error_message=f"{sheet_name} ????????",
                )
            ]
        header_values = [cls._cell_text(value) for value in rows[0]]
        header_map, header_errors = cls._build_header_map(sheet_meta, header_values)
        if header_errors:
            return [], header_errors
        parsed_rows = []
        for row_index, row_values in enumerate(rows[1:], start=2):
            normalized_row = {field_code: "" for field_code in sheet_meta["fields"]}
            has_value = False
            for col_index, header_value in enumerate(header_values):
                if not header_value:
                    continue
                field_code = header_map.get(col_index)
                if not field_code:
                    continue
                cell_value = cls._cell_text(row_values[col_index] if col_index < len(row_values) else "")
                normalized_row[field_code] = cell_value
                if cell_value:
                    has_value = True
            if not has_value:
                continue
            normalized_row["_sheet_name"] = sheet_name
            normalized_row["_source_row_no"] = row_index
            parsed_rows.append(normalized_row)
        if not parsed_rows:
            return [], [
                cls._make_error(
                    sheet_name=sheet_name,
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_SHEET_EMPTY",
                    error_message=f"{sheet_name} ???????????",
                )
            ]
        return parsed_rows, []

    @classmethod
    def _normalize_legacy_rows(cls, legacy_data):
        waybill_rows = {
            row["waybill_no"]: row
            for row in legacy_data.get("waybill_rows", [])
            if row.get("waybill_no")
        }
        customer_rows = {}
        for row in legacy_data.get("customer_rows", []):
            customer_rows[cls._legacy_customer_key(row)] = row

        normalized_rows = []
        customer_line_no_map = {}
        for row in legacy_data.get("goods_rows", []):
            waybill_row = waybill_rows.get(row.get("waybill_no"))
            customer_row = customer_rows.get(cls._legacy_customer_key(row))
            customer_line_key = cls._legacy_customer_key(row)
            customer_line_no = customer_line_no_map.get(customer_line_key)
            if not customer_line_no:
                customer_line_no = cls._build_legacy_customer_line_no(
                    row.get("waybill_no"),
                    row.get("store_no"),
                    row.get("customer_no"),
                    len(customer_line_no_map) + 1,
                )
                customer_line_no_map[customer_line_key] = customer_line_no
            normalized_rows.append(
                {
                    "warehouse_code": waybill_row.get("warehouse_code", "") if waybill_row else "",
                    "delivery_date": waybill_row.get("delivery_date", "") if waybill_row else "",
                    "wave_no": waybill_row.get("wave_no", "") if waybill_row else "",
                    "batch_no": waybill_row.get("batch_no", "") if waybill_row else "",
                    "waybill_no": row.get("waybill_no", ""),
                    "driver_name": waybill_row.get("driver_name", "") if waybill_row else "",
                    "driver_phone": waybill_row.get("driver_phone", "") if waybill_row else "",
                    "remark": row.get("remark", "") or (waybill_row.get("remark", "") if waybill_row else ""),
                    "customer_line_no": customer_line_no,
                    "customer_no": row.get("customer_no", "") or (customer_row.get("customer_no", "") if customer_row else ""),
                    "customer_name": customer_row.get("customer_name", "") if customer_row else "",
                    "store_no": row.get("store_no", "") or (customer_row.get("store_no", "") if customer_row else ""),
                    "store_name": customer_row.get("store_name", "") if customer_row else "",
                    "delivery_remark": customer_row.get("delivery_remark", "") if customer_row else "",
                    "signoff_requirement": customer_row.get("signoff_requirement", "") if customer_row else "",
                    "customer_ref": customer_row.get("customer_ref", "") if customer_row else "",
                    "goods_code": row.get("goods_code", ""),
                    "goods_name": row.get("goods_name", ""),
                    "spec": row.get("spec", ""),
                    "qty": row.get("qty", ""),
                    "package_count": row.get("package_count", ""),
                    "uom_name": row.get("uom_name", ""),
                    "weight": row.get("weight", ""),
                    "volume": row.get("volume", ""),
                    "temperature_zone": row.get("temperature_zone", ""),
                    "package_type": row.get("package_type", ""),
                    "_sheet_name": cls.PRIMARY_SHEET["sheet_name"],
                    "_source_row_no": row.get("_source_row_no", 0),
                }
            )
        return normalized_rows

    @classmethod
    def _build_header_map(cls, sheet_meta, header_values):
        normalized_headers = {
            cls._normalize_header(header_value): index
            for index, header_value in enumerate(header_values)
            if header_value
        }
        header_map = {}
        missing_fields = []
        required_fields = set(sheet_meta.get("required_fields", set()))
        for field_code in sheet_meta["fields"]:
            aliases = cls.FIELD_ALIASES.get(field_code, [field_code, cls.FIELD_LABELS.get(field_code, field_code)])
            matched_index = None
            for alias in aliases:
                alias_index = normalized_headers.get(cls._normalize_header(alias))
                if alias_index is not None:
                    matched_index = alias_index
                    break
            if matched_index is None:
                if field_code in required_fields:
                    missing_fields.append(cls.FIELD_LABELS.get(field_code, field_code))
                continue
            header_map[matched_index] = field_code

        if missing_fields:
            return {}, [
                cls._make_error(
                    sheet_name=sheet_meta["sheet_name"],
                    row_no=0,
                    field_code="template_file",
                    error_code="TEMPLATE_HEADER_MISMATCH",
                    error_message=f"{sheet_meta['sheet_name']} ?????????????{' / '.join(missing_fields)}?",
                )
            ]
        return header_map, []

    @classmethod
    def _validate_rows(cls, env, source_data):
        if source_data.get("input_mode") == cls.FORMAL_INPUT_MODE:
            return cls._validate_formal_rows(env, source_data)
        errors = []
        rows = source_data.get("rows", [])
        batches_by_no = cls._search_record_map(
            env,
            model_name="logistics.dispatch.batch",
            field_name="name",
            values={row["batch_no"] for row in rows if row.get("batch_no")},
        )
        waves_by_no = cls._search_record_map(
            env,
            model_name="logistics.dispatch.wave",
            field_name="name",
            values={row["wave_no"] for row in rows if row.get("wave_no")},
        )
        existing_waybills = cls._search_record_map(
            env,
            model_name="logistics.dispatch.waybill",
            field_name="name",
            values={row["waybill_no"] for row in rows if row.get("waybill_no")},
        )
        warehouses_by_code = cls._search_record_map(
            env,
            model_name="stock.warehouse",
            field_name="code",
            values={row["warehouse_code"] for row in rows if row.get("warehouse_code")},
        )

        waybill_rows_by_no = {}
        batch_rows_by_no = {}
        wave_rows_by_no = {}
        customer_rows_by_key = {}
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.PRIMARY_SHEET["required_fields"]))
            errors.extend(
                cls._validate_main_row(
                    row,
                    warehouses_by_code=warehouses_by_code,
                    existing_waybills=existing_waybills,
                    batches_by_no=batches_by_no,
                    waves_by_no=waves_by_no,
                )
            )
            waybill_no = row.get("waybill_no")
            if waybill_no and waybill_no in waybill_rows_by_no:
                errors.extend(
                    cls._validate_consistency_against_row(
                        row,
                        waybill_rows_by_no[waybill_no],
                        ["warehouse_code", "delivery_date", "wave_no", "batch_no", "vehicle_no", "driver_name", "driver_phone"],
                        error_code="WAYBILL_SNAPSHOT_CONFLICT",
                        label="同一运单号下的运单级字段必须一致",
                    )
                )
            elif waybill_no:
                waybill_rows_by_no[waybill_no] = row

            batch_no = row.get("batch_no")
            if batch_no and batch_no in batch_rows_by_no:
                errors.extend(
                    cls._validate_consistency_against_row(
                        row,
                        batch_rows_by_no[batch_no],
                        ["warehouse_code", "wave_no"],
                        error_code="BATCH_SNAPSHOT_CONFLICT",
                        label="同一批次号下的批次级字段必须一致",
                    )
                )
            elif batch_no:
                batch_rows_by_no[batch_no] = row

            wave_no = row.get("wave_no")
            if wave_no and wave_no in wave_rows_by_no:
                errors.extend(
                    cls._validate_consistency_against_row(
                        row,
                        wave_rows_by_no[wave_no],
                        ["warehouse_code", "delivery_date"],
                        error_code="WAVE_SNAPSHOT_CONFLICT",
                        label="同一波次号下的波次级字段必须一致",
                    )
                )
            elif wave_no:
                wave_rows_by_no[wave_no] = row

            customer_key = cls._customer_group_key(row)
            if not customer_key:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="customer_line_no",
                        error_code="CUSTOMER_LINE_KEY_MISSING",
                        error_message="门店节点必须至少提供 customer_line_no、store_no、store_name、customer_no 或 customer_name 之一。",
                    )
                )
            elif customer_key in customer_rows_by_key:
                errors.extend(
                    cls._validate_consistency_against_row(
                        row,
                        customer_rows_by_key[customer_key],
                        ["customer_no", "customer_name", "store_no", "store_name", "delivery_remark", "signoff_requirement", "customer_ref"],
                        error_code="CUSTOMER_LINE_SNAPSHOT_CONFLICT",
                        label="同一门店节点下的客户/门店快照字段必须一致",
                    )
                )
            else:
                customer_rows_by_key[customer_key] = row

        return errors

    @classmethod
    def _validate_required_fields(cls, row, required_fields):
        errors = []
        for field_code in required_fields:
            if row.get(field_code):
                continue
            errors.append(
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code="FIELD_REQUIRED",
                    error_message=f"{cls.FIELD_LABELS[field_code]}不能为空。",
                )
            )
        return errors

    @classmethod
    def _validate_main_row(cls, row, *, warehouses_by_code, existing_waybills, batches_by_no, waves_by_no):
        errors = []
        if row.get("delivery_date"):
            try:
                datetime.strptime(row["delivery_date"], "%Y-%m-%d")
            except ValueError:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="delivery_date",
                        error_code="FIELD_FORMAT_INVALID",
                        error_message="配送日期格式必须为 YYYY-MM-DD。",
                    )
                )
        if row.get("waybill_no") and row["waybill_no"] in existing_waybills:
            errors.append(
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code="waybill_no",
                    error_code="WAYBILL_NO_ALREADY_EXISTS",
                    error_message="运单号已存在，请更换运单号后重新导入。",
                )
            )
        batch_record = batches_by_no.get(row.get("batch_no"))
        if batch_record and row.get("warehouse_code") and batch_record.warehouse_id.code != row["warehouse_code"]:
            errors.append(
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code="batch_no",
                    error_code="BATCH_WAREHOUSE_CONFLICT",
                    error_message="批次号已存在，但所属仓库与导入文件不一致。",
                )
            )
        if batch_record and row.get("wave_no") and batch_record.wave_id and batch_record.wave_id.name != row["wave_no"]:
            errors.append(
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code="wave_no",
                    error_code="BATCH_WAVE_CONFLICT",
                    error_message="批次号已存在，但关联波次与导入文件不一致。",
                )
            )
        if row.get("warehouse_code") and row["warehouse_code"] not in warehouses_by_code:
            errors.append(
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code="warehouse_code",
                    error_code="WAREHOUSE_CODE_NOT_FOUND",
                    error_message="仓库编码不存在，请先维护仓库主数据。",
                )
            )
        wave_record = waves_by_no.get(row.get("wave_no"))
        if wave_record and row.get("warehouse_code") and wave_record.warehouse_id.code != row["warehouse_code"]:
            errors.append(
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code="wave_no",
                    error_code="WAVE_WAREHOUSE_CONFLICT",
                    error_message="波次号已存在，但所属仓库与导入文件不一致。",
                )
            )
        if wave_record and row.get("delivery_date"):
            wave_date = fields.Date.to_string(wave_record.dispatch_date) if wave_record.dispatch_date else ""
            if wave_date and wave_date != row["delivery_date"]:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="delivery_date",
                        error_code="WAVE_DATE_CONFLICT",
                        error_message="波次号已存在，但配送日期与导入文件不一致。",
                    )
                )
        errors.extend(cls._validate_goods_row(row))
        return errors

    @classmethod
    def _validate_goods_row(cls, row):
        errors = []
        errors.extend(cls._validate_positive_number(row, "qty", allow_zero=False))
        errors.extend(cls._validate_positive_number(row, "weight", allow_zero=True))
        errors.extend(cls._validate_positive_number(row, "volume", allow_zero=True))
        errors.extend(cls._validate_integer(row, "package_count", allow_zero=True))
        if row.get("temperature_zone") and row["temperature_zone"] not in cls.ALLOWED_TEMPERATURE_ZONES:
            errors.append(
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code="temperature_zone",
                    error_code="FIELD_ENUM_INVALID",
                    error_message="温层必须是 ambient / chilled / frozen / other 之一。",
                )
            )
        return errors

    @classmethod
    def _validate_consistency_against_row(cls, row, existing_row, field_codes, *, error_code, label):
        errors = []
        for field_code in field_codes:
            if row.get(field_code, "") == existing_row.get(field_code, ""):
                continue
            errors.append(
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code=error_code,
                    error_message=f"{label}，字段 {cls.FIELD_LABELS[field_code]} 存在冲突。",
                )
            )
        return errors

    @classmethod
    def _validate_positive_number(cls, row, field_code, *, allow_zero):
        value = row.get(field_code, "")
        if value == "":
            return []
        try:
            numeric_value = float(value)
        except ValueError:
            return [
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code="FIELD_FORMAT_INVALID",
                    error_message=f"{cls.FIELD_LABELS[field_code]}必须是数字。",
                )
            ]
        if allow_zero and numeric_value < 0:
            return [
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code="FIELD_VALUE_INVALID",
                    error_message=f"{cls.FIELD_LABELS[field_code]}不能小于 0。",
                )
            ]
        if not allow_zero and numeric_value <= 0:
            return [
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code="FIELD_VALUE_INVALID",
                    error_message=f"{cls.FIELD_LABELS[field_code]}必须大于 0。",
                )
            ]
        return []

    @classmethod
    def _validate_integer(cls, row, field_code, *, allow_zero):
        value = row.get(field_code, "")
        if value == "":
            return []
        try:
            numeric_value = int(float(value))
        except ValueError:
            return [
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code="FIELD_FORMAT_INVALID",
                    error_message=f"{cls.FIELD_LABELS[field_code]}必须是整数。",
                )
            ]
        if allow_zero and numeric_value < 0:
            return [
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code="FIELD_VALUE_INVALID",
                    error_message=f"{cls.FIELD_LABELS[field_code]}不能小于 0。",
                )
            ]
        return []

    @classmethod
    def _validate_formal_rows(cls, env, source_data):
        errors = []
        waybill_rows = source_data.get("waybill_rows", [])
        customer_line_rows = source_data.get("customer_line_rows", [])
        order_line_rows = source_data.get("order_line_rows", [])
        goods_line_rows = source_data.get("goods_line_rows", [])
        warehouses_by_code = cls._search_record_map(
            env,
            model_name="stock.warehouse",
            field_name="code",
            values={row.get("warehouse_code") for row in waybill_rows if row.get("warehouse_code")},
        )
        existing_waybills = cls._search_record_map(
            env,
            model_name="logistics.dispatch.waybill",
            field_name="name",
            values={row.get("waybill_no") for row in waybill_rows if row.get("waybill_no")},
        )
        batches_by_no = cls._search_record_map(
            env,
            model_name="logistics.dispatch.batch",
            field_name="name",
            values={row.get("batch_no") for row in waybill_rows if row.get("batch_no")},
        )
        waves_by_no = cls._search_record_map(
            env,
            model_name="logistics.dispatch.wave",
            field_name="name",
            values={row.get("wave_no") for row in waybill_rows if row.get("wave_no")},
        )
        errors.extend(
            cls._validate_formal_waybill_rows(
                waybill_rows,
                warehouses_by_code=warehouses_by_code,
                existing_waybills=existing_waybills,
                batches_by_no=batches_by_no,
                waves_by_no=waves_by_no,
            )
        )
        errors.extend(cls._validate_formal_customer_line_rows(customer_line_rows, waybill_rows))
        errors.extend(cls._validate_formal_order_line_rows(order_line_rows, customer_line_rows))
        errors.extend(cls._validate_formal_goods_line_rows(goods_line_rows, order_line_rows))
        return errors

    @classmethod
    def _validate_formal_waybill_rows(cls, rows, *, warehouses_by_code, existing_waybills, batches_by_no, waves_by_no):
        errors = []
        waybill_rows_by_no = {}
        batch_rows_by_no = {}
        wave_rows_by_no = {}
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.WAYBILL_SHEET["required_fields"]))
            errors.extend(cls._validate_formal_common_fields(row))
            waybill_no = (row.get("waybill_no") or "").strip()
            if waybill_no:
                if waybill_no in waybill_rows_by_no:
                    errors.append(
                        cls._make_error(
                            sheet_name=row["_sheet_name"],
                            row_no=row["_source_row_no"],
                            field_code="waybill_no",
                            error_code="WAYBILL_NO_DUPLICATED_IN_SHEET",
                            error_message="Waybill Sheet 中运单号不允许重复。",
                        )
                    )
                else:
                    waybill_rows_by_no[waybill_no] = row
            if row.get("warehouse_code") and row["warehouse_code"] not in warehouses_by_code:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="warehouse_code",
                        error_code="WAREHOUSE_CODE_NOT_FOUND",
                        error_message="仓库编码不存在，请先维护仓库主数据。",
                    )
                )
            if waybill_no and waybill_no in existing_waybills:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="waybill_no",
                        error_code="WAYBILL_NO_ALREADY_EXISTS",
                        error_message="运单号已存在，当前模式仅允许新建主链导入。",
                    )
                )
            batch_record = batches_by_no.get(row.get("batch_no"))
            if batch_record and row.get("warehouse_code") and batch_record.warehouse_id.code != row.get("warehouse_code"):
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="batch_no",
                        error_code="BATCH_WAREHOUSE_CONFLICT",
                        error_message="批次号已存在，但所属仓库与导入文件不一致。",
                    )
                )
            if batch_record and row.get("wave_no") and batch_record.wave_id and batch_record.wave_id.name != row.get("wave_no"):
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="wave_no",
                        error_code="BATCH_WAVE_CONFLICT",
                        error_message="批次号已存在，但关联波次与导入文件不一致。",
                    )
                )
            wave_record = waves_by_no.get(row.get("wave_no"))
            if wave_record and row.get("warehouse_code") and wave_record.warehouse_id.code != row.get("warehouse_code"):
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="wave_no",
                        error_code="WAVE_WAREHOUSE_CONFLICT",
                        error_message="波次号已存在，但所属仓库与导入文件不一致。",
                    )
                )
            if wave_record and row.get("delivery_date"):
                wave_date = fields.Date.to_string(wave_record.dispatch_date) if wave_record.dispatch_date else ""
                if wave_date and wave_date != row.get("delivery_date"):
                    errors.append(
                        cls._make_error(
                            sheet_name=row["_sheet_name"],
                            row_no=row["_source_row_no"],
                            field_code="delivery_date",
                            error_code="WAVE_DATE_CONFLICT",
                            error_message="波次号已存在，但配送日期与导入文件不一致。",
                        )
                    )
            batch_no = row.get("batch_no")
            if batch_no and batch_no in batch_rows_by_no:
                errors.extend(
                    cls._validate_consistency_against_row(
                        row,
                        batch_rows_by_no[batch_no],
                        ["warehouse_code", "wave_no"],
                        error_code="BATCH_SNAPSHOT_CONFLICT",
                        label="同一批次号下的批次级字段必须一致",
                    )
                )
            elif batch_no:
                batch_rows_by_no[batch_no] = row
            wave_no = row.get("wave_no")
            if wave_no and wave_no in wave_rows_by_no:
                errors.extend(
                    cls._validate_consistency_against_row(
                        row,
                        wave_rows_by_no[wave_no],
                        ["warehouse_code", "delivery_date"],
                        error_code="WAVE_SNAPSHOT_CONFLICT",
                        label="同一波次号下的波次级字段必须一致",
                    )
                )
            elif wave_no:
                wave_rows_by_no[wave_no] = row
        return errors

    @classmethod
    def _validate_formal_customer_line_rows(cls, rows, waybill_rows):
        errors = []
        waybill_keys = {(row.get("waybill_no") or "").strip() for row in waybill_rows if row.get("waybill_no")}
        customer_rows_by_key = {}
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.CUSTOMER_LINE_SHEET["required_fields"]))
            errors.extend(cls._validate_formal_common_fields(row))
            waybill_no = (row.get("waybill_no") or "").strip()
            customer_line_no = (row.get("customer_line_no") or "").strip()
            if waybill_no and waybill_no not in waybill_keys:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="waybill_no",
                        error_code="WAYBILL_REFERENCE_NOT_FOUND",
                        error_message="CustomerLine 关联的运单号未命中 Waybill Sheet。",
                    )
                )
            customer_key = (waybill_no, customer_line_no)
            if all(customer_key):
                if customer_key in customer_rows_by_key:
                    errors.append(
                        cls._make_error(
                            sheet_name=row["_sheet_name"],
                            row_no=row["_source_row_no"],
                            field_code="customer_line_no",
                            error_code="CUSTOMER_LINE_DUPLICATED_IN_WAYBILL",
                            error_message="同一运单下门店节点编号不允许重复。",
                        )
                    )
                else:
                    customer_rows_by_key[customer_key] = row
        return errors

    @classmethod
    def _validate_formal_order_line_rows(cls, rows, customer_line_rows):
        errors = []
        customer_keys = {
            ((row.get("waybill_no") or "").strip(), (row.get("customer_line_no") or "").strip())
            for row in customer_line_rows
            if row.get("waybill_no") and row.get("customer_line_no")
        }
        order_line_no_map = {}
        order_keys_by_parent = {}
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.ORDER_LINE_SHEET["required_fields"]))
            errors.extend(cls._validate_formal_common_fields(row))
            parent_key = ((row.get("waybill_no") or "").strip(), (row.get("customer_line_no") or "").strip())
            if all(parent_key) and parent_key not in customer_keys:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="customer_line_no",
                        error_code="CUSTOMER_LINE_REFERENCE_NOT_FOUND",
                        error_message="OrderLine 关联的门店节点未命中 CustomerLine Sheet。",
                    )
                )
            order_line_no = (row.get("order_line_no") or "").strip()
            if order_line_no:
                if order_line_no in order_line_no_map:
                    errors.append(
                        cls._make_error(
                            sheet_name=row["_sheet_name"],
                            row_no=row["_source_row_no"],
                            field_code="order_line_no",
                            error_code="ORDER_LINE_NO_DUPLICATED_IN_FILE",
                            error_message="订单行编号在同一文件内必须唯一。",
                        )
                    )
                else:
                    order_line_no_map[order_line_no] = row
            source_doc_no = (row.get("source_doc_no") or "").strip()
            sales_order_no = (row.get("sales_order_no") or "").strip()
            if not source_doc_no and not sales_order_no:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="source_doc_no",
                        error_code="IMPORT_ORDER_KEY_REQUIRED",
                        error_message="单据号 / 销售订单号 至少需要填写一个。",
                    )
                )
            parent_key_orders = order_keys_by_parent.setdefault(parent_key, {})
            for key_field in ("source_doc_no", "sales_order_no"):
                key_value = (row.get(key_field) or "").strip()
                if not key_value:
                    continue
                existing_row = parent_key_orders.get((key_field, key_value))
                if existing_row and cls._formal_order_signature(existing_row) != cls._formal_order_signature(row):
                    errors.append(
                        cls._make_error(
                            sheet_name=row["_sheet_name"],
                            row_no=row["_source_row_no"],
                            field_code=key_field,
                            error_code="ORDER_LINE_KEY_DUPLICATED_IN_PARENT",
                            error_message="同一门店节点下订单键重复且快照字段不一致。",
                        )
                    )
                else:
                    parent_key_orders[(key_field, key_value)] = row
        return errors

    @classmethod
    def _validate_formal_goods_line_rows(cls, rows, order_line_rows):
        errors = []
        order_line_nos = {(row.get("order_line_no") or "").strip() for row in order_line_rows if row.get("order_line_no")}
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.GOODS_LINE_SHEET["required_fields"]))
            errors.extend(cls._validate_formal_common_fields(row))
            order_line_no = (row.get("order_line_no") or "").strip()
            if order_line_no and order_line_no not in order_line_nos:
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="order_line_no",
                        error_code="ORDER_LINE_REFERENCE_NOT_FOUND",
                        error_message="GoodsLine 关联的订单行未命中 OrderLine Sheet。",
                    )
                )
            if not (row.get("product_name") or "").strip():
                errors.append(
                    cls._make_error(
                        sheet_name=row["_sheet_name"],
                        row_no=row["_source_row_no"],
                        field_code="product_name",
                        error_code="IMPORT_PRODUCT_NAME_REQUIRED",
                        error_message="商品名称不能为空。",
                    )
                )
        return errors

    @classmethod
    def _validate_formal_common_fields(cls, row):
        errors = []
        date_fields = ("delivery_date", "doc_date", "audited_at")
        datetime_fields = ("made_at",)
        phone_fields = ("contact_phone", "driver_phone", "registered_phone")
        boolean_fields = (
            "allow_cash_on_delivery",
            "illegal_parking_flag",
            "access_alley",
            "access_handcart",
            "access_pallet_exchange",
            "access_cooler_box_exchange",
            "delivery_week_mon",
            "delivery_week_tue",
            "delivery_week_wed",
            "delivery_week_thu",
            "delivery_week_fri",
            "delivery_week_sat",
            "delivery_week_sun",
            "above_standard_price_flag",
            "below_standard_price_flag",
        )
        positive_number_fields = {
            "doc_qty": False,
            "base_qty": True,
            "small_qty": True,
            "box_qty": True,
            "gift_qty": True,
            "exchange_qty": True,
            "unit_price": True,
            "small_unit_price": True,
            "amount": True,
            "settled_amount": True,
            "unsettled_amount": True,
            "tax_amount": True,
            "amount_ex_tax": True,
            "cost_amount": True,
            "gross_profit": True,
            "gross_profit_rate": True,
            "unit_weight": True,
            "unit_volume": True,
            "total_weight": True,
            "total_volume": True,
            "longitude": True,
            "latitude": True,
            "free_parking_minutes": True,
            "parking_fee_per_hour": True,
        }
        integer_fields = {
            "stop_seq_in_waybill": True,
            "upstairs_floor_count": True,
        }
        for field_code in date_fields:
            errors.extend(cls._validate_date_field(row, field_code))
        for field_code in datetime_fields:
            errors.extend(cls._validate_datetime_field(row, field_code))
        for field_code in phone_fields:
            errors.extend(cls._validate_phone_field(row, field_code))
        for field_code in boolean_fields:
            errors.extend(cls._validate_boolean_field(row, field_code))
        for field_code, allow_zero in positive_number_fields.items():
            errors.extend(cls._validate_positive_number(row, field_code, allow_zero=allow_zero))
        for field_code, allow_zero in integer_fields.items():
            errors.extend(cls._validate_integer(row, field_code, allow_zero=allow_zero))
        return errors

    @classmethod
    def _validate_date_field(cls, row, field_code):
        value = (row.get(field_code) or "").strip()
        if not value:
            return []
        try:
            datetime.strptime(value, "%Y-%m-%d")
        except ValueError:
            return [
                cls._make_error(
                    sheet_name=row["_sheet_name"],
                    row_no=row["_source_row_no"],
                    field_code=field_code,
                    error_code="FIELD_FORMAT_INVALID",
                    error_message=f"{cls.FIELD_LABELS.get(field_code, field_code)}格式必须为 YYYY-MM-DD。",
                )
            ]
        return []

    @classmethod
    def _validate_datetime_field(cls, row, field_code):
        value = (row.get(field_code) or "").strip()
        if not value:
            return []
        for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                datetime.strptime(value, pattern)
                return []
            except ValueError:
                continue
        return [
            cls._make_error(
                sheet_name=row["_sheet_name"],
                row_no=row["_source_row_no"],
                field_code=field_code,
                error_code="FIELD_FORMAT_INVALID",
                error_message=f"{cls.FIELD_LABELS.get(field_code, field_code)}格式不正确。",
            )
        ]

    @classmethod
    def _validate_phone_field(cls, row, field_code):
        value = (row.get(field_code) or "").strip()
        if not value:
            return []
        allowed_chars = set("0123456789+-() ")
        if all(char in allowed_chars for char in value):
            return []
        return [
            cls._make_error(
                sheet_name=row["_sheet_name"],
                row_no=row["_source_row_no"],
                field_code=field_code,
                error_code="FIELD_FORMAT_INVALID",
                error_message=f"{cls.FIELD_LABELS.get(field_code, field_code)}格式不正确。",
            )
        ]

    @classmethod
    def _validate_boolean_field(cls, row, field_code):
        value = (row.get(field_code) or "").strip()
        if not value:
            return []
        normalized = value.upper()
        if normalized in cls.BOOLEAN_TRUE_VALUES or normalized in cls.BOOLEAN_FALSE_VALUES:
            return []
        return [
            cls._make_error(
                sheet_name=row["_sheet_name"],
                row_no=row["_source_row_no"],
                field_code=field_code,
                error_code="FIELD_FORMAT_INVALID",
                error_message=f"{cls.FIELD_LABELS.get(field_code, field_code)}仅支持 是/否、Y/N、TRUE/FALSE、1/0。",
            )
        ]

    @classmethod
    def _formal_order_signature(cls, row):
        return tuple(
            (row.get(field_code) or "").strip()
            for field_code in (
                "source_doc_no",
                "sales_order_no",
                "source_ref_no",
                "third_party_doc_no",
                "doc_type",
                "business_type",
                "doc_source",
                "doc_date",
                "department_name",
                "channel_name",
                "salesperson_name",
                "order_remark",
            )
        )

    @classmethod
    def _execute_formal_task_import(cls, env, task, source_data):
        task_line_map = {line.business_key: line for line in task.task_line_ids}
        error_line_model = env["logistics.import.error.line"].sudo()
        wave_model = env["logistics.dispatch.wave"].sudo()
        batch_model = env["logistics.dispatch.batch"].sudo()
        waybill_model = env["logistics.dispatch.waybill"].sudo()
        customer_line_model = env["logistics.dispatch.waybill.customer.line"].sudo()
        order_line_model = env["logistics.dispatch.waybill.order.line"].sudo()
        goods_line_model = env["logistics.dispatch.waybill.customer.goods.line"].sudo()
        partner_model = env["res.partner"].sudo()
        customer_profile_model = env["logistics.customer.profile"].sudo()
        store_profile_model = env["logistics.store.profile"].sudo()

        waybill_rows = source_data.get("waybill_rows", [])
        customer_line_rows = source_data.get("customer_line_rows", [])
        order_line_rows = source_data.get("order_line_rows", [])
        goods_line_rows = source_data.get("goods_line_rows", [])

        warehouse_map = cls._search_record_map(
            env,
            model_name="stock.warehouse",
            field_name="code",
            values={row.get("warehouse_code") for row in waybill_rows if row.get("warehouse_code")},
        )
        partner_by_external_code = cls._search_partner_map(
            env,
            model_domain=[("is_logistics_partner", "=", True)],
            field_name="external_customer_code",
            values={row.get("external_customer_code") for row in customer_line_rows if row.get("external_customer_code")},
        )
        partner_by_name = cls._search_partner_map(
            env,
            model_domain=[("is_logistics_partner", "=", True)],
            field_name="name",
            values={row.get("customer_name") for row in customer_line_rows if row.get("customer_name")},
        )
        wave_record_map = cls._search_record_map(
            env,
            model_name="logistics.dispatch.wave",
            field_name="name",
            values={row.get("wave_no") for row in waybill_rows if row.get("wave_no")},
        )
        batch_record_map = cls._search_record_map(
            env,
            model_name="logistics.dispatch.batch",
            field_name="name",
            values={row.get("batch_no") for row in waybill_rows if row.get("batch_no")},
        )
        existing_waybill_map = cls._search_record_map(
            env,
            model_name="logistics.dispatch.waybill",
            field_name="name",
            values={row.get("waybill_no") for row in waybill_rows if row.get("waybill_no")},
        )

        created_waybill_map = {}
        created_customer_line_map = {}
        created_order_line_map = {}

        success_count = 0
        fail_count = 0
        skipped_count = 0
        created_wave_count = 0
        created_batch_count = 0
        created_waybill_count = 0
        created_customer_line_count = 0
        created_order_line_count = 0
        created_goods_line_count = 0

        for row in waybill_rows:
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                warehouse = warehouse_map.get(row.get("warehouse_code"))
                if not warehouse:
                    raise ValidationError("仓库编码不存在。")
                wave = wave_record_map.get(row.get("wave_no"))
                created_flags = []
                if not wave:
                    wave = wave_model.create(
                        {
                            "wave_no": row.get("wave_no"),
                            "dispatch_date": row.get("delivery_date") or False,
                            "warehouse_id": warehouse.id,
                            "organization_name_snapshot": row.get("organization_name") or False,
                        }
                    )
                    wave_record_map[row.get("wave_no")] = wave
                    created_wave_count += 1
                    created_flags.append("新建波次")
                batch = batch_record_map.get(row.get("batch_no"))
                if not batch:
                    batch = batch_model.create(
                        {
                            "batch_no": row.get("batch_no"),
                            "wave_id": wave.id,
                            "warehouse_id": warehouse.id,
                            "route_name_snapshot": row.get("route_name") or False,
                            "driver_name_snapshot": row.get("driver_name") or False,
                            "driver_phone_snapshot": row.get("driver_phone") or False,
                            "remark": row.get("delivery_remark") or False,
                        }
                    )
                    batch_record_map[row.get("batch_no")] = batch
                    created_batch_count += 1
                    created_flags.append("新建批次")
                else:
                    batch_updates = {}
                    if row.get("route_name") and not batch.route_name_snapshot:
                        batch_updates["route_name_snapshot"] = row.get("route_name")
                    if row.get("driver_name") and not batch.driver_name_snapshot:
                        batch_updates["driver_name_snapshot"] = row.get("driver_name")
                    if row.get("driver_phone") and not batch.driver_phone_snapshot:
                        batch_updates["driver_phone_snapshot"] = row.get("driver_phone")
                    if row.get("delivery_remark") and not batch.remark:
                        batch_updates["remark"] = row.get("delivery_remark")
                    if batch_updates:
                        batch.write(batch_updates)
                if row.get("waybill_no") in existing_waybill_map:
                    raise ValidationError("运单号已存在，当前模式仅允许新建主链导入。")
                waybill = waybill_model.create(
                    cls._build_formal_waybill_create_vals(row, warehouse=warehouse, batch=batch)
                )
                existing_waybill_map[row.get("waybill_no")] = waybill
                created_waybill_map[row.get("waybill_no")] = waybill
                created_waybill_count += 1
                created_flags.append("新建运单")
                success_count += 1
                cls._update_task_line(
                    task_line,
                    status="success",
                    message=" / ".join(created_flags) + "，本行导入成功。",
                    target_model="logistics.dispatch.waybill",
                    target_res_id=waybill.id,
                )
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"Waybill 导入失败：{exc}")
                error_line_model.create(
                    cls._make_runtime_error_line_vals(
                        task=task,
                        task_line=task_line,
                        row=row,
                        field_name="waybill_no",
                        error_message=f"Waybill 导入失败：{exc}",
                    )
                )

        for row in customer_line_rows:
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                waybill = created_waybill_map.get(row.get("waybill_no"))
                if not waybill:
                    raise ValidationError("未找到关联运单。")
                partner_record = cls._ensure_formal_partner_profile_records(
                    row,
                    partner_model=partner_model,
                    customer_profile_model=customer_profile_model,
                    store_profile_model=store_profile_model,
                    partner_by_external_code=partner_by_external_code,
                    partner_by_name=partner_by_name,
                )
                customer_line = customer_line_model.create(
                    cls._build_formal_customer_line_create_vals(
                        row,
                        waybill=waybill,
                        partner_record=partner_record,
                    )
                )
                created_customer_line_map[(row.get("waybill_no"), row.get("customer_line_no"))] = customer_line
                created_customer_line_count += 1
                success_count += 1
                cls._update_task_line(
                    task_line,
                    status="success",
                    message="新建门店节点，本行导入成功。",
                    target_model="logistics.dispatch.waybill.customer.line",
                    target_res_id=customer_line.id,
                )
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"CustomerLine 导入失败：{exc}")
                error_line_model.create(
                    cls._make_runtime_error_line_vals(
                        task=task,
                        task_line=task_line,
                        row=row,
                        field_name="customer_line_no",
                        error_message=f"CustomerLine 导入失败：{exc}",
                    )
                )

        for row in order_line_rows:
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                customer_line = created_customer_line_map.get((row.get("waybill_no"), row.get("customer_line_no")))
                if not customer_line:
                    raise ValidationError("未找到关联门店节点。")
                order_line = order_line_model.create(
                    cls._build_formal_order_line_create_vals(
                        row,
                        waybill=customer_line.waybill_id,
                        customer_line=customer_line,
                    )
                )
                created_order_line_map[row.get("order_line_no")] = order_line
                created_order_line_count += 1
                success_count += 1
                cls._update_task_line(
                    task_line,
                    status="success",
                    message="新建订单行，本行导入成功。",
                    target_model="logistics.dispatch.waybill.order.line",
                    target_res_id=order_line.id,
                )
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"OrderLine 导入失败：{exc}")
                error_line_model.create(
                    cls._make_runtime_error_line_vals(
                        task=task,
                        task_line=task_line,
                        row=row,
                        field_name="order_line_no",
                        error_message=f"OrderLine 导入失败：{exc}",
                    )
                )

        for row in goods_line_rows:
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                order_line = created_order_line_map.get(row.get("order_line_no"))
                if not order_line:
                    raise ValidationError("未找到关联订单行。")
                goods_line = goods_line_model.create(
                    cls._build_formal_goods_line_create_vals(
                        row,
                        order_line=order_line,
                    )
                )
                created_goods_line_count += 1
                success_count += 1
                cls._update_task_line(
                    task_line,
                    status="success",
                    message="新建货物行，本行导入成功。",
                    target_model="logistics.dispatch.waybill.customer.goods.line",
                    target_res_id=goods_line.id,
                )
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"GoodsLine 导入失败：{exc}")
                error_line_model.create(
                    cls._make_runtime_error_line_vals(
                        task=task,
                        task_line=task_line,
                        row=row,
                        field_name="order_line_no",
                        error_message=f"GoodsLine 导入失败：{exc}",
                    )
                )

        task_status = "success"
        if fail_count and success_count:
            task_status = "partial_failed"
        elif fail_count and not success_count:
            task_status = "failed"
        summary_message = (
            f"本次任务已完成，共 {task.total_count} 行，"
            f"成功 {success_count} 行，失败 {fail_count} 行，跳过 {skipped_count} 行。"
        )
        return {
            "task_status": task_status,
            "success_count": success_count,
            "fail_count": fail_count,
            "skipped_count": skipped_count,
            "created_wave_count": created_wave_count,
            "created_batch_count": created_batch_count,
            "created_waybill_count": created_waybill_count,
            "created_customer_line_count": created_customer_line_count,
            "created_order_line_count": created_order_line_count,
            "created_goods_line_count": created_goods_line_count,
            "summary_message": summary_message,
        }

    @classmethod
    def _build_formal_waybill_create_vals(cls, row, *, warehouse, batch):
        return {
            "waybill_no": row.get("waybill_no"),
            "delivery_date": row.get("delivery_date") or False,
            "batch_id": batch.id,
            "warehouse_id": warehouse.id,
            "organization_name_snapshot": row.get("organization_name") or False,
            "route_name_snapshot": row.get("route_name") or False,
            "delivery_remark_snapshot": row.get("delivery_remark") or False,
            "remark": row.get("delivery_remark") or False,
        }

    @classmethod
    def _build_formal_customer_line_create_vals(cls, row, *, waybill, partner_record=False):
        vals = {
            "waybill_id": waybill.id,
            "customer_line_no": row.get("customer_line_no") or False,
            "internal_customer_code_snapshot": row.get("customer_seq_no") or False,
            "external_customer_code_snapshot": row.get("external_customer_code") or False,
            "customer_name_snapshot": row.get("customer_name") or False,
            "contact_name_snapshot": row.get("contact_name") or False,
            "contact_phone_snapshot": row.get("contact_phone") or row.get("registered_phone") or False,
            "address_full_snapshot": row.get("address_full") or row.get("registered_address") or False,
            "longitude_snapshot": cls._safe_float(row.get("longitude"), default=0.0) if row.get("longitude") else False,
            "latitude_snapshot": cls._safe_float(row.get("latitude"), default=0.0) if row.get("latitude") else False,
            "address_region_json_snapshot": cls._build_address_region_snapshot(row) or False,
            "stop_seq_in_waybill": cls._safe_int(row.get("stop_seq_in_waybill"), default=0),
            "delivery_access_flags_snapshot": cls._build_delivery_access_flags_snapshot(row) or False,
            "upstairs_floor_count_snapshot": cls._safe_int(row.get("upstairs_floor_count"), default=0),
            "basement_height_limit_text_snapshot": row.get("basement_height_limit_text") or False,
            "delivery_note": cls._build_delivery_window_text(row) or row.get("receive_time_slots_text") or False,
        }
        if partner_record:
            vals["partner_id"] = partner_record.id
            vals["customer_id"] = partner_record.id
        return vals

    @classmethod
    def _build_formal_order_line_create_vals(cls, row, *, waybill, customer_line):
        source_doc_no = (row.get("source_doc_no") or row.get("sales_order_no") or "").strip() or False
        return {
            "waybill_id": waybill.id,
            "customer_line_id": customer_line.id,
            "source_doc_no": source_doc_no,
            "sales_order_no": row.get("sales_order_no") or False,
            "source_ref_no": row.get("source_ref_no") or False,
            "third_party_doc_no": row.get("third_party_doc_no") or False,
            "doc_type": row.get("doc_type") or False,
            "business_type": row.get("business_type") or False,
            "doc_source": row.get("doc_source") or False,
            "doc_date": cls._normalize_date_value(row.get("doc_date")),
            "audited_at": cls._normalize_date_value(row.get("audited_at")),
            "department_name_snapshot": row.get("department_name") or False,
            "channel_name_snapshot": row.get("channel_name") or False,
            "salesperson_name_snapshot": row.get("salesperson_name") or False,
            "payment_status": cls._normalize_selection_value(row.get("payment_status"), {"unpaid", "partial_paid", "paid"}),
            "settlement_status": cls._normalize_selection_value(row.get("settlement_status"), {"unsettled", "partial_settled", "settled"}),
            "doc_status": cls._normalize_selection_value(row.get("doc_status"), {"draft", "confirmed", "cancelled"}),
            "logistics_status": cls._normalize_selection_value(row.get("logistics_status"), {"pending_outbound", "outbounded", "delivering", "signed", "abnormal"}),
            "maker_name": row.get("maker_name") or False,
            "auditor_name": row.get("auditor_name") or False,
            "made_at": cls._normalize_datetime_value(row.get("made_at")),
            "order_remark": row.get("order_remark") or False,
            "custom_field_1": row.get("custom_field_1") or False,
        }

    @classmethod
    def _build_formal_goods_line_create_vals(cls, row, *, order_line):
        doc_qty = cls._safe_float(row.get("doc_qty"), default=0.0)
        box_qty = cls._safe_int(row.get("box_qty"), default=0)
        total_weight = cls._safe_float(row.get("total_weight"), default=0.0)
        total_volume = cls._safe_float(row.get("total_volume"), default=0.0)
        return {
            "customer_line_id": order_line.customer_line_id.id,
            "order_line_id": order_line.id,
            "external_product_code_snapshot": row.get("external_product_code") or False,
            "product_name_snapshot": row.get("product_name") or False,
            "spec_snapshot": row.get("spec") or False,
            "barcode_snapshot": row.get("barcode") or False,
            "brand_name_snapshot": row.get("brand_name") or False,
            "category_name_snapshot": row.get("category_name") or False,
            "base_unit_name": row.get("base_unit_name") or False,
            "doc_unit_name": row.get("doc_unit_name") or False,
            "small_unit_name": row.get("small_unit_name") or False,
            "base_qty": cls._safe_float(row.get("base_qty"), default=0.0),
            "doc_qty": doc_qty,
            "small_qty": cls._safe_float(row.get("small_qty"), default=0.0),
            "box_qty": box_qty,
            "gift_qty": cls._safe_float(row.get("gift_qty"), default=0.0),
            "exchange_qty": cls._safe_float(row.get("exchange_qty"), default=0.0),
            "unit_price": cls._safe_float(row.get("unit_price"), default=0.0),
            "small_unit_price": cls._safe_float(row.get("small_unit_price"), default=0.0),
            "amount": cls._safe_float(row.get("amount"), default=0.0),
            "settled_amount": cls._safe_float(row.get("settled_amount"), default=0.0),
            "unsettled_amount": cls._safe_float(row.get("unsettled_amount"), default=0.0),
            "tax_amount": cls._safe_float(row.get("tax_amount"), default=0.0),
            "amount_ex_tax": cls._safe_float(row.get("amount_ex_tax"), default=0.0),
            "cost_amount": cls._safe_float(row.get("cost_amount"), default=0.0),
            "gross_profit": cls._safe_float(row.get("gross_profit"), default=0.0),
            "gross_profit_rate": cls._safe_float(row.get("gross_profit_rate"), default=0.0),
            "above_standard_price_flag": bool(cls._normalize_boolean_value(row.get("above_standard_price_flag"))),
            "below_standard_price_flag": bool(cls._normalize_boolean_value(row.get("below_standard_price_flag"))),
            "unit_weight": cls._safe_float(row.get("unit_weight"), default=0.0),
            "unit_volume": cls._safe_float(row.get("unit_volume"), default=0.0),
            "total_weight": total_weight,
            "total_volume": total_volume,
            "line_remark": row.get("line_remark") or False,
            "goods_code": row.get("external_product_code") or False,
            "goods_name": row.get("product_name") or False,
            "specification": row.get("spec") or False,
            "quantity": doc_qty,
            "package_count": box_qty,
            "uom_name": row.get("doc_unit_name") or False,
            "weight": total_weight,
            "volume": total_volume,
            "remark": row.get("line_remark") or False,
        }

    @classmethod
    def _ensure_formal_partner_profile_records(
        cls,
        row,
        *,
        partner_model,
        customer_profile_model,
        store_profile_model,
        partner_by_external_code,
        partner_by_name,
    ):
        partner = cls._find_formal_partner_record(
            row,
            partner_by_external_code=partner_by_external_code,
            partner_by_name=partner_by_name,
        )
        partner_vals = cls._build_formal_partner_vals(row)
        if partner:
            if partner_vals:
                partner.write(partner_vals)
        else:
            partner = partner_model.create(partner_vals)
        if partner.external_customer_code:
            partner_by_external_code[partner.external_customer_code] = partner
        if partner.name:
            partner_by_name[partner.name] = partner
        cls._ensure_profile_record(customer_profile_model, partner)
        cls._ensure_profile_record(store_profile_model, partner)
        return partner

    @classmethod
    def _find_formal_partner_record(cls, row, *, partner_by_external_code, partner_by_name):
        external_customer_code = (row.get("external_customer_code") or "").strip()
        if external_customer_code and external_customer_code in partner_by_external_code:
            return partner_by_external_code[external_customer_code]
        customer_name = (row.get("customer_name") or "").strip()
        if customer_name and not external_customer_code:
            return partner_by_name.get(customer_name)
        return False

    @classmethod
    def _build_formal_partner_vals(cls, row):
        partner_name = (
            (row.get("customer_name") or "").strip()
            or (row.get("external_customer_code") or "").strip()
            or (row.get("customer_line_no") or "").strip()
        )
        vals = {
            "name": partner_name,
            "is_logistics_partner": True,
        }
        external_customer_code = (row.get("external_customer_code") or "").strip()
        customer_seq_no = (row.get("customer_seq_no") or "").strip()
        if external_customer_code:
            vals["external_customer_code"] = external_customer_code
        if customer_seq_no:
            vals["logistics_customer_code"] = customer_seq_no
        elif external_customer_code:
            vals["logistics_customer_code"] = external_customer_code
        if external_customer_code:
            vals["logistics_store_code"] = external_customer_code
        elif customer_seq_no:
            vals["logistics_store_code"] = customer_seq_no

        for source_field, target_field in (
            ("contact_name", "contact_name"),
            ("organization_name", "organization_name"),
            ("department_name", "department_name"),
            ("salesperson_name", "salesperson_name"),
            ("channel_name", "channel_name"),
            ("invoice_type", "invoice_type"),
            ("route_preference", "route_preference"),
            ("warehouse_preference", "warehouse_preference"),
            ("receive_start_time", "receive_start_time"),
            ("receive_end_time", "receive_end_time"),
            ("receive_time_slots_text", "receive_time_slots_text"),
            ("no_receive_time_slots_text", "no_receive_time_slots_text"),
            ("parking_location_text", "parking_location_text"),
            ("parking_mode_text", "parking_mode_text"),
            ("unload_entrance_text", "unload_entrance_text"),
            ("unload_location_text", "unload_location_text"),
            ("basement_height_limit_text", "basement_height_limit_text"),
        ):
            value = (row.get(source_field) or "").strip()
            if value:
                vals[target_field] = value

        contact_phone = ((row.get("contact_phone") or "").strip() or (row.get("registered_phone") or "").strip())
        if contact_phone:
            vals["contact_phone"] = contact_phone
        address_full = ((row.get("address_full") or "").strip() or (row.get("registered_address") or "").strip())
        if address_full:
            vals["address_full"] = address_full
        if row.get("longitude"):
            vals["partner_longitude"] = cls._safe_float(row.get("longitude"), default=0.0)
        if row.get("latitude"):
            vals["partner_latitude"] = cls._safe_float(row.get("latitude"), default=0.0)

        address_region_json = cls._build_address_region_snapshot(row)
        if address_region_json:
            vals["address_region_json"] = address_region_json

        customer_level = cls._normalize_customer_level_value(row.get("customer_level"))
        if customer_level:
            vals["logistics_customer_level"] = customer_level
        customer_status = cls._normalize_customer_status_value(row.get("customer_status"))
        if customer_status:
            vals["customer_status"] = customer_status

        for source_field, target_field in (
            ("allow_cash_on_delivery", "allow_cash_on_delivery"),
            ("internal_counterparty_flag", "internal_counterparty_flag"),
            ("illegal_parking_flag", "illegal_parking_flag"),
        ):
            raw_value = row.get(source_field)
            if raw_value not in (None, ""):
                vals[target_field] = bool(cls._normalize_boolean_value(raw_value))

        for source_field, target_field in (
            ("free_parking_minutes", "free_parking_minutes"),
            ("upstairs_floor_count", "upstairs_floor_count"),
        ):
            raw_value = row.get(source_field)
            if raw_value not in (None, ""):
                vals[target_field] = cls._safe_int(raw_value, default=0)

        if row.get("parking_fee_per_hour") not in (None, ""):
            vals["parking_fee_per_hour"] = cls._safe_float(row.get("parking_fee_per_hour"), default=0.0)

        delivery_week_flags = cls._build_delivery_week_flags(row)
        if delivery_week_flags:
            vals["delivery_week_flags"] = delivery_week_flags
        delivery_access_flags = cls._build_delivery_access_flags_snapshot(row)
        if delivery_access_flags:
            vals["delivery_access_flags"] = delivery_access_flags

        delivery_window_text = cls._build_delivery_window_text(row)
        if delivery_window_text:
            vals["delivery_window_text"] = delivery_window_text
        return vals

    @classmethod
    def _ensure_profile_record(cls, profile_model, partner):
        profile = profile_model.search([("partner_id", "=", partner.id)], limit=1)
        if profile:
            return profile
        return profile_model.create({"partner_id": partner.id})

    @classmethod
    def _build_delivery_window_text(cls, row):
        receive_time_slots_text = (row.get("receive_time_slots_text") or "").strip()
        if receive_time_slots_text:
            return receive_time_slots_text
        receive_start_time = (row.get("receive_start_time") or "").strip()
        receive_end_time = (row.get("receive_end_time") or "").strip()
        if receive_start_time and receive_end_time:
            return f"{receive_start_time}-{receive_end_time}"
        return False

    @classmethod
    def _build_delivery_week_flags(cls, row):
        day_mapping = {
            "delivery_week_mon": "mon",
            "delivery_week_tue": "tue",
            "delivery_week_wed": "wed",
            "delivery_week_thu": "thu",
            "delivery_week_fri": "fri",
            "delivery_week_sat": "sat",
            "delivery_week_sun": "sun",
        }
        flags = [flag for field_name, flag in day_mapping.items() if cls._normalize_boolean_value(row.get(field_name))]
        return ",".join(flags)

    @classmethod
    def _normalize_customer_level_value(cls, value):
        normalized = (value or "").strip().lower()
        mapping = {
            "standard": "standard",
            "标准": "standard",
            "vip": "vip",
            "战略": "strategic",
            "strategic": "strategic",
        }
        return mapping.get(normalized, False)

    @classmethod
    def _normalize_customer_status_value(cls, value):
        normalized = (value or "").strip().lower()
        mapping = {
            "normal": "normal",
            "正常": "normal",
            "active": "normal",
            "enabled": "normal",
            "disabled": "disabled",
            "停用": "disabled",
            "inactive": "disabled",
            "frozen": "frozen",
            "冻结": "frozen",
            "paused": "frozen",
        }
        return mapping.get(normalized, False)

    @classmethod
    def _normalize_date_value(cls, value):
        value = (value or "").strip()
        if not value:
            return False
        return value

    @classmethod
    def _normalize_datetime_value(cls, value):
        value = (value or "").strip()
        if not value:
            return False
        for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(value, pattern)
                return parsed.strftime("%Y-%m-%d %H:%M:%S")
            except ValueError:
                continue
        return False

    @classmethod
    def _normalize_selection_value(cls, value, allowed_values):
        normalized = (value or "").strip()
        if not normalized or normalized not in allowed_values:
            return False
        return normalized

    @classmethod
    def _normalize_boolean_value(cls, value):
        normalized = (value or "").strip().upper()
        if not normalized:
            return False
        if normalized in cls.BOOLEAN_TRUE_VALUES:
            return True
        if normalized in cls.BOOLEAN_FALSE_VALUES:
            return False
        return False

    @classmethod
    def _build_address_region_snapshot(cls, row):
        region = {
            "province_name": row.get("province_name") or False,
            "city_name": row.get("city_name") or False,
            "district_name": row.get("district_name") or False,
        }
        if any(region.values()):
            return region
        return False

    @classmethod
    def _build_delivery_access_flags_snapshot(cls, row):
        flags = []
        mapping = {
            "access_alley": "access_alley",
            "access_handcart": "access_handcart",
            "access_pallet_exchange": "access_pallet_exchange",
            "access_cooler_box_exchange": "access_cooler_box_exchange",
        }
        for field_code, flag_name in mapping.items():
            if cls._normalize_boolean_value(row.get(field_code)):
                flags.append(flag_name)
        return ",".join(flags)

    @classmethod
    def _search_partner_map(cls, env, *, model_domain, field_name, values):
        if not values:
            return {}
        records = env["res.partner"].sudo().search(model_domain + [(field_name, "in", list(values))])
        return {getattr(record, field_name): record for record in records if getattr(record, field_name)}

    @classmethod
    def _search_record_map(cls, env, *, model_name, field_name, values):
        if not values:
            return {}
        records = env[model_name].sudo().search([(field_name, "in", list(values))])
        return {getattr(record, field_name): record for record in records if getattr(record, field_name)}

    @classmethod
    def _make_error(cls, *, sheet_name, row_no, field_code, error_code, error_message):
        return {
            "sheet_name": sheet_name,
            "row_no": row_no,
            "field_code": field_code,
            "field_label": cls.FIELD_LABELS.get(field_code, field_code),
            "error_code": error_code,
            "error_message": error_message,
        }

    @classmethod
    def _build_precheck_result(cls, source_data, errors):
        failed_row_keys = {
            (error.get("sheet_name"), error.get("row_no"))
            for error in errors
            if error.get("row_no")
        }
        total_row_count = cls._count_total_rows(source_data)
        failed_row_count = len(failed_row_keys)
        if errors and failed_row_count == 0:
            failed_row_count = 1
        return {
            "template_code": cls.TEMPLATE_CODE,
            "template_version": cls.TEMPLATE_VERSION,
            "task_no": False,
            "import_batch_no": False,
            "precheck_token": False,
            "error_report_url": False,
            "total_row_count": total_row_count,
            "passed_row_count": max(total_row_count - failed_row_count, 0),
            "failed_row_count": failed_row_count,
            "can_confirm_import": total_row_count > 0 and not errors,
            "errors": sorted(
                errors,
                key=lambda item: (
                    item.get("sheet_name", ""),
                    item.get("row_no", 0),
                    item.get("field_code", ""),
                    item.get("error_code", ""),
                ),
            ),
        }

    @classmethod
    def _count_total_rows(cls, source_data):
        return sum(1 for _ in cls._iter_source_rows(source_data))

    @classmethod
    def _get_batch_by_precheck_token(cls, env, precheck_token):
        if not precheck_token:
            raise ValidationError("请提供预校验令牌。")
        batch = env["logistics.import.batch"].sudo().search([("precheck_token", "=", precheck_token)], limit=1)
        if not batch:
            raise ValidationError("未找到对应的预校验批次。")
        return batch

    @classmethod
    def _get_batch_for_error_report(cls, env, *, task_no="", import_batch_no="", precheck_token=""):
        batch_model = env["logistics.import.batch"].sudo()
        if task_no:
            batch = batch_model.search([("name", "=", task_no)], limit=1)
            if batch:
                return batch
        if import_batch_no:
            batch = batch_model.search([("name", "=", import_batch_no)], limit=1)
            if batch:
                return batch
        if precheck_token:
            batch = batch_model.search([("precheck_token", "=", precheck_token)], limit=1)
            if batch:
                return batch
        raise ValidationError("未找到可导出错误报告的导入批次。")

    @classmethod
    def _apply_import(cls, env, source_data, stats):
        waybill_model = env["logistics.dispatch.waybill"].sudo()
        customer_line_model = env["logistics.dispatch.waybill.customer.line"].sudo()
        goods_line_model = env["logistics.dispatch.waybill.customer.goods.line"].sudo()

        waybill_rows = source_data["waybill_rows"]
        customer_rows = source_data["customer_rows"]
        goods_rows = source_data["goods_rows"]
        warehouse_map = cls._search_record_map(
            env,
            model_name="stock.warehouse",
            field_name="code",
            values={row["warehouse_code"] for row in waybill_rows if row.get("warehouse_code")},
        )

        existing_waybills = waybill_model.search([("name", "in", [row["waybill_no"] for row in waybill_rows])])
        existing_waybill_map = {record.name: record for record in existing_waybills}
        waybill_record_map = {}
        for row in waybill_rows:
            waybill_no = row["waybill_no"]
            if waybill_no in existing_waybill_map:
                stats["skipped_record_count"] += 1
                waybill_record_map[waybill_no] = existing_waybill_map[waybill_no]
                continue
            vals = {
                "waybill_no": waybill_no,
                "delivery_date": row.get("delivery_date") or False,
                "remark": row.get("remark") or False,
            }
            if row.get("batch_no"):
                vals["batch_no"] = row["batch_no"]
            if row.get("warehouse_code") and row["warehouse_code"] in warehouse_map:
                vals["warehouse_id"] = warehouse_map[row["warehouse_code"]].id
            waybill = waybill_model.create(vals)
            waybill_record_map[waybill_no] = waybill
            stats["created_waybill_count"] += 1

        customer_line_record_map = {}
        for row in customer_rows:
            waybill = waybill_record_map.get(row["waybill_no"])
            if not waybill or waybill.name in existing_waybill_map:
                stats["skipped_record_count"] += 1
                continue
            customer_line = customer_line_model.create(
                {
                    "waybill_id": waybill.id,
                    "customer_no": row["customer_no"],
                    "customer_name": row.get("customer_name") or False,
                    "store_no": row.get("store_no") or False,
                    "store_name": row.get("store_name") or False,
                    "delivery_note": row.get("delivery_remark") or False,
                    "signoff_requirement": row.get("signoff_requirement") or False,
                    "customer_ref": row.get("customer_ref") or False,
                }
            )
            customer_line_record_map[cls._customer_key(row)] = customer_line
            stats["created_customer_line_count"] += 1

        for row in goods_rows:
            customer_line = customer_line_record_map.get(cls._customer_key(row))
            if not customer_line:
                stats["skipped_record_count"] += 1
                continue
            goods_line_model.create(
                {
                    "customer_line_id": customer_line.id,
                    "goods_code": row.get("goods_code") or False,
                    "goods_name": row["goods_name"],
                    "specification": row.get("spec") or False,
                    "quantity": cls._safe_float(row.get("qty"), default=0.0),
                    "package_count": cls._safe_int(row.get("package_count"), default=0),
                    "uom_name": row.get("uom_name") or False,
                    "weight": cls._safe_float(row.get("weight"), default=0.0),
                    "volume": cls._safe_float(row.get("volume"), default=0.0),
                    "temperature_zone": row.get("temperature_zone") or "ambient",
                    "package_type": row.get("package_type") or False,
                    "remark": row.get("remark") or False,
                }
            )
            stats["created_goods_line_count"] += 1

    @classmethod
    def _build_result_payload(cls, batch):
        batch.ensure_one()
        file_ext = cls._guess_file_ext(batch.file_name)
        return {
            "task_no": batch.name,
            "import_batch_no": batch.name,
            "object_type": cls.OBJECT_TYPE,
            "object_type_label": cls.OBJECT_TYPE_LABEL,
            "status": batch.state,
            "status_label": cls.STATUS_LABELS.get(batch.state, batch.state),
            "total_count": batch.total_row_count,
            "success_count": batch.passed_row_count,
            "fail_count": batch.failed_row_count,
            "summary_message": cls._build_summary_message(batch),
            "template_code": batch.template_code,
            "template_version": batch.template_version,
            "file_name": batch.file_name,
            "total_row_count": batch.total_row_count,
            "passed_row_count": batch.passed_row_count,
            "failed_row_count": batch.failed_row_count,
            "confirmed_at": fields.Datetime.to_string(batch.confirmed_at) if batch.confirmed_at else False,
            "started_at": fields.Datetime.to_string(batch.confirmed_at) if batch.confirmed_at else False,
            "finished_at": fields.Datetime.to_string(batch.finished_at) if batch.finished_at else False,
            "created_waybill_count": batch.created_waybill_count,
            "created_customer_line_count": batch.created_customer_line_count,
            "created_goods_line_count": batch.created_goods_line_count,
            "updated_record_count": batch.updated_record_count,
            "skipped_record_count": batch.skipped_record_count,
            "failed_record_count": batch.failed_record_count,
            "business_summary": {
                "written_waybill_count": batch.created_waybill_count,
                "written_customer_line_count": batch.created_customer_line_count,
                "written_goods_line_count": batch.created_goods_line_count,
            },
            "source_file": {
                "source_file_no": False,
                "file_name": batch.file_name,
                "file_ext": file_ext,
                "uploaded_at": False,
            },
            "operator": False,
            "error_report": {
                "download_ready": bool(batch.error_report_url),
                "download_url": batch.error_report_url or False,
            },
            "next_actions": {
                "refresh": True,
                "go_import_center": True,
                "go_waybill_list": True,
                "review_by_task": True,
            },
            "error_report_url": batch.error_report_url,
            "failure_reason": batch.failure_reason,
        }

    @classmethod
    def _build_error_report_url(cls, batch):
        batch.ensure_one()
        return f"/api/admin/logistics/imports/tasks/{batch.name}/error-report"

    @classmethod
    def _build_template_download_url(cls, *, template_locale):
        return (
            "/api/admin/logistics/imports/waybill-standard/template/download"
            f"?template_code={cls.TEMPLATE_CODE}"
            f"&template_version={cls.TEMPLATE_VERSION}"
            f"&template_locale={template_locale}"
        )

    @classmethod
    def _empty_source_data(cls):
        return {"rows": [], "input_mode": "single_sheet"}

    @classmethod
    def _cell_text(cls, value):
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")
        return str(value).strip()

    @classmethod
    def _guess_file_ext(cls, file_name):
        if not file_name or "." not in file_name:
            return False
        return file_name.rsplit(".", 1)[-1].lower()

    @classmethod
    def _build_summary_message(cls, batch):
        batch.ensure_one()
        if batch.state == "finished":
            return (
                f"本次任务已完成，共 {batch.total_row_count} 行，"
                f"成功 {batch.passed_row_count} 行，失败 {batch.failed_row_count} 行。"
            )
        if batch.state == "failed":
            return batch.failure_reason or "本次任务执行失败。"
        if batch.state == "importing":
            return "任务执行中，请稍后刷新结果。"
        if batch.state == "prechecked":
            return "任务已完成预校验，等待正式导入。"
        if batch.state == "expired":
            return "任务已失效，请重新发起导入。"
        return "请继续核对当前导入任务状态。"

    @classmethod
    def _get_batch_by_task_ref(cls, env, *, task_no="", import_batch_no=""):
        batch_ref = (task_no or import_batch_no or "").strip()
        if not batch_ref:
            raise ValidationError("请提供导入任务号。")
        batch = env["logistics.import.batch"].sudo().search([("name", "=", batch_ref)], limit=1)
        if not batch:
            raise ValidationError("未找到对应的导入任务。")
        return batch

    @classmethod
    def _get_task_by_task_no(cls, env, task_no):
        task_ref = (task_no or "").strip()
        if not task_ref:
            return False
        return env["logistics.import.task"].sudo().search([("task_no", "=", task_ref)], limit=1)

    @classmethod
    def _build_task_error_report_url(cls, task):
        task.ensure_one()
        return f"/api/admin/logistics/imports/tasks/{task.task_no}/error-report"

    @classmethod
    def _write_source_file_bytes(cls, *, source_file_no, file_name, raw_bytes):
        suffix = f".{cls._guess_file_ext(file_name) or 'xlsx'}"
        storage_dir = Path(tempfile.gettempdir()) / "odoo_logistics_imports"
        storage_dir.mkdir(parents=True, exist_ok=True)
        storage_path = storage_dir / f"{source_file_no}{suffix}"
        storage_path.write_bytes(raw_bytes or b"")
        return str(storage_path)

    @classmethod
    def _load_source_file_bytes(cls, source_file):
        source_file.ensure_one()
        storage_path = (source_file.storage_path or "").strip()
        if not storage_path:
            raise ValidationError("当前任务未找到源文件存储路径。")
        file_path = Path(storage_path)
        if not file_path.exists():
            raise ValidationError("当前任务源文件不存在，请重新上传。")
        return file_path.read_bytes()

    @classmethod
    def _iter_source_rows(cls, source_data):
        if source_data.get("input_mode") == cls.FORMAL_INPUT_MODE:
            for sheet_meta in cls.SHEETS:
                for row in source_data.get(sheet_meta["key"], []):
                    yield row
            return
        for row in source_data.get("rows", []):
            yield row

    @classmethod
    def _group_errors_by_row(cls, errors):
        grouped = {}
        for error in errors:
            row_key = (error.get("sheet_name"), error.get("row_no"))
            grouped.setdefault(row_key, []).append(error)
        return grouped

    @classmethod
    def _build_row_business_key(cls, row):
        sheet_name = row.get("_sheet_name") or "????"
        row_no = row.get("_source_row_no") or 0
        if sheet_name == cls.WAYBILL_SHEET["sheet_name"]:
            detail_parts = [row.get("waybill_no")]
        elif sheet_name == cls.CUSTOMER_LINE_SHEET["sheet_name"]:
            detail_parts = [row.get("waybill_no"), row.get("customer_line_no"), row.get("customer_name")]
        elif sheet_name == cls.ORDER_LINE_SHEET["sheet_name"]:
            detail_parts = [row.get("order_line_no"), row.get("source_doc_no") or row.get("sales_order_no")]
        elif sheet_name == cls.GOODS_LINE_SHEET["sheet_name"]:
            detail_parts = [row.get("order_line_no"), row.get("product_name")]
        else:
            detail_parts = [
                row.get("waybill_no"),
                row.get("customer_line_no") or row.get("store_no") or row.get("customer_no"),
                row.get("goods_name"),
            ]
        detail = " / ".join(filter(None, detail_parts)) or "--"
        business_key = f"{sheet_name}@{row_no} {detail}"
        return business_key[:128]

    @classmethod
    def _update_task_line(cls, task_line, *, status, message, target_model=False, target_res_id=False):
        if not task_line:
            return
        task_line.sudo().write(
            {
                "status": status,
                "message": message,
                "target_model": target_model or False,
                "target_res_id": target_res_id or False,
            }
        )

    @classmethod
    def _make_runtime_error_line_vals(cls, *, task, task_line, row, field_name, error_message):
        raw_value = row.get(field_name) if row else False
        return {
            "task_id": task.id,
            "task_line_id": task_line.id if task_line else False,
            "source_row_no": row.get("_source_row_no", 0) if row else 0,
            "field_name": field_name,
            "raw_value": raw_value or False,
            "mapped_value": raw_value or False,
            "error_code": "IMPORT_TARGET_WRITE_FAILED",
            "error_message": error_message,
        }

    @classmethod
    def _build_source_file_payload(cls, source_file):
        source_file.ensure_one()
        return {
            "source_file_no": source_file.source_file_no,
            "file_name": source_file.file_name,
            "file_ext": source_file.file_ext or cls._guess_file_ext(source_file.file_name),
            "uploaded_at": fields.Datetime.to_string(source_file.uploaded_at) if source_file.uploaded_at else False,
        }

    @classmethod
    def _build_operator_payload(cls, operator):
        if not operator:
            return False
        return {
            "id": operator.id,
            "name": operator.name,
        }

    @classmethod
    def _get_task_line_status_counts(cls, task):
        counts = {"success": 0, "failed": 0, "skipped": 0, "pending": 0}
        for line in task.task_line_ids:
            counts[line.status] = counts.get(line.status, 0) + 1
        return counts

    @classmethod
    def _get_task_target_counts(cls, task):
        counts = {"waybill": 0, "customer_line": 0, "order_line": 0, "goods_line": 0}
        for line in task.task_line_ids:
            if line.status != "success":
                continue
            if line.target_model == "logistics.dispatch.waybill":
                counts["waybill"] += 1
            elif line.target_model == "logistics.dispatch.waybill.customer.line":
                counts["customer_line"] += 1
            elif line.target_model == "logistics.dispatch.waybill.order.line":
                counts["order_line"] += 1
            elif line.target_model == "logistics.dispatch.waybill.customer.goods.line":
                counts["goods_line"] += 1
        return counts

    @classmethod
    def _get_task_business_counts(cls, task):
        counts = {"wave": 0, "batch": 0, "waybill": 0, "customer_line": 0, "order_line": 0, "goods_line": 0}
        source_file = task.source_file_id
        if not source_file:
            return counts
        try:
            source_bytes = cls._load_source_file_bytes(source_file)
            source_data, parse_errors = cls._parse_workbook(source_bytes)
        except Exception:
            return counts
        if parse_errors:
            return counts
        task_line_status = {line.business_key: line.status for line in task.task_line_ids}
        wave_keys = set()
        batch_keys = set()
        waybill_keys = set()
        customer_keys = set()
        order_keys = set()
        goods_count = 0
        for row in cls._iter_source_rows(source_data):
            if task_line_status.get(cls._build_row_business_key(row)) != "success":
                continue
            sheet_name = row.get("_sheet_name")
            if sheet_name == cls.WAYBILL_SHEET["sheet_name"]:
                if row.get("wave_no"):
                    wave_keys.add(row["wave_no"])
                if row.get("batch_no"):
                    batch_keys.add(row["batch_no"])
                if row.get("waybill_no"):
                    waybill_keys.add(row["waybill_no"])
            elif sheet_name == cls.CUSTOMER_LINE_SHEET["sheet_name"]:
                customer_key = ((row.get("waybill_no") or "").strip(), (row.get("customer_line_no") or "").strip())
                if all(customer_key):
                    customer_keys.add(customer_key)
            elif sheet_name == cls.ORDER_LINE_SHEET["sheet_name"]:
                if row.get("order_line_no"):
                    order_keys.add(row["order_line_no"])
            elif sheet_name == cls.GOODS_LINE_SHEET["sheet_name"]:
                goods_count += 1
            else:
                if row.get("wave_no"):
                    wave_keys.add(row["wave_no"])
                if row.get("batch_no"):
                    batch_keys.add(row["batch_no"])
                if row.get("waybill_no"):
                    waybill_keys.add(row["waybill_no"])
                customer_key = cls._customer_group_key(row)
                if customer_key:
                    customer_keys.add(customer_key)
                goods_count += 1
        counts["wave"] = len(wave_keys)
        counts["batch"] = len(batch_keys)
        counts["waybill"] = len(waybill_keys)
        counts["customer_line"] = len(customer_keys)
        counts["order_line"] = len(order_keys)
        counts["goods_line"] = goods_count
        return counts

    @classmethod
    def _build_task_summary_message(cls, task, counts):
        if task.status == "pending":
            return (
                f"预校验完成，共 {task.total_count} 行，"
                f"通过 {task.success_count} 行，失败 {task.fail_count} 行。"
            )
        if task.status == "running":
            return "任务执行中，请稍后刷新结果。"
        return (
            f"本次任务已完成，共 {task.total_count} 行，"
            f"成功 {task.success_count} 行，失败 {task.fail_count} 行，跳过 {counts.get('skipped', 0)} 行。"
        )

    @classmethod
    def _build_task_failure_reason(cls, task):
        if task.status == "failed" and task.error_line_ids:
            return task.error_line_ids.sorted(key=lambda rec: rec.id)[0].error_message
        if task.status == "partial_failed" and task.error_line_ids:
            return task.error_line_ids.sorted(key=lambda rec: rec.id)[0].error_message
        return False

    @classmethod
    def _get_selection_label(cls, selection, value):
        selection_map = dict(selection or [])
        return selection_map.get(value, value)

    @classmethod
    def _normalize_header(cls, header_name):
        return (header_name or "").strip().lower().replace("-", "_")

    @classmethod
    def _customer_group_key(cls, row):
        waybill_no = (row.get("waybill_no") or "").strip()
        identity = (
            (row.get("customer_line_no") or "").strip()
            or (row.get("store_no") or "").strip()
            or (row.get("store_name") or "").strip()
            or (row.get("customer_no") or "").strip()
            or (row.get("customer_name") or "").strip()
        )
        if not waybill_no or not identity:
            return False
        return waybill_no, identity

    @classmethod
    def _legacy_customer_key(cls, row):
        return cls._customer_group_key(
            {
                "waybill_no": row.get("waybill_no"),
                "customer_line_no": "",
                "store_no": row.get("store_no"),
                "store_name": row.get("store_name"),
                "customer_no": row.get("customer_no"),
                "customer_name": row.get("customer_name"),
            }
        )

    @classmethod
    def _build_legacy_customer_line_no(cls, waybill_no, store_no, customer_no, sequence):
        seed = (store_no or customer_no or "").strip()
        if seed:
            seed = seed.replace(" ", "_")[:32]
        else:
            seed = f"{sequence:03d}"
        return f"{(waybill_no or 'WB').strip()}-CL-{seed}"

    @classmethod
    def _get_effective_customer_line_no(cls, row, *, customer_line_no_map, customer_line_seq_by_waybill):
        customer_key = cls._customer_group_key(row)
        if customer_key in customer_line_no_map:
            return customer_line_no_map[customer_key]
        explicit_no = (row.get("customer_line_no") or "").strip()
        if explicit_no:
            customer_line_no_map[customer_key] = explicit_no
            return explicit_no
        waybill_no = customer_key[0]
        next_seq = customer_line_seq_by_waybill.get(waybill_no, 0) + 1
        customer_line_seq_by_waybill[waybill_no] = next_seq
        generated_no = f"{waybill_no}-CL-{next_seq:03d}"
        customer_line_no_map[customer_key] = generated_no
        return generated_no

    @classmethod
    def _build_customer_line_create_vals(cls, row, *, waybill, customer_line_no, customers_by_code, stores_by_code):
        customer_record = customers_by_code.get(row.get("customer_no"))
        store_record = stores_by_code.get(row.get("store_no"))
        vals = {
            "waybill_id": waybill.id,
            "customer_line_no": customer_line_no,
            "internal_customer_code_snapshot": row.get("customer_no") or False,
            "external_customer_code_snapshot": row.get("customer_ref") or False,
            "customer_name_snapshot": row.get("store_name") or row.get("customer_name") or False,
            "delivery_note": row.get("delivery_remark") or False,
            "signoff_requirement": row.get("signoff_requirement") or False,
            "customer_ref": row.get("customer_ref") or False,
        }
        if customer_record:
            vals["customer_id"] = customer_record.id
        if store_record:
            vals["partner_id"] = store_record.id
            if "customer_id" not in vals and store_record.parent_id:
                vals["customer_id"] = store_record.parent_id.id
        return vals

    @classmethod
    def _customer_key(cls, row):
        return (
            (row.get("waybill_no") or "").strip(),
            (row.get("customer_no") or "").strip(),
            (row.get("store_no") or "").strip(),
        )

    @classmethod
    def _goods_key(cls, row):
        return (
            (row.get("waybill_no") or "").strip(),
            (row.get("customer_no") or "").strip(),
            (row.get("store_no") or "").strip(),
            (row.get("goods_code") or "").strip(),
            (row.get("goods_name") or "").strip(),
            (row.get("spec") or "").strip(),
        )

    @classmethod
    def _safe_int(cls, value, *, default=0):
        if value in (None, ""):
            return default
        return int(float(value))

    @classmethod
    def _safe_float(cls, value, *, default=0.0):
        if value in (None, ""):
            return default
        return float(value)
