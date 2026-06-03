import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from odoo import fields
from odoo.exceptions import ValidationError

from .waybill_standard_import_service_v2 import WaybillStandardImportService


class Phase5WorkbookImportService(WaybillStandardImportService):
    TEMPLATE_CODE = "TSL-IMPORT-PHASE5-WORKBOOK-V1"
    TEMPLATE_VERSION = "v1"
    LEGACY_TEMPLATE_CODES = set()
    LEGACY_TEMPLATE_VERSIONS = set()
    OBJECT_TYPE = "phase5_workbook"
    OBJECT_TYPE_LABEL = "Phase5 Workbook Import"
    TEMPLATE_FILE_NAME = "TSL-IMPORT-PHASE5-WORKBOOK-V1.xlsx"
    TEMPLATE_VARIANTS = {
        "zh_CN": {
            "label": "五期导入模板（五张样例表）",
            "file_name": TEMPLATE_FILE_NAME,
            "field_locale": "zh_CN",
        }
    }
    DEFAULT_TEMPLATE_LOCALE = "zh_CN"
    INPUT_MODE = "phase5_workbook"
    FLOAT_PATTERN = re.compile(r"-?\d+(?:\.\d+)?")
    WEEKDAY_CODES = (
        ("周一配送", "mon"),
        ("周二配送", "tue"),
        ("周三配送", "wed"),
        ("周四配送", "thu"),
        ("周五配送", "fri"),
        ("周六配送", "sat"),
        ("周日配送", "sun"),
    )

    PRODUCT_SHEET = {
        "key": "product_rows",
        "file_kind": "product_profile",
        "sheet_name": "商品资料",
        "fields": [
            "名称", "规格", "编号", "品牌", "类别", "全级分类", "换算关系",
            "小单位标准价", "中单位标准价", "大单位标准价",
            "小单位采购价", "中单位采购价", "大单位采购价",
            "小单位建议零售价", "中单位建议零售价", "大单位建议零售价",
            "小单位商品条码", "中单位商品条码", "大单位商品条码",
            "小单位销售单位", "中单位销售单位", "大单位销售单位",
            "默认供应商", "采购员", "销项税率(%)", "长(cm)", "宽(cm)", "高(cm)",
            "体积(m³)", "体积单位（大单位）", "毛重(kg)", "毛重单位（大单位）",
            "标签", "起订量", "备注", "大单位数量", "小单位毛重（kg）", "中单位毛重（kg）",
            "小单位体积(m³)", "中单位体积(m³)", "品牌方名称", "商品名称", "是否可以压货",
            "小单位数量", "中单位数量",
        ],
        "required_fields": {"编号", "名称", "规格"},
    }
    CUSTOMER_SHEET = {
        "key": "customer_rows",
        "file_kind": "customer_profile",
        "sheet_name": "客户资料",
        "fields": [
            "客户编号", "客户名称", "客户等级", "所属组织", "部门", "联系人", "手机号码",
            "省", "市", "区", "街道", "详细地址", "业务员", "线路", "渠道",
            "是否允许货到付款", "经度", "纬度", "外部编码", "标签", "状态", "客户状态",
            "内部往来单位", "发票类型", "注册电话", "注册地址", "客户序号", "上楼层数", "违停",
            "地库限高", "免费停车时长_分钟", "停车费_小时", "开始收货时间", "截止收货时间",
            "收货时间段", "不收货时间段", "停车位置", "停车方式", "卸货入口", "卸货位置",
            "周一配送", "周二配送", "周三配送", "周四配送", "周五配送", "周六配送", "周日配送",
        ],
        "required_fields": {"客户编号", "客户名称", "详细地址"},
    }
    ORDER_DETAIL_SHEET = {
        "key": "order_detail_rows",
        "file_kind": "order_detail",
        "sheet_name": "排线订单详情",
        "fields": [
            "批次号", "排线序号", "运单号", "销售订单号", "客户名称", "订单备注", "收货地址",
            "整件数", "散件数", "总重量_kg", "总体积_m3", "集货位", "司机姓名", "车牌号",
        ],
        "required_fields": {"运单号", "销售订单号", "客户名称", "整件数", "散件数"},
    }
    STORE_DETAIL_SHEET = {
        "key": "store_detail_rows",
        "file_kind": "store_detail",
        "sheet_name": "排线门店详情",
        "fields": [
            "批次号", "排线序号", "运单号", "客户名称", "收货地址", "联系人", "联系电话", "集货位",
            "不收货时间段", "门店备注", "整件数", "散件数", "总重量_kg", "总体积_m3", "经度", "纬度", "司机姓名", "车牌号",
        ],
        "required_fields": {"运单号", "客户名称", "收货地址"},
    }
    GOODS_TRIPLET_SHEET = {
        "key": "goods_triplet_rows",
        "file_kind": "store_goods_triplet",
        "sheet_name": "门店货物三联单",
        "fields": [
            "批次号", "排线序号", "运单号", "销售订单号", "客户名称", "行号", "商品名称", "商品编号", "商品条码",
            "包装规格", "发货数量", "小单位销售单位", "小单位销售单价", "金额", "收货地址", "联系人", "联系电话", "业务员", "业务员联系方式",
        ],
        "required_fields": {"运单号", "销售订单号", "商品编号", "商品条码", "包装规格", "发货数量", "小单位销售单位", "金额"},
    }
    SHEETS = [PRODUCT_SHEET, CUSTOMER_SHEET, ORDER_DETAIL_SHEET, STORE_DETAIL_SHEET, GOODS_TRIPLET_SHEET]
    FIELD_LABELS = {field: field for sheet_meta in SHEETS for field in sheet_meta["fields"]}
    FIELD_LABELS["template_file"] = "模板文件"
    FIELD_ALIASES = {field: [field] for sheet_meta in SHEETS for field in sheet_meta["fields"]}

    @classmethod
    def _build_template_download_url(cls, *, template_locale):
        return (
            "/api/admin/logistics/imports/phase5-workbook/template/download"
            f"?template_code={cls.TEMPLATE_CODE}"
            f"&template_version={cls.TEMPLATE_VERSION}"
            f"&template_locale={template_locale}"
        )

    @classmethod
    def _empty_source_data(cls):
        return {"rows": [], "input_mode": cls.INPUT_MODE, "file_kind": ""}

    @classmethod
    def load_template_bytes(cls, *, template_locale=""):
        cls.get_template_variant(template_locale)
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        samples = cls._build_template_samples()
        for sheet_meta in cls.SHEETS:
            worksheet = workbook.create_sheet(sheet_meta["sheet_name"])
            worksheet.append(sheet_meta["fields"])
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
            for sample_row in samples[sheet_meta["key"]]:
                worksheet.append([sample_row.get(field_code, "") for field_code in sheet_meta["fields"]])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @classmethod
    def _build_template_samples(cls):
        return {
            "product_rows": [{"名称": "矿泉水", "规格": "550ml*24瓶", "编号": "P50001", "品牌": "示例品牌", "类别": "饮料", "小单位标准价": "3.50", "小单位采购价": "2.80", "小单位建议零售价": "4.00", "小单位商品条码": "6900000000001", "小单位销售单位": "瓶", "默认供应商": "示例供应商", "备注": "最小可用示例", "大单位数量": "24", "小单位数量": "1", "中单位数量": "12"}],
            "customer_rows": [{"客户编号": "C50001", "客户名称": "世纪大道门店", "联系人": "张三", "手机号码": "13910000001", "省": "上海市", "市": "上海市", "区": "浦东新区", "街道": "世纪大道", "详细地址": "上海市浦东新区世纪大道100号", "业务员": "业务员A", "渠道": "直营", "经度": "121.54420", "纬度": "31.22110", "外部编码": "EXT-C50001", "不收货时间段": "12:00-13:00", "周一配送": "是", "周三配送": "是", "周五配送": "是"}],
            "order_detail_rows": [{"批次号": "PL-BT-20260429-01", "排线序号": "1", "运单号": "PL-WB-20260429-01", "销售订单号": "SO-20260429-01", "客户名称": "世纪大道门店", "订单备注": "优先配送", "收货地址": "上海市浦东新区世纪大道100号", "整件数": "10", "散件数": "2", "总重量_kg": "120.5", "总体积_m3": "1.28", "集货位": "A-01", "司机姓名": "王司机", "车牌号": "沪A12345"}],
            "store_detail_rows": [{"批次号": "PL-BT-20260429-01", "排线序号": "1", "运单号": "PL-WB-20260429-01", "客户名称": "世纪大道门店", "收货地址": "上海市浦东新区世纪大道100号", "联系人": "张三", "联系电话": "13910000001", "不收货时间段": "12:00-13:00", "门店备注": "走西侧卸货口", "整件数": "10", "散件数": "2", "总重量_kg": "120.5", "总体积_m3": "1.28", "经度": "121.54420", "纬度": "31.22110"}],
            "goods_triplet_rows": [{"批次号": "PL-BT-20260429-01", "排线序号": "1", "运单号": "PL-WB-20260429-01", "销售订单号": "SO-20260429-01", "客户名称": "世纪大道门店", "行号": "1", "商品名称": "矿泉水", "商品编号": "P50001", "商品条码": "6900000000001", "包装规格": "550ml*24瓶", "发货数量": "2箱", "小单位销售单位": "瓶", "小单位销售单价": "3.50", "金额": "168.00", "收货地址": "上海市浦东新区世纪大道100号", "联系人": "张三", "联系电话": "13910000001", "业务员": "业务员A", "业务员联系方式": "13919990000"}],
        }

    @classmethod
    def _parse_workbook(cls, raw_bytes):
        if not raw_bytes:
            return cls._empty_source_data(), [cls._make_error(sheet_name="模板文件", row_no=0, field_code="template_file", error_code="PRECHECK_PARSE_FAILED", error_message="模板文件内容为空，无法执行预校验。")]
        try:
            workbook = load_workbook(io.BytesIO(raw_bytes), data_only=True)
        except Exception:
            return cls._empty_source_data(), [cls._make_error(sheet_name="模板文件", row_no=0, field_code="template_file", error_code="TEMPLATE_FILE_TYPE_INVALID", error_message="五期导入仅支持 .xlsx 文件。")]
        sheet_meta = False
        worksheet = False
        for sheet_name in workbook.sheetnames:
            candidate = workbook[sheet_name]
            rows = list(candidate.iter_rows(values_only=True, min_row=1, max_row=1))
            if not rows:
                continue
            header_values = [cls._cell_text(value) for value in rows[0]]
            sheet_meta = cls._match_sheet_meta(header_values)
            if sheet_meta:
                worksheet = candidate
                break
        if not worksheet or not sheet_meta:
            return cls._empty_source_data(), [cls._make_error(sheet_name="模板文件", row_no=0, field_code="template_file", error_code="TEMPLATE_AND_FILE_MISMATCH", error_message="未识别到五期导入支持的表头，请上传五期五张样例表之一。")]
        parsed_rows, errors = cls._parse_sheet_rows(worksheet, sheet_meta, sheet_name_override=worksheet.title)
        if errors:
            return cls._empty_source_data(), errors
        for row in parsed_rows:
            row["_file_kind"] = sheet_meta["file_kind"]
        return {"rows": parsed_rows, "input_mode": cls.INPUT_MODE, "file_kind": sheet_meta["file_kind"]}, []

    @classmethod
    def _match_sheet_meta(cls, header_values):
        normalized = [str(item or "").strip() for item in header_values]
        for sheet_meta in cls.SHEETS:
            if normalized == sheet_meta["fields"]:
                return sheet_meta
        return False

    @classmethod
    def _validate_rows(cls, env, source_data):
        del env
        validators = {
            "product_profile": cls._validate_product_rows,
            "customer_profile": cls._validate_customer_rows,
            "order_detail": cls._validate_order_detail_rows,
            "store_detail": cls._validate_store_detail_rows,
            "store_goods_triplet": cls._validate_goods_triplet_rows,
        }
        validator = validators.get(source_data.get("file_kind"))
        return validator(source_data.get("rows", [])) if validator else []

    @classmethod
    def _validate_required_fields(cls, row, required_fields):
        errors = []
        for field_code in required_fields:
            if (row.get(field_code) or "").strip():
                continue
            errors.append(cls._make_error(sheet_name=row["_sheet_name"], row_no=row["_source_row_no"], field_code=field_code, error_code="FIELD_REQUIRED", error_message=f"{field_code} 不能为空。"))
        return errors

    @classmethod
    def _validate_numeric_field(cls, row, field_code, *, allow_negative=False):
        raw_value = (row.get(field_code) or "").strip()
        if not raw_value:
            return []
        try:
            value = float(raw_value)
        except Exception:
            return [cls._make_error(sheet_name=row["_sheet_name"], row_no=row["_source_row_no"], field_code=field_code, error_code="FIELD_FORMAT_INVALID", error_message=f"{field_code} 必须是数字。")]
        if not allow_negative and value < 0:
            return [cls._make_error(sheet_name=row["_sheet_name"], row_no=row["_source_row_no"], field_code=field_code, error_code="FIELD_VALUE_INVALID", error_message=f"{field_code} 不能小于 0。")]
        return []

    @classmethod
    def _validate_product_rows(cls, rows):
        errors = []
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.PRODUCT_SHEET["required_fields"]))
            for field_name in ("小单位标准价", "小单位采购价", "小单位建议零售价", "大单位数量", "小单位数量", "中单位数量"):
                errors.extend(cls._validate_numeric_field(row, field_name))
        return errors

    @classmethod
    def _validate_customer_rows(cls, rows):
        errors = []
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.CUSTOMER_SHEET["required_fields"]))
            errors.extend(cls._validate_numeric_field(row, "经度", allow_negative=True))
            errors.extend(cls._validate_numeric_field(row, "纬度", allow_negative=True))
        return errors

    @classmethod
    def _validate_order_detail_rows(cls, rows):
        errors = []
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.ORDER_DETAIL_SHEET["required_fields"]))
            for field_name in ("整件数", "散件数", "总重量_kg", "总体积_m3"):
                errors.extend(cls._validate_numeric_field(row, field_name))
        return errors

    @classmethod
    def _validate_store_detail_rows(cls, rows):
        errors = []
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.STORE_DETAIL_SHEET["required_fields"]))
            for field_name in ("整件数", "散件数", "总重量_kg", "总体积_m3"):
                errors.extend(cls._validate_numeric_field(row, field_name))
            errors.extend(cls._validate_numeric_field(row, "经度", allow_negative=True))
            errors.extend(cls._validate_numeric_field(row, "纬度", allow_negative=True))
        return errors

    @classmethod
    def _validate_goods_triplet_rows(cls, rows):
        errors = []
        for row in rows:
            errors.extend(cls._validate_required_fields(row, cls.GOODS_TRIPLET_SHEET["required_fields"]))
            errors.extend(cls._validate_numeric_field(row, "小单位销售单价"))
            errors.extend(cls._validate_numeric_field(row, "金额"))
        return errors

    @classmethod
    def _build_row_business_key(cls, row):
        file_kind = row.get("_file_kind") or "phase5"
        row_no = row.get("_source_row_no") or 0
        if file_kind == "product_profile":
            detail = row.get("编号") or row.get("名称") or "--"
        elif file_kind == "customer_profile":
            detail = row.get("客户编号") or row.get("客户名称") or "--"
        elif file_kind == "order_detail":
            detail = " / ".join(filter(None, [row.get("运单号"), row.get("销售订单号"), row.get("客户名称")])) or "--"
        elif file_kind == "store_detail":
            detail = " / ".join(filter(None, [row.get("运单号"), row.get("客户名称")])) or "--"
        else:
            detail = " / ".join(filter(None, [row.get("运单号"), row.get("销售订单号"), row.get("商品编号") or row.get("商品名称")])) or "--"
        return f"{file_kind}@{row_no} {detail}"[:128]

    @classmethod
    def _execute_task_import(cls, env, task, source_data):
        file_kind = source_data.get("file_kind")
        if file_kind == "product_profile":
            return cls._execute_product_import(env, task, source_data)
        if file_kind == "customer_profile":
            return cls._execute_customer_import(env, task, source_data)
        if file_kind == "order_detail":
            return cls._execute_order_detail_import(env, task, source_data)
        if file_kind == "store_detail":
            return cls._execute_store_detail_import(env, task, source_data)
        if file_kind == "store_goods_triplet":
            return cls._execute_goods_triplet_import(env, task, source_data)
        raise ValidationError("未识别到可执行的五期导入文件类型。")

    @classmethod
    def _execute_product_import(cls, env, task, source_data):
        task_line_map = {line.business_key: line for line in task.task_line_ids}
        error_line_model = env["logistics.import.error.line"].sudo()
        product_model = env["product.template"].sudo()
        product_unit_model = env["logistics.product.unit"].sudo()
        success_count = 0
        fail_count = 0
        for row in source_data.get("rows", []):
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                product_code = (row.get("编号") or "").strip()
                product = product_model.search([("external_product_code", "=", product_code)], limit=1)
                product_name = (row.get("商品名称") or row.get("名称") or "").strip()
                product_vals = {
                    "external_product_code": product_code,
                    "product_name": product_name,
                    "name": product_name,
                    "brand_name": cls._false_if_blank(row.get("品牌")),
                    "category_name": cls._false_if_blank(row.get("类别")),
                    "product_tag": cls._false_if_blank(row.get("标签")),
                    "default_barcode": cls._false_if_blank(row.get("小单位商品条码")),
                    "base_unit_name": cls._fallback_unit_name(row.get("小单位销售单位")),
                    "default_weight": cls._safe_float(row.get("小单位毛重（kg）"), default=0.0),
                    "default_volume": cls._safe_float(row.get("小单位体积(m³)"), default=0.0),
                    "delivery_requirement_text": cls._false_if_blank(row.get("备注")),
                }
                if product:
                    product.write(product_vals)
                else:
                    product = product_model.create(product_vals)
                sale_unit_name = cls._fallback_unit_name(row.get("小单位销售单位"))
                spec_desc = cls._false_if_blank(row.get("规格")) or "未定义规格"
                unit = product_unit_model.search([("product_tmpl_id", "=", product.id), ("spec_desc", "=", spec_desc), ("sale_unit_name", "=", sale_unit_name)], limit=1)
                unit_vals = {
                    "product_tmpl_id": product.id,
                    "sku_code": product_code or False,
                    "spec_desc": spec_desc,
                    "full_category_name": cls._false_if_blank(row.get("全级分类")),
                    "unit_name": sale_unit_name,
                    "sale_unit_name": sale_unit_name,
                    "convert_to_base": cls._parse_float_like_text(row.get("换算关系")),
                    "barcode": cls._false_if_blank(row.get("小单位商品条码")),
                    "standard_price": cls._safe_float(row.get("小单位标准价"), default=0.0),
                    "purchase_price": cls._safe_float(row.get("小单位采购价"), default=0.0),
                    "retail_price": cls._safe_float(row.get("小单位建议零售价"), default=0.0),
                    "sale_tax_rate": cls._safe_float(row.get("销项税率(%)"), default=0.0),
                    "length_cm": cls._safe_float(row.get("长(cm)"), default=0.0),
                    "width_cm": cls._safe_float(row.get("宽(cm)"), default=0.0),
                    "height_cm": cls._safe_float(row.get("高(cm)"), default=0.0),
                    "weight": cls._safe_float(row.get("毛重(kg)"), default=0.0),
                    "small_unit_weight": cls._safe_float(row.get("小单位毛重（kg）"), default=0.0),
                    "middle_unit_weight": cls._safe_float(row.get("中单位毛重（kg）"), default=0.0),
                    "volume": cls._safe_float(row.get("体积(m³)"), default=0.0),
                    "small_unit_volume": cls._safe_float(row.get("小单位体积(m³)"), default=0.0),
                    "middle_unit_volume": cls._safe_float(row.get("中单位体积(m³)"), default=0.0),
                    "volume_unit_large": cls._false_if_blank(row.get("体积单位（大单位）")),
                    "gross_weight_unit_large": cls._false_if_blank(row.get("毛重单位（大单位）")),
                    "small_unit_qty": cls._safe_float(row.get("小单位数量"), default=0.0),
                    "middle_unit_qty": cls._safe_float(row.get("中单位数量"), default=0.0),
                    "large_unit_qty": cls._safe_float(row.get("大单位数量"), default=0.0),
                    "brand_owner_name": cls._false_if_blank(row.get("品牌方名称")),
                    "default_vendor_name": cls._false_if_blank(row.get("默认供应商")),
                    "buyer_name": cls._false_if_blank(row.get("采购员")),
                    "min_order_qty": cls._safe_float(row.get("起订量"), default=0.0),
                    "can_press_stock": cls._normalize_boolean_value(row.get("是否可以压货")),
                    "unit_remark": cls._false_if_blank(row.get("备注")),
                }
                if unit:
                    unit.write(unit_vals)
                else:
                    unit = product_unit_model.create(unit_vals)
                success_count += 1
                cls._update_task_line(task_line, status="success", message="商品与规格信息写入成功。", target_model="logistics.product.unit", target_res_id=unit.id)
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"商品导入失败：{exc}")
                error_line_model.create(cls._make_runtime_error_line_vals(task=task, task_line=task_line, row=row, field_name="编号", error_message=f"商品导入失败：{exc}"))
        return cls._build_task_stats(task, success_count, fail_count, 0)

    @classmethod
    def _execute_customer_import(cls, env, task, source_data):
        task_line_map = {line.business_key: line for line in task.task_line_ids}
        error_line_model = env["logistics.import.error.line"].sudo()
        partner_model = env["res.partner"].sudo()
        customer_profile_model = env["logistics.customer.profile"].sudo()
        store_profile_model = env["logistics.store.profile"].sudo()
        success_count = 0
        fail_count = 0
        for row in source_data.get("rows", []):
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                customer_code = (row.get("客户编号") or "").strip()
                partner = partner_model.search(["|", ("logistics_customer_code", "=", customer_code), ("logistics_store_code", "=", customer_code)], limit=1)
                partner_vals = {
                    "name": cls._false_if_blank(row.get("客户名称")) or customer_code,
                    "customer_name": cls._false_if_blank(row.get("客户名称")) or customer_code,
                    "is_logistics_partner": True,
                    "is_logistics_customer": True,
                    "is_logistics_store": True,
                    "logistics_customer_code": customer_code,
                    "logistics_store_code": customer_code,
                    "internal_customer_code": cls._false_if_blank(row.get("客户序号")),
                    "external_customer_code": cls._false_if_blank(row.get("外部编码")),
                    "contact_name": cls._false_if_blank(row.get("联系人")),
                    "contact_phone": cls._false_if_blank(row.get("手机号码") or row.get("注册电话")),
                    "address_full": cls._false_if_blank(row.get("详细地址") or row.get("注册地址")),
                    "address_region_json": cls._build_partner_address_region_json(row) or False,
                    "organization_name": cls._false_if_blank(row.get("所属组织")),
                    "department_name": cls._false_if_blank(row.get("部门")),
                    "salesperson_name": cls._false_if_blank(row.get("业务员")),
                    "channel_name": cls._false_if_blank(row.get("渠道")),
                    "allow_cash_on_delivery": bool(cls._normalize_boolean_value(row.get("是否允许货到付款"))),
                    "internal_counterparty_flag": bool(cls._normalize_boolean_value(row.get("内部往来单位"))),
                    "invoice_type": cls._false_if_blank(row.get("发票类型")),
                    "logistics_customer_level": cls._normalize_customer_level(row.get("客户等级")),
                    "customer_status": cls._normalize_customer_status(row.get("客户状态")),
                    "receive_start_time": cls._false_if_blank(row.get("开始收货时间")) or "",
                    "receive_end_time": cls._false_if_blank(row.get("截止收货时间")) or "",
                    "receive_time_slots_text": cls._false_if_blank(row.get("收货时间段")) or "",
                    "no_receive_time_slots_text": cls._false_if_blank(row.get("不收货时间段")) or "",
                    "delivery_week_flags": cls._build_delivery_week_flags(row),
                    "illegal_parking_flag": bool(cls._normalize_boolean_value(row.get("违停"))),
                    "free_parking_minutes": cls._safe_int(row.get("免费停车时长_分钟"), default=0),
                    "parking_fee_per_hour": cls._safe_float(row.get("停车费_小时"), default=0.0),
                    "parking_location_text": cls._false_if_blank(row.get("停车位置")) or "",
                    "parking_mode_text": cls._false_if_blank(row.get("停车方式")) or "",
                    "unload_entrance_text": cls._false_if_blank(row.get("卸货入口")) or "",
                    "unload_location_text": cls._false_if_blank(row.get("卸货位置")) or "",
                    "upstairs_floor_count": cls._safe_int(row.get("上楼层数"), default=0),
                    "basement_height_limit_text": cls._false_if_blank(row.get("地库限高")) or "",
                    "city": cls._false_if_blank(row.get("市")),
                    "street": cls._false_if_blank(row.get("街道") or row.get("详细地址")),
                    "partner_longitude": cls._safe_float(row.get("经度"), default=0.0) if row.get("经度") else False,
                    "partner_latitude": cls._safe_float(row.get("纬度"), default=0.0) if row.get("纬度") else False,
                }
                partner_vals = {key: value for key, value in partner_vals.items() if key in partner_model._fields}
                if partner:
                    partner.write(partner_vals)
                else:
                    partner = partner_model.create(partner_vals)
                cls._ensure_profile_record(customer_profile_model, partner.id)
                cls._ensure_profile_record(store_profile_model, partner.id)
                success_count += 1
                cls._update_task_line(task_line, status="success", message="客户主数据写入成功。", target_model="res.partner", target_res_id=partner.id)
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"客户导入失败：{exc}")
                error_line_model.create(cls._make_runtime_error_line_vals(task=task, task_line=task_line, row=row, field_name="客户编号", error_message=f"客户导入失败：{exc}"))
        return cls._build_task_stats(task, success_count, fail_count, 0)

    @classmethod
    def _execute_store_detail_import(cls, env, task, source_data):
        task_line_map = {line.business_key: line for line in task.task_line_ids}
        error_line_model = env["logistics.import.error.line"].sudo()
        success_count = 0
        fail_count = 0
        for row in source_data.get("rows", []):
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                waybill = cls._find_or_create_waybill(env, row)
                customer_line = cls._find_or_create_customer_line(env, row, waybill=waybill)
                cls._update_partner_from_store_snapshot(customer_line.partner_id or customer_line.customer_id, row)
                customer_line.write({
                    "customer_name_snapshot": cls._false_if_blank(row.get("客户名称")) or customer_line.customer_name_snapshot,
                    "contact_name_snapshot": cls._false_if_blank(row.get("联系人")) or customer_line.contact_name_snapshot,
                    "contact_phone_snapshot": cls._false_if_blank(row.get("联系电话")) or customer_line.contact_phone_snapshot,
                    "address_full_snapshot": cls._false_if_blank(row.get("收货地址")) or customer_line.address_full_snapshot,
                    "longitude_snapshot": cls._safe_float(row.get("经度"), default=0.0) if row.get("经度") else customer_line.longitude_snapshot,
                    "latitude_snapshot": cls._safe_float(row.get("纬度"), default=0.0) if row.get("纬度") else customer_line.latitude_snapshot,
                    "delivery_note": cls._false_if_blank(row.get("门店备注")) or customer_line.delivery_note,
                })
                success_count += 1
                cls._update_task_line(task_line, status="success", message="门店快照写入成功。", target_model="logistics.dispatch.waybill.customer.line", target_res_id=customer_line.id)
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"门店详情导入失败：{exc}")
                error_line_model.create(cls._make_runtime_error_line_vals(task=task, task_line=task_line, row=row, field_name="运单号", error_message=f"门店详情导入失败：{exc}"))
        return cls._build_task_stats(task, success_count, fail_count, 0)

    @classmethod
    def _execute_order_detail_import(cls, env, task, source_data):
        task_line_map = {line.business_key: line for line in task.task_line_ids}
        error_line_model = env["logistics.import.error.line"].sudo()
        order_line_model = env["logistics.dispatch.waybill.order.line"].sudo()
        success_count = 0
        fail_count = 0
        for row in source_data.get("rows", []):
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                waybill = cls._find_or_create_waybill(env, row)
                customer_line = cls._find_or_create_customer_line(env, row, waybill=waybill)
                sales_order_no = (row.get("销售订单号") or "").strip()
                order_line = order_line_model.search([("customer_line_id", "=", customer_line.id), "|", ("sales_order_no", "=", sales_order_no), ("source_doc_no", "=", sales_order_no)], limit=1)
                vals = {
                    "waybill_id": waybill.id,
                    "customer_line_id": customer_line.id,
                    "source_doc_no": sales_order_no,
                    "sales_order_no": sales_order_no,
                    "order_remark": cls._false_if_blank(row.get("订单备注")),
                    "goods_summary": cls._false_if_blank(row.get("收货地址")),
                    "whole_package_count": cls._safe_float(row.get("整件数"), default=0.0),
                    "loose_package_count": cls._safe_float(row.get("散件数"), default=0.0),
                    "weight_summary": cls._safe_float(row.get("总重量_kg"), default=0.0),
                    "volume_summary": cls._safe_float(row.get("总体积_m3"), default=0.0),
                    "gathering_location": cls._false_if_blank(row.get("集货位")),
                }
                if order_line:
                    order_line.write(vals)
                else:
                    order_line = order_line_model.create(vals)
                success_count += 1
                cls._update_task_line(task_line, status="success", message="订单快照写入成功。", target_model="logistics.dispatch.waybill.order.line", target_res_id=order_line.id)
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"订单详情导入失败：{exc}")
                error_line_model.create(cls._make_runtime_error_line_vals(task=task, task_line=task_line, row=row, field_name="销售订单号", error_message=f"订单详情导入失败：{exc}"))
        return cls._build_task_stats(task, success_count, fail_count, 0)

    @classmethod
    def _execute_goods_triplet_import(cls, env, task, source_data):
        task_line_map = {line.business_key: line for line in task.task_line_ids}
        error_line_model = env["logistics.import.error.line"].sudo()
        goods_line_model = env["logistics.dispatch.waybill.customer.goods.line"].sudo()
        success_count = 0
        fail_count = 0
        for row in source_data.get("rows", []):
            task_line = task_line_map.get(cls._build_row_business_key(row))
            try:
                waybill = cls._find_or_create_waybill(env, row)
                customer_line = cls._find_or_create_customer_line(env, row, waybill=waybill)
                order_line = cls._find_or_create_order_line(env, row, waybill=waybill, customer_line=customer_line)
                cls._update_customer_line_contact_snapshot(customer_line, row)
                if row.get("业务员") or row.get("业务员联系方式"):
                    order_line.write({
                        "salesperson_name_snapshot": cls._false_if_blank(row.get("业务员")) or order_line.salesperson_name_snapshot,
                        "salesperson_phone": cls._false_if_blank(row.get("业务员联系方式")) or order_line.salesperson_phone,
                    })
                goods_line = cls._find_existing_goods_line(goods_line_model, order_line, row)
                qty_value = cls._parse_float_like_text(row.get("发货数量")) or 1.0
                sequence = cls._safe_int(row.get("行号"), default=0)
                vals = {
                    "customer_line_id": customer_line.id,
                    "order_line_id": order_line.id,
                    "sequence": sequence,
                    "external_product_code_snapshot": cls._false_if_blank(row.get("商品编号")),
                    "product_name_snapshot": cls._false_if_blank(row.get("商品名称")),
                    "spec_snapshot": cls._false_if_blank(row.get("包装规格")),
                    "barcode_snapshot": cls._false_if_blank(row.get("商品条码")),
                    "small_unit_name": cls._false_if_blank(row.get("小单位销售单位")),
                    "delivery_qty_text": cls._false_if_blank(row.get("发货数量")),
                    "small_unit_price": cls._safe_float(row.get("小单位销售单价"), default=0.0),
                    "amount": cls._safe_float(row.get("金额"), default=0.0),
                    "goods_code": cls._false_if_blank(row.get("商品编号")),
                    "goods_name": cls._false_if_blank(row.get("商品名称")) or "未命名商品",
                    "specification": cls._false_if_blank(row.get("包装规格")),
                    "quantity": qty_value,
                    "package_count": 0,
                    "uom_name": cls._false_if_blank(row.get("小单位销售单位")),
                    "weight": 0.0,
                    "volume": 0.0,
                    "remark": False,
                }
                if goods_line:
                    goods_line.write(vals)
                else:
                    goods_line = goods_line_model.create(vals)
                success_count += 1
                cls._update_task_line(task_line, status="success", message="货物快照写入成功。", target_model="logistics.dispatch.waybill.customer.goods.line", target_res_id=goods_line.id)
            except Exception as exc:
                fail_count += 1
                cls._update_task_line(task_line, status="failed", message=f"三联单导入失败：{exc}")
                error_line_model.create(cls._make_runtime_error_line_vals(task=task, task_line=task_line, row=row, field_name="商品编号", error_message=f"三联单导入失败：{exc}"))
        return cls._build_task_stats(task, success_count, fail_count, 0)

    @classmethod
    def _find_or_create_waybill(cls, env, row):
        waybill_model = env["logistics.dispatch.waybill"].sudo()
        batch_model = env["logistics.dispatch.batch"].sudo()
        waybill_no = (row.get("运单号") or "").strip()
        if not waybill_no:
            raise ValidationError("运单号不能为空。")
        waybill = waybill_model.search([("name", "=", waybill_no)], limit=1)
        batch = False
        batch_no = (row.get("批次号") or "").strip()
        if batch_no:
            batch = batch_model.search([("name", "=", batch_no)], limit=1)
        if not waybill:
            create_vals = {"name": waybill_no}
            if batch:
                create_vals["batch_id"] = batch.id
                create_vals["warehouse_id"] = batch.warehouse_id.id
            waybill = waybill_model.create(create_vals)
        elif batch and not waybill.batch_id:
            waybill.write({"batch_id": batch.id, "warehouse_id": batch.warehouse_id.id})
        return waybill

    @classmethod
    def _find_or_create_customer_line(cls, env, row, *, waybill):
        customer_line_model = env["logistics.dispatch.waybill.customer.line"].sudo()
        partner = cls._resolve_or_create_partner(env, row)
        customer_name = cls._false_if_blank(row.get("客户名称"))
        address = cls._false_if_blank(row.get("收货地址") or row.get("详细地址"))
        domain = [("waybill_id", "=", waybill.id)]
        if partner:
            domain.append(("partner_id", "=", partner.id))
        elif customer_name:
            domain.append(("customer_name_snapshot", "=", customer_name))
        customer_line = customer_line_model.search(domain, limit=1)
        if not customer_line and customer_name and address:
            customer_line = customer_line_model.search([("waybill_id", "=", waybill.id), ("customer_name_snapshot", "=", customer_name), ("address_full_snapshot", "=", address)], limit=1)
        vals = {
            "waybill_id": waybill.id,
            "partner_id": partner.id if partner else False,
            "customer_id": partner.id if partner else False,
            "customer_name_snapshot": customer_name or False,
            "contact_name_snapshot": cls._false_if_blank(row.get("联系人")) or False,
            "contact_phone_snapshot": cls._false_if_blank(row.get("联系电话") or row.get("手机号码")) or False,
            "address_full_snapshot": address or False,
            "longitude_snapshot": cls._safe_float(row.get("经度"), default=0.0) if row.get("经度") else False,
            "latitude_snapshot": cls._safe_float(row.get("纬度"), default=0.0) if row.get("纬度") else False,
            "delivery_note": cls._false_if_blank(row.get("门店备注")) or False,
        }
        if customer_line:
            customer_line.write({key: value for key, value in vals.items() if value not in (False, None, "")})
            return customer_line
        vals["customer_line_no"] = customer_line_model._generate_customer_line_no(waybill.id)
        return customer_line_model.create(vals)

    @classmethod
    def _find_or_create_order_line(cls, env, row, *, waybill, customer_line):
        order_line_model = env["logistics.dispatch.waybill.order.line"].sudo()
        sales_order_no = (row.get("销售订单号") or "").strip()
        order_line = order_line_model.search([("customer_line_id", "=", customer_line.id), "|", ("sales_order_no", "=", sales_order_no), ("source_doc_no", "=", sales_order_no)], limit=1)
        if order_line:
            return order_line
        return order_line_model.create({"waybill_id": waybill.id, "customer_line_id": customer_line.id, "source_doc_no": sales_order_no, "sales_order_no": sales_order_no})

    @classmethod
    def _find_existing_goods_line(cls, goods_line_model, order_line, row):
        sequence = cls._safe_int(row.get("行号"), default=0)
        goods_code = (row.get("商品编号") or "").strip()
        if sequence:
            goods_line = goods_line_model.search([("order_line_id", "=", order_line.id), ("sequence", "=", sequence)], limit=1)
            if goods_line:
                return goods_line
        if goods_code:
            return goods_line_model.search([("order_line_id", "=", order_line.id), ("goods_code", "=", goods_code)], limit=1)
        return False

    @classmethod
    def _resolve_or_create_partner(cls, env, row):
        partner_model = env["res.partner"].sudo()
        code = (row.get("客户编号") or "").strip()
        customer_name = (row.get("客户名称") or "").strip()
        partner = False
        if code:
            partner = partner_model.search(["|", ("logistics_customer_code", "=", code), ("logistics_store_code", "=", code)], limit=1)
        if not partner and customer_name:
            partner = partner_model.search([("name", "=", customer_name)], limit=1)
        if partner:
            return partner
        if not customer_name:
            return False
        vals = {
            "name": customer_name,
            "customer_name": customer_name,
            "is_logistics_partner": True,
            "is_logistics_customer": True,
            "is_logistics_store": True,
            "logistics_customer_code": code or False,
            "logistics_store_code": code or False,
            "address_full": cls._false_if_blank(row.get("收货地址") or row.get("详细地址")),
            "contact_name": cls._false_if_blank(row.get("联系人")),
            "contact_phone": cls._false_if_blank(row.get("联系电话") or row.get("手机号码")),
        }
        vals = {key: value for key, value in vals.items() if key in partner_model._fields}
        partner = partner_model.create(vals)
        cls._ensure_profile_record(env["logistics.customer.profile"].sudo(), partner.id)
        cls._ensure_profile_record(env["logistics.store.profile"].sudo(), partner.id)
        return partner

    @classmethod
    def _update_partner_from_store_snapshot(cls, partner, row):
        if not partner:
            return
        vals = {}
        if row.get("客户名称"):
            vals["name"] = row.get("客户名称")
            vals["customer_name"] = row.get("客户名称")
        if row.get("联系人"):
            vals["contact_name"] = row.get("联系人")
        if row.get("联系电话"):
            vals["contact_phone"] = row.get("联系电话")
        if row.get("收货地址"):
            vals["address_full"] = row.get("收货地址")
        if row.get("不收货时间段"):
            vals["no_receive_time_slots_text"] = row.get("不收货时间段")
        if row.get("经度") and "partner_longitude" in partner._fields:
            vals["partner_longitude"] = cls._safe_float(row.get("经度"), default=0.0)
        if row.get("纬度") and "partner_latitude" in partner._fields:
            vals["partner_latitude"] = cls._safe_float(row.get("纬度"), default=0.0)
        if vals:
            partner.write(vals)

    @classmethod
    def _update_customer_line_contact_snapshot(cls, customer_line, row):
        vals = {}
        if row.get("收货地址"):
            vals["address_full_snapshot"] = row.get("收货地址")
        if row.get("联系人"):
            vals["contact_name_snapshot"] = row.get("联系人")
        if row.get("联系电话"):
            vals["contact_phone_snapshot"] = row.get("联系电话")
        if vals:
            customer_line.write(vals)

    @classmethod
    def _ensure_profile_record(cls, profile_model, partner_id):
        if not partner_id:
            return
        if profile_model.search([("partner_id", "=", partner_id)], limit=1):
            return
        profile_model.create({"partner_id": partner_id})

    @classmethod
    def _build_partner_address_region_json(cls, row):
        payload = {
            "province": cls._false_if_blank(row.get("省")),
            "city": cls._false_if_blank(row.get("市")),
            "district": cls._false_if_blank(row.get("区")),
            "street": cls._false_if_blank(row.get("街道")),
        }
        return {key: value for key, value in payload.items() if value}

    @classmethod
    def _build_delivery_week_flags(cls, row):
        flags = []
        for field_label, code in cls.WEEKDAY_CODES:
            if cls._normalize_boolean_value(row.get(field_label)):
                flags.append(code)
        return ",".join(flags)

    @classmethod
    def _normalize_customer_level(cls, raw_value):
        value = (raw_value or "").strip().lower()
        if value == "vip":
            return "vip"
        if value in {"战略", "strategic"}:
            return "strategic"
        if value:
            return "standard"
        return False

    @classmethod
    def _normalize_customer_status(cls, raw_value):
        value = (raw_value or "").strip().lower()
        if value in {"inactive", "停用"}:
            return "inactive"
        if value in {"paused", "暂停"}:
            return "paused"
        if value:
            return "normal"
        return False

    @classmethod
    def _parse_float_like_text(cls, raw_value):
        text = (raw_value or "").strip()
        if not text:
            return 0.0
        matched = cls.FLOAT_PATTERN.search(text.replace(",", ""))
        if not matched:
            return 0.0
        try:
            return float(matched.group())
        except Exception:
            return 0.0

    @classmethod
    def _fallback_unit_name(cls, raw_value):
        return (raw_value or "").strip() or "未配置"

    @classmethod
    def _false_if_blank(cls, raw_value):
        text = (raw_value or "").strip()
        return text or False

    @classmethod
    def _build_task_stats(cls, task, success_count, fail_count, skipped_count):
        task_status = "success"
        if fail_count and success_count:
            task_status = "partial_failed"
        elif fail_count and not success_count:
            task_status = "failed"
        return {
            "task_status": task_status,
            "success_count": success_count,
            "fail_count": fail_count,
            "skipped_count": skipped_count,
            "summary_message": f"五期导入完成，共 {task.total_count} 行，成功 {success_count} 行，失败 {fail_count} 行，跳过 {skipped_count} 行。",
        }
