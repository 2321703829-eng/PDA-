import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from odoo import fields
from odoo.exceptions import ValidationError


class DriverRouteExcelExportError(ValidationError):
    def __init__(self, error_code, message):
        self.error_code = error_code
        super().__init__(message)


class DriverRouteExcelExportService:
    DOWNLOAD_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    SHEET_NAME = "司机路线清单"
    FILE_NAME_TEMPLATE = "driver_route_{delivery_date}.xlsx"
    HEADERS = [
        "批次号",
        "送货顺序",
        "运单号",
        "客户名称",
        "门店名称",
        "联系人",
        "联系电话",
        "地址",
        "客户/门店备注",
    ]
    CRITICAL_FIELDS = {
        "批次号": "批次号",
        "送货顺序": "送货顺序",
        "运单号": "运单号",
        "客户名称": "客户名称",
        "门店名称": "门店名称",
        "联系电话": "联系电话",
        "地址": "地址",
    }
    COLUMN_WIDTHS = {
        "A": 18,
        "B": 10,
        "C": 22,
        "D": 18,
        "E": 18,
        "F": 14,
        "G": 18,
        "H": 42,
        "I": 28,
    }
    NOTE_FILL = PatternFill(fill_type="solid", fgColor="FFF3CD")
    NOTE_FONT = Font(bold=True, color="9A6700")
    HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
    HEADER_FONT = Font(bold=True, color="1F2937")
    ERROR_FILL = PatternFill(fill_type="solid", fgColor="FDE2E1")

    @classmethod
    def export_by_delivery_date(cls, env, *, delivery_date="", file_locale="zh_CN"):
        del file_locale
        cls._ensure_driver_route_export_access(env)
        normalized_date = cls._normalize_delivery_date(delivery_date)
        waybills = cls._get_waybills(env, normalized_date)
        if not waybills:
            raise DriverRouteExcelExportError(
                "DRIVER_ROUTE_EXPORT_EMPTY",
                f"No route data found for delivery_date {normalized_date}.",
            )

        customer_lines = cls._get_customer_lines(env, waybills)
        rows, anomaly_summary = cls._build_rows(waybills, customer_lines)
        workbook_bytes = cls._build_workbook_bytes(normalized_date, rows, anomaly_summary)
        return {
            "file_name": cls.FILE_NAME_TEMPLATE.format(delivery_date=normalized_date),
            "content_type": cls.DOWNLOAD_CONTENT_TYPE,
            "file_bytes": workbook_bytes,
        }

    @classmethod
    def _ensure_driver_route_export_access(cls, env):
        if env.user.has_group("logistics_dispatch.group_logistics_export_user"):
            return
        raise DriverRouteExcelExportError(
            "EXPORT_PERMISSION_DENIED",
            "You do not have permission to export the driver route workbook.",
        )

    @classmethod
    def _normalize_delivery_date(cls, delivery_date):
        if not delivery_date:
            raise DriverRouteExcelExportError("DELIVERY_DATE_REQUIRED", "delivery_date is required.")
        normalized = fields.Date.to_date(delivery_date)
        if not normalized:
            raise DriverRouteExcelExportError("DELIVERY_DATE_INVALID", f"Invalid delivery_date: {delivery_date}")
        return normalized.isoformat()

    @classmethod
    def _get_waybills(cls, env, delivery_date):
        return env["logistics.dispatch.waybill"].sudo().search(
            [("delivery_date", "=", delivery_date)],
            order="batch_id asc, route_seq asc, name asc, id asc",
        )

    @classmethod
    def _get_customer_lines(cls, env, waybills):
        if not waybills:
            return env["logistics.dispatch.waybill.customer.line"].browse()
        return env["logistics.dispatch.waybill.customer.line"].sudo().search(
            [("waybill_id", "in", waybills.ids)],
            order="waybill_id asc, stop_seq_in_waybill asc, sequence asc, id asc",
        )

    @classmethod
    def _build_rows(cls, waybills, customer_lines):
        lines_by_waybill = {}
        for line in customer_lines:
            lines_by_waybill.setdefault(line.waybill_id.id, []).append(line)

        rows = []
        anomaly_fields = set()
        waybills_without_stops = []

        for waybill in waybills:
            lines = lines_by_waybill.get(waybill.id, [])
            if not lines:
                waybills_without_stops.append(waybill.name or str(waybill.id))
                continue
            for line in lines:
                row = cls._build_row_dict(waybill, line)
                row_anomalies = cls._collect_row_anomalies(row)
                if row_anomalies:
                    anomaly_fields.update(row_anomalies)
                rows.append(
                    {
                        "values": row,
                        "anomalies": row_anomalies,
                    }
                )

        rows.sort(key=cls._sort_key)
        anomaly_summary = cls._build_anomaly_summary(rows, anomaly_fields, waybills_without_stops)
        return rows, anomaly_summary

    @classmethod
    def _build_row_dict(cls, waybill, line):
        partner = line.partner_id or line.customer_id or waybill.partner_id or waybill.customer_id or waybill.store_id
        batch = waybill.batch_id

        customer_name = (
            (line.customer_name_snapshot or "").strip()
            or (line.customer_name or "").strip()
            or (line.partner_name or "").strip()
            or (waybill.customer_name or "").strip()
            or (waybill.partner_name or "").strip()
        )
        store_name = (
            (line.store_name or "").strip()
            or (partner.name or "").strip()
            or customer_name
        )
        contact_name = (
            (line.contact_name_snapshot or "").strip()
            or ((partner.contact_name or "").strip() if partner else "")
        )
        contact_phone = (
            (line.contact_phone_snapshot or "").strip()
            or ((partner.contact_phone or "").strip() if partner else "")
        )
        address = (
            (line.address_full_snapshot or "").strip()
            or ((partner.address_full or "").strip() if partner else "")
            or ((partner.contact_address or "").strip() if partner else "")
        )
        remark = (
            (line.delivery_note or "").strip()
            or ((partner.logistics_service_note or "").strip() if partner else "")
            or (waybill.delivery_remark_snapshot or "").strip()
            or (waybill.remark or "").strip()
            or ((batch.remark or "").strip() if batch else "")
        )

        return {
            "批次号": (batch.name or "").strip() if batch else "",
            "送货顺序": line.stop_seq_in_waybill or "",
            "运单号": (waybill.name or "").strip(),
            "客户名称": customer_name,
            "门店名称": store_name,
            "联系人": contact_name,
            "联系电话": contact_phone,
            "地址": address,
            "客户/门店备注": remark,
        }

    @classmethod
    def _collect_row_anomalies(cls, row):
        anomalies = []
        for key, label in cls.CRITICAL_FIELDS.items():
            value = row.get(key)
            if value in (None, "", False):
                anomalies.append(label)
        return anomalies

    @classmethod
    def _sort_key(cls, row_item):
        values = row_item["values"]
        stop_seq = values.get("送货顺序")
        normalized_stop_seq = stop_seq if isinstance(stop_seq, int) and stop_seq > 0 else 999999
        return (
            values.get("批次号") or "ZZZ",
            normalized_stop_seq,
            values.get("运单号") or "",
            values.get("客户名称") or "",
            values.get("门店名称") or "",
        )

    @classmethod
    def _build_anomaly_summary(cls, rows, anomaly_fields, waybills_without_stops):
        row_issue_count = sum(1 for row in rows if row["anomalies"])
        warnings = []
        if row_issue_count:
            warnings.append(
                f"本表包含 {row_issue_count} 条需人工复查的停靠点，请重点检查 {', '.join(sorted(anomaly_fields))}。"
            )
        if waybills_without_stops:
            warnings.append(f"另有 {len(waybills_without_stops)} 个运单缺少停靠点数据，请人工复查。")
        return {
            "has_warning": bool(warnings),
            "message": "异常提醒：" + ("；".join(warnings) if warnings else "无。"),
        }

    @classmethod
    def _build_workbook_bytes(cls, delivery_date, rows, anomaly_summary):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = cls.SHEET_NAME

        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cls.HEADERS))
        note_cell = sheet.cell(row=1, column=1, value=f"导出日期：{delivery_date}；{anomaly_summary['message']}")
        note_cell.fill = cls.NOTE_FILL
        note_cell.font = cls.NOTE_FONT
        note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 24

        for index, header in enumerate(cls.HEADERS, start=1):
            cell = sheet.cell(row=2, column=index, value=header)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_index, row_item in enumerate(rows, start=3):
            values = row_item["values"]
            anomalies = set(row_item["anomalies"])
            for col_index, header in enumerate(cls.HEADERS, start=1):
                cell = sheet.cell(row=row_index, column=col_index, value=values.get(header, ""))
                cell.alignment = Alignment(vertical="top", wrap_text=header in ("地址", "客户/门店备注"))
                if header in anomalies:
                    cell.fill = cls.ERROR_FILL

        sheet.freeze_panes = "A3"
        last_row = max(sheet.max_row, 2)
        sheet.auto_filter.ref = f"A2:I{last_row}"
        for column_letter, width in cls.COLUMN_WIDTHS.items():
            sheet.column_dimensions[column_letter].width = width

        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()
