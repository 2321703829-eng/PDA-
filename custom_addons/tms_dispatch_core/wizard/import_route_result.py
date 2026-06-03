import base64
import io

from odoo import _, api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:
    openpyxl = None


class ImportRouteResult(models.TransientModel):
    _name = "tms.import.route.result"
    _description = "导入排线结果"

    excel_file = fields.Binary(string="排线结果 Excel", required=True, attachment=True)
    file_name = fields.Char(string="文件名")
    update_existing = fields.Boolean(string="更新已有批次", default=False,
        help="若勾选, 同名批次号将被覆盖; 否则只新增。")

    def action_import(self):
        if not openpyxl:
            raise UserError(_("请安装 openpyxl 库: pip install openpyxl"))
        if not self.excel_file:
            raise UserError(_("请先上传排线结果 Excel 文件"))

        data = base64.b64decode(self.excel_file)
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        sheet = None
        for name in wb.sheetnames:
            if "明细" in name or "detail" in name.lower():
                sheet = wb[name]
                break
        if not sheet:
            sheet = wb[wb.sheetnames[0]]

        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        if not rows:
            raise UserError(_("排线结果为空, 请检查 Excel 内容"))

        # 列: 车辆序号, 顺序, 站号, 收货单位, 收货地址, 经度, 纬度, 重量, 件数, 方数
        groups = {}
        for row in rows:
            if not row[0]:
                continue
            vehicle_no = str(int(float(str(row[0])))) if row[0] else "0"
            if vehicle_no not in groups:
                groups[vehicle_no] = []
            groups[vehicle_no].append({
                "stop_seq": int(float(str(row[1]))) if row[1] else 1,
                "stop_code": str(row[2]) if row[2] else "",
                "store_name": str(row[3])[:100] if row[3] else "",
                "address": str(row[4])[:200] if row[4] else "",
                "lng": float(str(row[5])) if row[5] else 0,
                "lat": float(str(row[6])) if row[6] else 0,
                "weight": float(str(row[7])) if row[7] else 0,
                "waybill_no": "",
            })

        Batch = self.env["logistics.route.planning.batch"]
        StopLine = self.env["logistics.route.planning.stop.line"]
        created = 0

        for vehicle_no, stops in groups.items():
            batch_no = f"PC{int(vehicle_no):03d}-{fields.Date.today().strftime('%Y%m%d')}-IMP"
            # 查重
            existing = Batch.search([("batch_no", "=", batch_no)], limit=1)
            if existing and self.update_existing:
                existing.stop_line_ids.unlink()
                batch = existing
                batch.write({"stop_count": len(stops)})
            elif existing:
                continue
            else:
                batch = Batch.create({
                    "batch_no": batch_no,
                    "delivery_date": fields.Date.today(),
                    "entry_mode": "import_sheet",
                    "planning_state": "result_ready",
                    "route_status": "route_planned",
                    "warehouse_id": 1,
                    "stop_count": len(stops),
                })

            for s in stops:
                StopLine.create({
                    "batch_id": batch.id,
                    "stop_seq": s["stop_seq"],
                    "store_name": s["store_name"],
                    "waybill_no": s["waybill_no"] or f"YD{int(vehicle_no):02d}{s['stop_seq']:03d}",
                    "longitude": s["lng"],
                    "latitude": s["lat"],
                    "address_detail": s["address"],
                })
            created += 1

        return {
            "type": "ir.actions.act_window",
            "name": _("排线调度"),
            "res_model": "logistics.route.planning.batch",
            "view_mode": "list,form",
            "domain": [],
        }
