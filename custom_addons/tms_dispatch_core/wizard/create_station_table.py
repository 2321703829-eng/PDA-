import base64
import io
import json
import logging
import urllib.parse
import urllib.request

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

AMAP_KEY = "f3d2d7ffa1a949b52d4735a6f014c871"
AMAP_GEOCODE_URL = "https://restapi.amap.com/v3/geocode/geo"


class CreateStationTable(models.TransientModel):
    _name = "tms.create.station.table"
    _description = "生成排线站点表"

    sales_file = fields.Binary(string="销售明细表", required=True, attachment=True)
    sales_file_name = fields.Char(string="销售表文件名")
    store_file = fields.Binary(string="门店信息表", required=True, attachment=True)
    store_file_name = fields.Char(string="门店表文件名")
    output_name = fields.Char(string="输出表名", default="排线站点表")
    state = fields.Selection([("draft", "草稿"), ("processing", "处理中"), ("done", "完成")], default="draft")
    log = fields.Text(string="处理日志", readonly=True)
    station_count = fields.Integer(string="站点数", readonly=True)
    result_batch_id = fields.Many2one("logistics.route.planning.batch", string="结果批次", readonly=True)

    def _read_excel(self, data, sheet_idx=0):
        try:
            import openpyxl
        except ImportError:
            raise UserError(_("服务器未安装 openpyxl, 请联系管理员。"))
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = wb[wb.sheetnames[sheet_idx]]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise UserError(_("Excel 为空"))
        headers = [str(c) if c else "" for c in rows[0]]
        return headers, rows[1:]

    def _geocode(self, address, city="广东省"):
        """调用高德 API 获取经纬度"""
        params = urllib.parse.urlencode({
            "key": AMAP_KEY, "address": address, "city": city, "output": "json"
        })
        url = f"{AMAP_GEOCODE_URL}?{params}"
        try:
            req = urllib.request.Request(url)
            with urllib.request.urlopen(req, timeout=5) as resp:
                data = json.loads(resp.read().decode())
                if data.get("status") == "1" and data.get("geocodes"):
                    loc = data["geocodes"][0]["location"].split(",")
                    return float(loc[0]), float(loc[1])
        except Exception as e:
            _logger.warning("Geocode failed for %s: %s", address, str(e))
        return 0.0, 0.0

    def action_process(self):
        if not self.sales_file or not self.store_file:
            raise UserError(_("请上传销售明细表和门店信息表"))
        self.write({"state": "processing", "log": "开始处理...\n"})

        # 1) 解析销售明细表
        sales_data = base64.b64decode(self.sales_file)
        headers, rows = self._read_excel(sales_data)
        store_col = None
        sales_cols = {}
        for i, h in enumerate(headers):
            if "门店" in h or "店铺" in h or "名称" in h:
                store_col = i
            if "单据日期" in h:
                sales_cols["date"] = i
            if "箱数" in h:
                sales_cols["boxes"] = i
            if "重量" in h and "kg" in h.lower():
                sales_cols["weight"] = i

        if store_col is None:
            raise UserError(_("销售表未找到门店/店铺列, 表头: %s") % str(headers[:10]))

        # 汇总: 门店 → 总重量/箱数
        from collections import defaultdict
        store_demand = defaultdict(lambda: {"weight": 0.0, "boxes": 0})
        for row in rows:
            if not row or len(row) <= store_col:
                continue
            name = str(row[store_col]).strip() if row[store_col] else ""
            if not name:
                continue
            wt = float(str(row[sales_cols.get("weight", 0)] or 0)) if sales_cols.get("weight") is not None else 0
            bx = int(float(str(row[sales_cols.get("boxes", 0)] or 0))) if sales_cols.get("boxes") is not None else 0
            store_demand[name]["weight"] += wt
            store_demand[name]["boxes"] += bx

        log = self.log + f"解析完成: {len(store_demand)} 个门店\n"
        self.write({"log": log})

        # 2) 解析门店信息表, geocode
        store_data = base64.b64decode(self.store_file)
        h2, r2 = self._read_excel(store_data)
        name_col = addr_col = None
        for i, h in enumerate(h2):
            if "名称" in h or "门店" in h or "客户" in h:
                name_col = i
            if "地址" in h:
                addr_col = i
        if name_col is None:
            raise UserError(_("门店表未找到名称列: %s") % str(h2[:10]))

        geocoded = {}
        total = min(len(r2), 500)  # 限制 500 条以免超时
        for idx, row in enumerate(r2[:total]):
            name = str(row[name_col]).strip() if row[name_col] else ""
            addr = str(row[addr_col]).strip() if addr_col is not None else ""
            if not name:
                continue
            search_addr = addr or name
            lng, lat = self._geocode(search_addr)
            geocoded[name] = {"lng": lng, "lat": lat, "address": addr}
            if (idx + 1) % 20 == 0:
                self.write({"log": log + f"Geocode: {idx+1}/{total}\n"})

        log = self.log + f"Geocode 完成: {len(geocoded)} 条\n"
        self.write({"log": log})

        # 3) 合并生成站点表 → route.planning.batch
        batch = self.env["logistics.route.planning.batch"].create({
            "batch_no": f"STATION-{fields.Date.today().strftime('%Y%m%d')}-{fields.Datetime.now().strftime('%H%M')}",
            "delivery_date": fields.Date.today(),
            "entry_mode": "import_sheet",
            "planning_state": "result_ready",
            "route_status": "route_planned",
            "warehouse_id": 1,
        })

        seq = 0
        for store_name, demand in store_demand.items():
            geo = geocoded.get(store_name, {})
            lng = geo.get("lng", 0) or 0
            lat = geo.get("lat", 0) or 0
            addr = geo.get("address", "")
            seq += 1
            self.env["logistics.route.planning.stop.line"].create({
                "batch_id": batch.id,
                "stop_seq": seq,
                "store_name": store_name,
                "waybill_no": f"YD-STN-{seq:04d}",
                "longitude": lng,
                "latitude": lat,
                "address_detail": addr,
                "goods_weight_total": demand["weight"],
                "package_count_total": demand["boxes"],
            })

        batch.stop_count = seq
        self.write({
            "state": "done",
            "station_count": seq,
            "result_batch_id": batch.id,
            "log": log + f"\n完成! {seq} 个站点已生成, 批次: {batch.batch_no}"
        })

        return {
            "type": "ir.actions.act_window",
            "name": _("排线调度"),
            "res_model": "logistics.route.planning.batch",
            "view_mode": "list,form",
            "domain": [],
        }
